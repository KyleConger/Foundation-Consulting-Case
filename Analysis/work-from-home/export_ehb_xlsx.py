# -*- coding: utf-8 -*-
"""Export WFH chart/table contents to EH&B Excel workbook. Read-only on CSVs."""
import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"c:\Users\Owner\Desktop\MAN6930 Case")
OUT = ROOT / "Analysis" / "work-from-home" / "out"
DEST_DIR = ROOT / "Decisions" / "EH&B"
DEST = DEST_DIR / "Work from Home Rates.xlsx"

DEST_DIR.mkdir(parents=True, exist_ok=True)

with open(OUT / "summary.json", encoding="utf-8") as f:
    summary = json.load(f)


def load_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


acs_rows = load_csv(OUT / "acs2024_wfh_by_state.csv")
atus_rows = load_csv(OUT / "atus2024_work_at_home_breaks.csv")
abs_rows = load_csv(OUT / "abs2023_workhome_by_firm_size.csv")

us = next(r for r in acs_rows if r["geo_level"] == "nation")
states = [r for r in acs_rows if r["geo_level"] == "state"]
states_sorted = sorted(states, key=lambda r: -float(r["wfh_rate_pct"]))

cite_acs = summary["citations"]["acs"]
cite_atus = summary["citations"]["atus"]
cite_abs = summary["citations"]["abs"]

header_font = Font(bold=True)
header_fill = PatternFill("solid", fgColor="F2F2F2")
wrap = Alignment(wrap_text=True, vertical="top")
pct_1dp = "0.0"
int_fmt = "#,##0"
full_prec = "0.000000000000000"


def style_header(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def round1(x):
    return round(float(x), 1)


def as_int(x):
    if x is None or x == "":
        return None
    return int(float(x))


def as_float(x):
    if x is None or x == "":
        return None
    return float(x)


def display_geo(name):
    return "D.C." if name == "District of Columbia" else name


wb = Workbook()

# ----- 1. Notes — three definitions side by side -----
ws = wb.active
ws.title = "Notes"
ws.append(
    [
        "Field",
        "ACS 2024 (usually WFH)",
        "ATUS 2024 (any work at home on days worked)",
        "ABS 2023 (share of firms with any WFH employees)",
    ]
)

notes_rows = [
    (
        "Definition (do not mix)",
        "Percent of workers 16+ whose usual means of transportation to work was "
        "“Worked from home” (journey-to-work). Headline US rate: 13.3%.",
        "Percent of employed persons (15+) who did any work at home on days they "
        "worked. US overall on that question: 32.5%. Seniority figures are proxies "
        "(occupation, earnings, education), not job titles.",
        "Percent of employer FIRMS that had any employees who worked from home "
        "(EMPSZFI <500 vs 500+). Not a worker rate. Firms <500: 35.8%; firms 500+: 77.5%.",
    ),
    ("Agency", cite_acs["agency"], cite_atus["agency"], cite_abs["agency"]),
    ("Program", cite_acs["program"], cite_atus["program"], cite_abs["program"]),
    ("Table / file", cite_acs["table"], cite_atus["table"], cite_abs["table"]),
    ("Year", cite_acs["year"], cite_atus["year"], cite_abs["year"]),
    (
        "Exact question / variable",
        cite_acs["variable_or_question"],
        cite_atus["variable_or_question"],
        cite_abs["variable_or_question"],
    ),
    ("URL used", cite_acs["url"], cite_atus["url"], cite_abs["url"]),
    (
        "Secondary URL",
        cite_acs["url_secondary"],
        cite_atus["url_secondary"],
        cite_abs["url_secondary"],
    ),
    ("Funded by", cite_acs["funded_by"], cite_atus["funded_by"], cite_abs["funded_by"]),
    ("Origin", cite_acs["origin"], cite_atus["origin"], cite_abs["origin"]),
    ("Confidence", cite_acs["confidence"], cite_atus["confidence"], cite_abs["confidence"]),
    ("Notes", cite_acs["notes"], cite_atus["notes"], cite_abs["notes"]),
    (
        "Canvas",
        "Chtr Work from Home Rates (chtr-work-from-home-rates.canvas.tsx)",
        "Same canvas — ATUS section only",
        "Same canvas — ABS firm-size section only",
    ),
    (
        "Source CSVs / script",
        "Analysis/work-from-home/out/acs2024_wfh_by_state.csv · "
        "build_wfh_rates.py (CSV is source of truth if canvas differs)",
        "Analysis/work-from-home/out/atus2024_work_at_home_breaks.csv",
        "Analysis/work-from-home/out/abs2023_workhome_by_firm_size.csv",
    ),
    (
        "Sheet guide",
        "United States = ACS national rate. States = all 50+DC sorted by rate.",
        "Seniority proxies = ATUS breaks; column is the ATUS question.",
        "Firm size = ABS <500 vs 500+ share of firms (EMPSZFI 655 / 657).",
    ),
]
for row in notes_rows:
    ws.append(list(row))

style_header(ws, 4)
set_widths(ws, [28, 42, 48, 48])
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=4):
    for cell in row:
        cell.alignment = wrap
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 72 if r in (2, 7, 13) else 36

# ----- 2. United States -----
ws = wb.create_sheet("United States")
ws.append(
    [
        "Geography",
        "Usually WFH (numerator)",
        "Workers 16+ (denominator)",
        "Usually WFH % (exhibit, 1 decimal)",
        "Usually WFH % (full precision)",
        "Source URL",
        "Secondary URL",
        "Agency",
        "Program",
        "Table",
        "Year",
        "Question / variable",
    ]
)
ws.append(
    [
        us["geography"],
        as_int(us["worked_from_home"]),
        as_int(us["workers_16_plus"]),
        round1(us["wfh_rate_pct"]),
        float(us["wfh_rate_pct"]),
        us["url"],
        us["url_secondary"],
        us["agency"],
        us["program"],
        us["table"],
        us["year"],
        us["variable_or_question"],
    ]
)
for col in (2, 3):
    ws.cell(2, col).number_format = int_fmt
ws.cell(2, 4).number_format = pct_1dp
ws.cell(2, 5).number_format = full_prec
style_header(ws, 12)
set_widths(ws, [16, 22, 22, 18, 22, 55, 45, 22, 40, 40, 8, 55])
ws.row_dimensions[2].height = 60
for col in (6, 7, 12):
    ws.cell(2, col).alignment = wrap

# ----- 3. States -----
ws = wb.create_sheet("States")
headers = [
    "Rank (by WFH rate)",
    "State / DC",
    "Usually WFH % (exhibit, 1 decimal)",
    "Usually WFH % (full precision)",
    "Workers 16+ (base)",
    "Worked from home",
    "Source URL",
    "Secondary URL",
]
ws.append(headers)
for i, r in enumerate(states_sorted, 1):
    ws.append(
        [
            i,
            display_geo(r["geography"]),
            round1(r["wfh_rate_pct"]),
            float(r["wfh_rate_pct"]),
            as_int(r["workers_16_plus"]),
            as_int(r["worked_from_home"]),
            r["url"],
            r["url_secondary"],
        ]
    )
    row = ws.max_row
    ws.cell(row, 3).number_format = pct_1dp
    ws.cell(row, 4).number_format = full_prec
    for col in (5, 6):
        ws.cell(row, col).number_format = int_fmt
style_header(ws, len(headers))
set_widths(ws, [14, 22, 18, 22, 18, 16, 55, 45])

# ----- 4. Seniority proxies (ATUS) -----
ws = wb.create_sheet("Seniority proxies")
atus_question = (
    "Percent who did any work at home on days they worked "
    "(% of employed persons who worked on an average day)"
)
ws.append(
    [
        "Break type",
        "Group",
        "Source table",
        atus_question + " — exhibit (1 decimal)",
        atus_question + " — full precision",
        "Worked at home (000s)",
        "Worked on average day (000s)",
        "Total employed (000s)",
        "Seniority proxy?",
        "Source URL",
        "Secondary URL (PDF)",
        "Agency",
        "Year",
        "Notes",
    ]
)
# Story order: overall, mgmt, service, earnings, education, hours
# Story-first order by substring match, then remaining rows
story_substrings = [
    "Total, 15 years and over",
    "Management, business, and financial operations",
    "Professional and related",
    "Service",
    "highest quartile",
    "lowest quartile",
    "Bachelor's degree and higher",
    "High school graduates, no college",
    "Full-time workers",
    "Part-time workers",
]
ordered = []
seen = set()


def match_atus(substr: str):
    for r in atus_rows:
        if substr.lower() in r["group"].lower() and id(r) not in seen:
            # Prefer exact-ish "Service" over longer groups containing the word
            if substr.lower() == "service" and r["group"].strip().lower() != "service":
                continue
            return r
    return None


for substr in story_substrings:
    hit = match_atus(substr)
    if hit is not None:
        ordered.append(hit)
        seen.add(id(hit))
for r in atus_rows:
    if id(r) not in seen:
        ordered.append(r)
        seen.add(id(r))

for r in ordered:
    ws.append(
        [
            r["break_type"],
            r["group"],
            r["source_table"],
            round1(r["published_pct_one_decimal"]),
            float(r["pct_worked_at_home_of_those_who_worked"]),
            as_float(r["worked_at_home_000s"]),
            as_float(r["worked_on_average_day_000s"]),
            as_float(r["total_employed_000s"]),
            "Yes" if r["seniority_proxy"] in ("True", "true", True) else "No",
            r["url"],
            r["url_secondary"],
            r["agency"],
            r["year"],
            r["notes"],
        ]
    )
    row = ws.max_row
    ws.cell(row, 4).number_format = pct_1dp
    ws.cell(row, 5).number_format = full_prec
    for col in (6, 7, 8):
        ws.cell(row, col).number_format = "#,##0.0"
style_header(ws, 14)
set_widths(ws, [16, 48, 12, 18, 18, 16, 18, 16, 12, 50, 50, 22, 8, 50])
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=14, max_col=14):
    for cell in row:
        cell.alignment = wrap

# ----- 5. Firm size (ABS) -----
ws = wb.create_sheet("Firm size")
ws.append(
    [
        "EMPSZFI code",
        "Employment size of firm",
        "BUSCHAR",
        "Business characteristic",
        "% of employer firms (exhibit, 1 decimal)",
        "% of employer firms (full / recomputed)",
        "Employer firms",
        "Employees",
        "% of employees",
        "Note",
        "Source URL",
        "Secondary URL",
        "Agency",
        "Table",
        "Year",
        "Question / variable",
    ]
)
for r in abs_rows:
    note = (
        "Share of firms (not workers) that had any employees who worked from home"
        if r["buschar"] == "EWA"
        else (
            "Total reporting denominator for the firm-size bin"
            if r["buschar"] == "EWTR"
            else "Complement: firms that did not have WFH employees"
        )
    )
    recomputed = r.get("recomputed_pct_of_total_reporting") or ""
    full_val = float(recomputed) if recomputed not in ("", None) else float(r["pct_of_employer_firms"])
    ws.append(
        [
            r["empszfi"],
            r["employment_size_of_firm"],
            r["buschar"],
            r["business_characteristic"],
            round1(r["pct_of_employer_firms"]),
            full_val,
            as_float(r["employer_firms"]),
            as_float(r["employees"]),
            round1(r["pct_of_employees"]) if r.get("pct_of_employees") else None,
            note,
            r["url"],
            r["url_secondary"],
            r["agency"],
            r["table"],
            r["year"],
            r["variable_or_question"],
        ]
    )
    row = ws.max_row
    ws.cell(row, 5).number_format = pct_1dp
    ws.cell(row, 6).number_format = full_prec
    for col in (7, 8):
        ws.cell(row, col).number_format = int_fmt
    if ws.cell(row, 9).value is not None:
        ws.cell(row, 9).number_format = pct_1dp
style_header(ws, 16)
set_widths(ws, [12, 36, 10, 48, 16, 18, 14, 14, 14, 42, 50, 45, 22, 40, 8, 50])
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=10, max_col=16):
    for cell in row:
        cell.alignment = wrap

wb.save(DEST)
print("Saved:", DEST)
print("Sheets:", wb.sheetnames)

# ----- Verification -----
checks = []
checks.append(("ACS US %", round1(us["wfh_rate_pct"]), 13.3))
checks.append(("ACS US workers", as_int(us["workers_16_plus"]), 165360450))
checks.append(("ACS US WFH", as_int(us["worked_from_home"]), 22026372))

dc = next(r for r in states_sorted if r["geography"] == "District of Columbia")
ms = next(r for r in states_sorted if r["geography"] == "Mississippi")
checks.append(("DC %", round1(dc["wfh_rate_pct"]), 22.9))
checks.append(("DC base", as_int(dc["workers_16_plus"]), 388136))
checks.append(("MS %", round1(ms["wfh_rate_pct"]), 6.2))
checks.append(("MS base", as_int(ms["workers_16_plus"]), 1302251))

def atus_pct(group_substr):
    matches = [r for r in atus_rows if group_substr.lower() in r["group"].lower()]
    assert matches, group_substr
    return round1(matches[0]["published_pct_one_decimal"])


checks.append(("ATUS overall", atus_pct("Total, 15 years"), 32.5))
checks.append(("ATUS mgmt", atus_pct("Management, business"), 48.1))
def atus_pct_exact(group_exact):
    matches = [r for r in atus_rows if r["group"].strip().lower() == group_exact.lower()]
    assert matches, group_exact
    return round1(matches[0]["published_pct_one_decimal"])


checks.append(("ATUS service", atus_pct_exact("Service"), 10.5))
checks.append(("ATUS top earn", atus_pct("highest quartile"), 51.0))
checks.append(("ATUS bot earn", atus_pct("lowest quartile"), 13.3))
checks.append(("ATUS bach", atus_pct("Bachelor's degree and higher"), 50.0))
checks.append(("ATUS HS", atus_pct("High school graduates, no college"), 17.8))

ewa_lt = next(r for r in abs_rows if r["empszfi"] == "655" and r["buschar"] == "EWA")
ewa_ge = next(r for r in abs_rows if r["empszfi"] == "657" and r["buschar"] == "EWA")
checks.append(("ABS <500", round1(ewa_lt["pct_of_employer_firms"]), 35.8))
checks.append(("ABS 500+", round1(ewa_ge["pct_of_employer_firms"]), 77.5))

with open(DEST, "rb") as f:
    magic = f.read(2)
checks.append(("PK zip header", magic, b"PK"))

wb2 = load_workbook(DEST, data_only=True)
us_s = wb2["United States"]
checks.append(("xlsx US D2", us_s["D2"].value, 13.3))
st = wb2["States"]
checks.append(("xlsx DC B2", st["B2"].value, "D.C."))
checks.append(("xlsx DC C2", st["C2"].value, 22.9))
checks.append(("xlsx DC E2", st["E2"].value, 388136))
checks.append(("xlsx MS B52", st["B52"].value, "Mississippi"))
checks.append(("xlsx MS C52", st["C52"].value, 6.2))
checks.append(("xlsx MS E52", st["E52"].value, 1302251))

# Find ATUS rows in sheet by group substring
sp = wb2["Seniority proxies"]
sp_map = {}
for row in sp.iter_rows(min_row=2, max_row=sp.max_row, values_only=True):
    if row[1]:
        sp_map[row[1]] = row[3]

def find_sp(substr):
    for k, v in sp_map.items():
        if substr.lower() in k.lower():
            return v
    return None


checks.append(("xlsx ATUS overall", find_sp("Total, 15"), 32.5))
checks.append(("xlsx ATUS mgmt", find_sp("Management, business"), 48.1))
checks.append(("xlsx ATUS top", find_sp("highest quartile"), 51.0))
checks.append(("xlsx ATUS bot", find_sp("lowest quartile"), 13.3))
def find_sp_exact(name):
    for k, v in sp_map.items():
        if k.strip().lower() == name.lower():
            return v
    return None


checks.append(("xlsx ATUS service", find_sp_exact("Service"), 10.5))
checks.append(("xlsx ATUS bach", find_sp("Bachelor's degree and higher"), 50.0))
checks.append(("xlsx ATUS HS", find_sp("High school graduates, no college"), 17.8))

fs = wb2["Firm size"]
fs_vals = {}
for row in fs.iter_rows(min_row=2, max_row=fs.max_row, values_only=True):
    if row[2] == "EWA":
        fs_vals[row[0]] = row[4]
checks.append(("xlsx ABS 655", fs_vals.get("655"), 35.8))
checks.append(("xlsx ABS 657", fs_vals.get("657"), 77.5))

print("\nVerification:")
all_ok = True
for name, got, expected in checks:
    ok = got == expected
    if not ok:
        all_ok = False
    print(f"  [{'OK' if ok else 'FAIL'}] {name}: got={got!r} expected={expected!r}")
print("ALL PASS" if all_ok else "SOME FAILED")
print("File size:", DEST.stat().st_size)
print("Sheet names:", wb2.sheetnames)
