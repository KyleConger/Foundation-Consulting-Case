"""
Product mix, product trends, and unit economics that bear on EH&B.

EH&B here is a corporate/broker door that opens a billed residential
Spectrum Internet account. This script does not size that channel.
It recomputes shares from filing figures and leaves unpublished
unit-economics cells blank.

Inputs are transcribed from PrimarySources/readable and tagged
RETRIEVED. Shares and stock-chain checks are DERIVED in this file.
No take-up, subsidy, broker commission, CAC, or contribution per PSU.
"""

from __future__ import annotations

import csv
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"c:\Users\Owner\Desktop\MAN6930 Case")
OUT = ROOT / "Analysis" / "ehb-unit-economics"
MASTER = ROOT / "Decisions" / "EH&B" / "MASTER.xlsx"

# Preserve every sheet that existed before this append.
PRESERVE = [
    "TOC",
    "Emp-Notes",
    "Emp-US",
    "Emp-States",
    "Emp-MetrosHiLo",
    "Emp-MetrosFull",
    "Emp-MSAFull",
    "Emp-CutoffSens",
    "WFH-Notes",
    "WFH-US",
    "WFH-States",
    "WFH-Seniority",
    "WFH-FirmSize",
    "Brok-Notes",
    "Brok-Core",
    "Brok-Later",
    "Brok-InHand",
    "BrkDn-Notes",
    "BrkDn-Score",
    "BrkDn-M1",
    "BrkDn-M2",
    "BrkDn-M5",
    "BrkDn-M9",
    "BrkDn-M13",
    "BrkDn-Analogues",
    "CvC-Notes",
    "CvC-Leads",
    "CvC-Charter",
    "CvC-Comcast",
    "CvC-Neither",
    "Dens-Notes",
    "Dens-State",
    "Dens-Division",
    "Dens-A-CvC",
    "Dens-B-WFH",
    "Dens-C-Large",
    "State Combined",
]

NEW_SHEETS = ["Product Mix", "Product Trends", "Unit Economics"]

D = Decimal


def q1(x: Decimal) -> Decimal:
    """One decimal, half up, matching an Excel 0.0 display."""
    return x.quantize(D("0.1"), rounding=ROUND_HALF_UP)


def share(part: int, whole: int) -> Decimal:
    return D(part) / D(whole)


def pct_label(part: int, whole: int) -> str:
    return f"{q1(share(part, whole) * 100)}%"


# ---------------------------------------------------------------------------
# RETRIEVED inputs. Dollars in millions. Customers in thousands, as published.
# ---------------------------------------------------------------------------

# Q2 2026 residential product revenue that rolls into residential revenue,
# then the rest of company revenue. CHTR-Q2-2026-Earnings-Release-Ex99.1.md
Q2_REV = {
    "Internet": 5776,
    "Mobile service": 1095,
    "Video": 3149,
    "Voice": 331,
    "Commercial": 1865,
    "Advertising sales": 416,
    "Other": 894,
}
Q2_RESIDENTIAL = 10351  # Internet + mobile service + video + voice
Q2_TOTAL = 13526
Q2_CONNECTIVITY = 6871  # Internet + mobile service
Q2_SMALL = 1104
Q2_MID = 761

# Q1 2026. CHTR-Q1-2026-Earnings-Release-Ex99.1.md
Q1_REV = {
    "Internet": 5852,
    "Mobile service": 1052,
    "Video": 3252,
    "Voice": 338,
    "Commercial": 1839,
    "Advertising sales": 358,
    "Other": 906,
}
Q1_RESIDENTIAL = 10494
Q1_TOTAL = 13597
Q1_CONNECTIVITY = 6904

# FY2025. CHTR-FY2025-Earnings-Release-Ex99.1.md (full-year columns)
FY_REV = {
    "Internet": 23765,
    "Mobile service": 3762,
    "Video": 13703,
    "Voice": 1350,
    "Commercial": 7315,
    "Advertising sales": 1468,
    "Other": 3411,
}
FY_RESIDENTIAL = 42580
FY_TOTAL = 54774
FY_CONNECTIVITY = 27527

# Ending stock and quarterly net adds, residential, thousands.
# Source file is the Exhibit 99.1 that prints the quarter. 2024Q1–2024Q3
# are not in an Ex99.1 in this repo; both trending extracts agree.
QUARTERS = [
    "2024Q1",
    "2024Q2",
    "2024Q3",
    "2024Q4",
    "2025Q1",
    "2025Q2",
    "2025Q3",
    "2025Q4",
    "2026Q1",
    "2026Q2",
]

# Residential Internet ending customers and net adds (thousands).
INET_END = {
    "2024Q1": 28472,
    "2024Q2": 28318,
    "2024Q3": 28205,
    "2024Q4": 28034,
    "2025Q1": 27979,
    "2025Q2": 27868,
    "2025Q3": 27760,
    "2025Q4": 27641,
    "2026Q1": 27524,
    "2026Q2": 27358,
}
INET_NA = {
    "2024Q1": -72,
    "2024Q2": -154,
    "2024Q3": -113,
    "2024Q4": -171,
    "2025Q1": -55,
    "2025Q2": -111,
    "2025Q3": -108,
    "2025Q4": -119,
    "2026Q1": -117,
    "2026Q2": -166,
}

# Same print in CHTR-FY2025-Trending-Schedule.md (2024Q1–2025Q4).
INET_NA_FY_TREND = {
    "2024Q1": -72,
    "2024Q2": -154,
    "2024Q3": -113,
    "2024Q4": -171,
    "2025Q1": -55,
    "2025Q2": -111,
    "2025Q3": -108,
    "2025Q4": -119,
}

VID_NA = {
    "2024Q1": -392,
    "2024Q2": -393,
    "2024Q3": -281,
    "2024Q4": -110,
    "2025Q1": -167,
    "2025Q2": -73,
    "2025Q3": -64,
    "2025Q4": 49,
    "2026Q1": -51,
    "2026Q2": -11,
}
MOB_NA = {  # residential mobile LINES, not customers
    "2024Q1": 470,
    "2024Q2": 534,
    "2024Q3": 521,
    "2024Q4": 504,
    "2025Q1": 488,
    "2025Q2": 471,
    "2025Q3": 462,
    "2025Q4": 406,
    "2026Q1": 344,
    "2026Q2": 385,
}

# Residential voice ENDING customers (thousands). Net adds are not printed.
VOICE_END = {
    "2023Q4": 6712,  # FY2025 trending, so 2024Q1 change is defined
    "2024Q1": 6438,
    "2024Q2": 6170,
    "2024Q3": 5895,
    "2024Q4": 5636,
    "2025Q1": 5372,
    "2025Q2": 5161,
    "2025Q3": 4967,
    "2025Q4": 4832,
    "2026Q1": 4665,
    "2026Q2": 4494,
}

NA_SOURCE = {
    "2024Q1": "CHTR-Q1-2026-Trending-Schedule.md",
    "2024Q2": "CHTR-Q1-2026-Trending-Schedule.md",
    "2024Q3": "CHTR-Q1-2026-Trending-Schedule.md",
    "2024Q4": "CHTR-FY2025-Earnings-Release-Ex99.1.md",
    "2025Q1": "CHTR-Q1-2026-Earnings-Release-Ex99.1.md",
    "2025Q2": "CHTR-Q2-2026-Earnings-Release-Ex99.1.md",
    "2025Q3": "CHTR-FY2025-Earnings-Release-Ex99.1.md",
    "2025Q4": "CHTR-Q2-2026-Earnings-Release-Ex99.1.md",
    "2026Q1": "CHTR-Q2-2026-Earnings-Release-Ex99.1.md",
    "2026Q2": "CHTR-Q2-2026-Earnings-Release-Ex99.1.md",
}

# Stock at 30 Jun 2026. CHTR-Q2-2026-Earnings-Release-Ex99.1.md (thousands).
RES_REL_Q2 = 29276
SMB_REL_Q2 = 2223
TOTAL_REL_Q2 = 31499
PASSINGS_Q2 = 58981
PEN_PUBLISHED = D("53.4")  # percent, as printed
INET_RES_Q2 = 27358
VID_RES_Q2 = 12010
VOICE_RES_Q2 = 4494
MOB_LINES_RES_Q2 = 12099
INET_TOTAL_Q2 = 29388
INET_NA_TOTAL_Q2 = -172
INET_NA_TOTAL_Q1 = -120

# Blended monthly residential revenue per residential customer. Not Internet ARPU.
ARPU = {
    "2026Q2": D("117.52"),
    "2026Q1": D("118.44"),
    "2025Q4": D("117.19"),
    "2025Q2": D("119.70"),
}

# 10-K note 15 / FY Ex99.1. Millions.
PROG_FY25 = 8822
PROG_FY24 = 9653
OPEX_FY25 = 32739  # includes stock compensation $673
OPEX_FY24 = 33167
OPEX_COMPONENTS_FY25 = [8822, 6704, 5165, 3115, 3782, 673, 19, 4459]
INTEREST_FY25 = 5042  # interest expense, net
PRINCIPAL_B = D("93.8")
LEVERAGE = D("4.18")
CASH_M = 509

# Total Internet net adds the reader must not substitute for residential.
# Q1 residential -117 vs total -120; Q2 residential -166 vs total -172.


def voice_na(q: str) -> int:
    i = QUARTERS.index(q)
    prev = "2023Q4" if i == 0 else QUARTERS[i - 1]
    return VOICE_END[q] - VOICE_END[prev]


def check() -> None:
    assert sum(Q2_REV.values()) == Q2_TOTAL
    assert Q2_REV["Internet"] + Q2_REV["Mobile service"] == Q2_CONNECTIVITY
    assert (
        Q2_CONNECTIVITY + Q2_REV["Video"] + Q2_REV["Voice"] == Q2_RESIDENTIAL
    )
    assert Q2_SMALL + Q2_MID == Q2_REV["Commercial"]

    assert sum(Q1_REV.values()) == Q1_TOTAL
    assert Q1_REV["Internet"] + Q1_REV["Mobile service"] == Q1_CONNECTIVITY
    assert (
        Q1_CONNECTIVITY + Q1_REV["Video"] + Q1_REV["Voice"] == Q1_RESIDENTIAL
    )

    assert sum(FY_REV.values()) == FY_TOTAL
    assert FY_REV["Internet"] + FY_REV["Mobile service"] == FY_CONNECTIVITY
    assert (
        FY_CONNECTIVITY + FY_REV["Video"] + FY_REV["Voice"] == FY_RESIDENTIAL
    )

    assert sum(OPEX_COMPONENTS_FY25) == OPEX_FY25

    # Known anchors.
    assert q1(share(Q1_REV["Internet"], Q1_TOTAL) * 100) == D("43.0")
    assert q1(share(Q2_REV["Internet"], Q2_TOTAL) * 100) == D("42.7")
    assert INET_NA["2026Q1"] == -117
    assert INET_NA["2026Q2"] == -166
    assert INET_NA_TOTAL_Q1 == -120
    assert INET_NA_TOTAL_Q2 == -172
    # Residential + SMB = total on the two quarters the assignment names.
    assert INET_NA["2026Q1"] + (-3) == INET_NA_TOTAL_Q1
    assert INET_NA["2026Q2"] + (-6) == INET_NA_TOTAL_Q2

    for q, na in INET_NA_FY_TREND.items():
        assert INET_NA[q] == na, q

    # Ending stock equals prior ending plus this quarter's net adds.
    prev_end = INET_END["2024Q1"] - INET_NA["2024Q1"]
    for q in QUARTERS:
        assert INET_END[q] == prev_end + INET_NA[q], q
        prev_end = INET_END[q]
    assert sum(INET_NA[q] for q in QUARTERS) == -1186
    assert all(INET_NA[q] < 0 for q in QUARTERS)

    # Published penetration is the printed ratio, and it matches the counts
    # at one decimal. Passings include business sites; relationships include SMB.
    assert RES_REL_Q2 + SMB_REL_Q2 == TOTAL_REL_Q2
    recomputed_pen = q1(share(TOTAL_REL_Q2, PASSINGS_Q2) * 100)
    assert recomputed_pen == PEN_PUBLISHED == D("53.4")

    assert INET_RES_Q2 == INET_END["2026Q2"]
    assert q1(share(INET_RES_Q2, RES_REL_Q2) * 100) == D("93.4")
    assert q1(share(VID_RES_Q2, RES_REL_Q2) * 100) == D("41.0")
    assert q1(share(VOICE_RES_Q2, RES_REL_Q2) * 100) == D("15.4")

    assert q1(share(PROG_FY25, OPEX_FY25) * 100) == D("26.9")
    assert q1(share(PROG_FY24, OPEX_FY24) * 100) == D("29.1")

    # Q2 2026 total voice decline in the prose is 178 thousand =
    # residential derived −171 + small-business ending change −7.
    assert voice_na("2026Q2") == -171
    assert VOICE_END["2026Q1"] - VOICE_END["2026Q2"] == 171

    # Residential product shares of residential revenue round back to 100.0
    # at one decimal in Q2 and Q1.
    for rev, res in ((Q2_REV, Q2_RESIDENTIAL), (Q1_REV, Q1_RESIDENTIAL)):
        parts = ["Internet", "Mobile service", "Video", "Voice"]
        s = sum(q1(share(rev[p], res) * 100) for p in parts)
        assert s == D("100.0"), s

    q2_parts = list(Q2_REV)
    s = sum(q1(share(Q2_REV[p], Q2_TOTAL) * 100) for p in q2_parts)
    assert s == D("100.0"), s


def mix_rows() -> list[dict]:
    rows: list[dict] = []
    n = 1

    def add(period, product, revenue, company, residential, count, count_type, rel_base, tag, source, note):
        nonlocal n
        co = share(revenue, company) if revenue is not None and company else None
        rs = (
            share(revenue, residential)
            if revenue is not None and residential and product in ("Internet", "Mobile service", "Video", "Voice")
            else None
        )
        rel = (
            share(count, rel_base)
            if count is not None and rel_base and count_type == "residential customers"
            else None
        )
        rows.append(
            {
                "line_id": f"M{n}",
                "period": period,
                "product": product,
                "revenue_usd_millions": revenue if revenue is not None else "",
                "share_of_company_revenue": f"{co:.10f}" if co is not None else "",
                "share_of_company_revenue_1dp": f"{q1(co * 100)}%" if co is not None else "",
                "share_of_residential_revenue": f"{rs:.10f}" if rs is not None else "",
                "share_of_residential_revenue_1dp": f"{q1(rs * 100)}%" if rs is not None else "",
                "customers_or_lines_thousands": count if count is not None else "",
                "count_type": count_type,
                "share_of_residential_relationships": f"{rel:.10f}" if rel is not None else "",
                "share_of_residential_relationships_1dp": f"{q1(rel * 100)}%" if rel is not None else "",
                "tag": tag,
                "source_file": source,
                "note": note,
            }
        )
        n += 1

    q2_src = "CHTR-Q2-2026-Earnings-Release-Ex99.1.md"
    for product, rev in Q2_REV.items():
        count = None
        ctype = ""
        note = "Residential product line inside residential revenue." if product in (
            "Internet",
            "Mobile service",
            "Video",
            "Voice",
        ) else "Not a residential product line. Kept so Internet's share is of company revenue."
        if product == "Internet":
            count, ctype = INET_RES_Q2, "residential customers"
            note = "Revenue is the residential Internet line. Count is residential Internet customers at 30 Jun 2026. Not total Internet."
        elif product == "Video":
            count, ctype = VID_RES_Q2, "residential customers"
        elif product == "Voice":
            count, ctype = VOICE_RES_Q2, "residential customers"
        elif product == "Mobile service":
            count, ctype = MOB_LINES_RES_Q2, "residential mobile lines"
            note = "Service revenue only (devices sit in Other). Count is lines, not customers, so relationship share is blank."
        add(
            "Q2 2026",
            product,
            rev,
            Q2_TOTAL,
            Q2_RESIDENTIAL,
            count,
            ctype,
            RES_REL_Q2,
            "RETRIEVED",
            q2_src,
            note,
        )

    add(
        "Q2 2026",
        "Residential revenue (subtotal)",
        Q2_RESIDENTIAL,
        Q2_TOTAL,
        None,
        RES_REL_Q2,
        "residential relationships",
        None,
        "RETRIEVED",
        q2_src,
        "Internet + mobile service + video + voice. Relationship count is the denominator for product attachment, not a fifth product.",
    )
    add(
        "Q2 2026",
        "Total revenue",
        Q2_TOTAL,
        Q2_TOTAL,
        None,
        None,
        "",
        None,
        "RETRIEVED",
        q2_src,
        "Denominator for company-revenue shares.",
    )

    q1_src = "CHTR-Q1-2026-Earnings-Release-Ex99.1.md"
    for product in ("Internet", "Mobile service", "Video", "Voice"):
        add(
            "Q1 2026",
            product,
            Q1_REV[product],
            Q1_TOTAL,
            Q1_RESIDENTIAL,
            None,
            "",
            None,
            "RETRIEVED",
            q1_src,
            "Revenue only. Q1 Internet share of company revenue is the 43.0% anchor.",
        )
    add("Q1 2026", "Total revenue", Q1_TOTAL, Q1_TOTAL, None, None, "", None, "RETRIEVED", q1_src, "Denominator.")

    fy_src = "CHTR-FY2025-Earnings-Release-Ex99.1.md"
    for product in ("Internet", "Mobile service", "Video", "Voice"):
        add(
            "FY2025",
            product,
            FY_REV[product],
            FY_TOTAL,
            FY_RESIDENTIAL,
            None,
            "",
            None,
            "RETRIEVED",
            fy_src,
            "Full-year residential product revenue.",
        )
    add("FY2025", "Total revenue", FY_TOTAL, FY_TOTAL, None, None, "", None, "RETRIEVED", fy_src, "Denominator.")
    return rows


def trend_rows() -> list[dict]:
    rows = []
    for i, q in enumerate(QUARTERS, start=1):
        rows.append(
            {
                "line_id": f"T{i}",
                "quarter": q,
                "residential_internet_net_adds_thousands": INET_NA[q],
                "residential_video_net_adds_thousands": VID_NA[q],
                "residential_mobile_line_net_adds_thousands": MOB_NA[q],
                "residential_voice_net_adds_thousands": voice_na(q),
                "residential_internet_ending_thousands": INET_END[q],
                "tag_internet_video_mobile": "RETRIEVED",
                "tag_voice": "DERIVED",
                "source_file": NA_SOURCE[q],
                "note": (
                    "Voice net adds = change in published ending residential voice customers; "
                    "the filing does not print a voice net-add line. "
                    "Mobile is lines, not customers. Do not net lines against Internet customers. "
                    "Internet is residential, not total Internet."
                ),
            }
        )
    rows.append(
        {
            "line_id": "T11",
            "quarter": "Sum 2024Q1–2026Q2",
            "residential_internet_net_adds_thousands": sum(INET_NA.values()),
            "residential_video_net_adds_thousands": sum(VID_NA.values()),
            "residential_mobile_line_net_adds_thousands": sum(MOB_NA.values()),
            "residential_voice_net_adds_thousands": sum(voice_na(q) for q in QUARTERS),
            "residential_internet_ending_thousands": "",
            "tag_internet_video_mobile": "DERIVED",
            "tag_voice": "DERIVED",
            "source_file": "Sum of T1–T10",
            "note": "Ten-quarter sum. Not a net of mobile lines against Internet customers.",
        }
    )
    return rows


def unit_rows() -> list[dict]:
    def row(i, item, value, unit, period, tag, source, note):
        return {
            "line_id": f"U{i}",
            "item": item,
            "value": value,
            "unit": unit,
            "period": period,
            "tag": tag,
            "source_file": source,
            "note": note,
        }

    pen = share(TOTAL_REL_Q2, PASSINGS_Q2)
    prog25 = share(PROG_FY25, OPEX_FY25)
    prog24 = share(PROG_FY24, OPEX_FY24)
    blank = ""
    return [
        row(1, "Monthly residential revenue per residential customer", "117.52", "dollars per month", "Q2 2026", "RETRIEVED", "CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Total residential quarterly revenue / 3 / average residential customer relationships. Blended across products. Not Internet ARPU."),
        row(2, "Monthly residential revenue per residential customer", "118.44", "dollars per month", "Q1 2026", "RETRIEVED", "CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Same definition. Restated in the Q2 operating statistics."),
        row(3, "Monthly residential revenue per residential customer", "119.70", "dollars per month", "Q2 2025", "RETRIEVED", "CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Year-ago blended figure. The release states a 1.8% decline vs Q2 2026."),
        row(4, "Internet ARPU", blank, "dollars per month", "", "NOT IN FILING", "", "Not published. Do not use the blended residential figure as Internet ARPU."),
        row(5, "Video programming costs", PROG_FY25, "USD millions", "FY2025", "RETRIEVED", "CHTR-10-K-FY2025.md", "Note 15. FY Ex99.1 prints the same $8,822 million. 10-K prose says approximately $8.8 billion."),
        row(6, "Video programming costs", PROG_FY24, "USD millions", "FY2024", "RETRIEVED", "CHTR-10-K-FY2025.md", "Note 15. 10-K prose says approximately $9.7 billion."),
        row(7, "Total operating costs and expenses", OPEX_FY25, "USD millions", "FY2025", "RETRIEVED", "CHTR-10-K-FY2025.md", "Note 15 total, including stock compensation. Components sum to this figure."),
        row(8, "Total operating costs and expenses", OPEX_FY24, "USD millions", "FY2024", "RETRIEVED", "CHTR-10-K-FY2025.md", "Note 15 total."),
        row(9, "Programming share of operating costs", f"{prog25:.10f}", "share of opex", "FY2025", "DERIVED", "CHTR-10-K-FY2025.md", f"8822/32739 displays as {q1(prog25 * 100)}%. 10-K prose rounds to 27%."),
        row(10, "Programming share of operating costs", f"{prog24:.10f}", "share of opex", "FY2024", "DERIVED", "CHTR-10-K-FY2025.md", f"9653/33167 displays as {q1(prog24 * 100)}%. 10-K prose rounds to 29%."),
        row(11, "Interest expense, net", INTEREST_FY25, "USD millions", "FY2025", "RETRIEVED", "CHTR-10-K-FY2025.md", "Income statement. Equals $5.042 billion. Not an Internet cost allocation."),
        row(12, "Principal amount of debt", "93.8", "USD billions", "30 Jun 2026", "RETRIEVED", "CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Also stated in CHTR-10-Q-Q2-2026.md. Constraint on spending density as a consumer price war."),
        row(13, "Net debt / LTM Adjusted EBITDA", "4.18", "times", "30 Jun 2026", "RETRIEVED", "CHTR-10-Q-Q2-2026.md", "Charter's leverage ratio, as stated. Not recomputed here."),
        row(14, "Cash and cash equivalents", CASH_M, "USD millions", "30 Jun 2026", "RETRIEVED", "CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Balance sheet and the liquidity sentence."),
        row(15, "Estimated passings", PASSINGS_Q2, "thousands", "30 Jun 2026", "RETRIEVED", "CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Homes and business sites. Not a residential-only homes count."),
        row(16, "Customer-relationship penetration of passings", "0.534", "share", "30 Jun 2026", "RETRIEVED", "CHTR-Q2-2026-Earnings-Release-Ex99.1.md", f"Printed 53.4%. 31,499/58,981 displays as {q1(pen * 100)}%. Includes SMB customers. Not an employee-home take rate."),
        row(17, "Contribution per incremental residential Internet PSU", blank, "USD", "", "NOT IN FILING", "", "No dollar contribution per incremental PSU is in the filings. Left blank. Not estimated."),
        row(18, "CAC for a residential Internet PSU", blank, "USD", "", "NOT IN FILING", "", "Left blank. Not estimated."),
        row(19, "Employer or broker take-up of Internet-as-benefit", blank, "share", "", "NOT IN FILING", "", "Left blank. Not estimated. Channel M13 stays blank."),
        row(20, "Broker commission", blank, "USD", "", "NOT IN FILING", "", "Left blank. Not estimated."),
        row(21, "Subsidy tied to this channel", blank, "USD", "", "NOT IN FILING", "", "Left blank. Rural construction subsidies are a different channel and are not used here."),
    ]


def write_csv(name: str, rows: list[dict]) -> None:
    path = OUT / name
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def _fill_sheet(ws, title: str, source: str, headers: list[str], data: list[list], percent_cols: set[int]) -> None:
    ws.sheet_view.showGridLines = False
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    header_fill = PatternFill("solid", fgColor="F2F2F2")
    last = get_column_letter(len(headers))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    c1 = ws.cell(1, 1, title)
    c1.font = Font(bold=True, size=14, name="Calibri")
    c1.alignment = Alignment(wrap_text=True, vertical="center")
    c2 = ws.cell(2, 1, source)
    c2.font = Font(italic=True, size=10, name="Calibri", color="595959")
    c2.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 32
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(4, col, h)
        cell.font = Font(bold=True, size=10, name="Calibri")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{last}{4 + len(data)}"
    for r, row in enumerate(data, start=5):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(r, c, None if val == "" else val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=c == len(headers))
            if c in percent_cols and isinstance(val, float):
                cell.number_format = "0.0%"
            elif isinstance(val, float) and c not in percent_cols:
                cell.number_format = "#,##0.00"
            elif isinstance(val, int):
                cell.number_format = "#,##0"
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        width = 18
        if col == 1:
            width = 12
        if col == 2:
            width = 42
        if col == len(headers):
            width = 55
        ws.column_dimensions[letter].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = title
    ws.print_title_rows = "1:4"


def _num(v):
    if v == "" or v is None:
        return ""
    if isinstance(v, str) and v.replace(".", "", 1).isdigit():
        return float(v) if "." in v else int(v)
    return v


def append_master(mix, trends, units) -> tuple[int, int]:
    wb = load_workbook(MASTER)
    before = list(wb.sheetnames)
    kept = [s for s in before if s not in NEW_SHEETS]
    assert kept == PRESERVE, kept
    for name in NEW_SHEETS:
        if name in wb.sheetnames:
            del wb[name]
    # TOC row for this append — replace if re-run, else add under the catalog.
    toc = wb["TOC"]
    marker = "Product Mix, Product Trends, Unit Economics"
    existing = None
    for row in range(1, toc.max_row + 1):
        if toc.cell(row, 3).value == marker:
            existing = row
            break
    dest = existing or toc.max_row + 1
    toc.cell(dest, 1, "Analysis/ehb-unit-economics/ (appended; MASTER was not recompiled)")
    toc.cell(dest, 2, "2026-09-27")
    toc.cell(dest, 3, marker)
    toc.cell(dest, 4, "product_mix.csv, product_trends.csv, unit_economics.csv")
    note_row = dest + 1
    # Do not duplicate the note if the following row already carries it.
    note = "Appended 2026-09-27 from Analysis/ehb-unit-economics. _compile_master.py was not re-run. Earlier tabs are unchanged."
    if toc.cell(note_row, 1).value != note:
        # If we inserted by overwrite of an existing marker, the note may already sit below.
        if existing is None:
            toc.cell(note_row, 1, note)
    for cell_row in (dest, note_row):
        for col in range(1, 5):
            cell = toc.cell(cell_row, col)
            cell.font = Font(name="Calibri", size=10)

    mix_headers = [
        "Line",
        "Product",
        "Period",
        "Revenue ($ millions)",
        "Share of company revenue",
        "Share of residential revenue",
        "Customers or lines (000s)",
        "Count type",
        "Share of residential relationships",
        "Tag",
        "Source file",
        "Note",
    ]
    mix_data = []
    for r in mix:
        mix_data.append(
            [
                r["line_id"],
                r["product"],
                r["period"],
                _num(r["revenue_usd_millions"]),
                float(r["share_of_company_revenue"]) if r["share_of_company_revenue"] else "",
                float(r["share_of_residential_revenue"]) if r["share_of_residential_revenue"] else "",
                _num(r["customers_or_lines_thousands"]),
                r["count_type"],
                float(r["share_of_residential_relationships"]) if r["share_of_residential_relationships"] else "",
                r["tag"],
                r["source_file"],
                r["note"],
            ]
        )
    ws = wb.create_sheet("Product Mix")
    _fill_sheet(
        ws,
        "Residential Internet was 42.7% of Q2 2026 company revenue, and 93.4% of residential relationships already take it",
        "Sources: CHTR-Q2-2026-Earnings-Release-Ex99.1.md (Q2 revenue and 30 Jun 2026 customers); CHTR-Q1-2026-Earnings-Release-Ex99.1.md; CHTR-FY2025-Earnings-Release-Ex99.1.md. Shares recomputed in build_mix_trends_unit_economics.py. Percents display at one decimal. Analysis, not a recommendation to launch.",
        mix_headers,
        mix_data,
        percent_cols={5, 6, 9},
    )
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["K"].width = 48
    ws.column_dimensions["L"].width = 62

    tr_headers = [
        "Line",
        "Quarter",
        "Residential Internet net adds (000s)",
        "Residential video net adds (000s)",
        "Residential mobile line net adds (000s)",
        "Residential voice net adds (000s)",
        "Residential Internet ending (000s)",
        "Tag, Internet / video / mobile",
        "Tag, voice",
        "Source file",
        "Note",
    ]
    tr_data = []
    for r in trends:
        tr_data.append(
            [
                r["line_id"],
                r["quarter"],
                r["residential_internet_net_adds_thousands"],
                r["residential_video_net_adds_thousands"],
                r["residential_mobile_line_net_adds_thousands"],
                r["residential_voice_net_adds_thousands"],
                r["residential_internet_ending_thousands"] if r["residential_internet_ending_thousands"] != "" else "",
                r["tag_internet_video_mobile"],
                r["tag_voice"],
                r["source_file"],
                r["note"],
            ]
        )
    ws = wb.create_sheet("Product Trends")
    _fill_sheet(
        ws,
        "Residential Internet net adds stayed negative for 10 straight quarters, through −166,000 in Q2 2026",
        "Residential customers, thousands. 2024Q4–2026Q2 from Ex99.1 operating statistics (FY2025, Q1 2026, Q2 2026). 2024Q1–2024Q3 from CHTR-Q1-2026-Trending-Schedule.md, same prints in CHTR-FY2025-Trending-Schedule.md. Voice is DERIVED from ending stock. Mobile is lines. Not total Internet (−120k in Q1 2026; −172k in Q2 2026).",
        tr_headers,
        tr_data,
        percent_cols=set(),
    )
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 26
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["J"].width = 52
    ws.column_dimensions["K"].width = 62

    u_headers = ["Line", "Item", "Value", "Unit", "Period", "Tag", "Source file", "Note"]
    u_data = []
    for r in units:
        val = r["value"]
        if val != "" and val is not None and not isinstance(val, (int, float)):
            val = float(val)
        u_data.append(
            [r["line_id"], r["item"], val if val != "" else "", r["unit"], r["period"], r["tag"], r["source_file"], r["note"]]
        )
    ws = wb.create_sheet("Unit Economics")
    _fill_sheet(
        ws,
        "No contribution dollar per Internet PSU is published; principal is $93.8 billion at 4.18x",
        "CHTR-Q2-2026-Earnings-Release-Ex99.1.md; CHTR-10-Q-Q2-2026.md; CHTR-10-K-FY2025.md note 15 and the income statement. Blank rows are not zeros and not estimates. Analysis, not a recommendation to launch.",
        u_headers,
        u_data,
        percent_cols=set(),
    )
    # Format share rows (U9, U10, U16 are rows 13, 14, 20 on the sheet: header at row 4, data starts row 5 → U1 is row 5, so Un is row 4+n).
    for line_no in (9, 10, 16):
        cell = ws.cell(4 + line_no, 3)
        cell.number_format = "0.0%"
    # ARPU and leverage at two decimals; principal as published (93.8).
    for line_no in (1, 2, 3, 13):
        ws.cell(4 + line_no, 3).number_format = "0.00"
    ws.cell(4 + 12, 3).number_format = "0.0"
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["G"].width = 48
    ws.column_dimensions["H"].width = 70

    after_names = list(wb.sheetnames)
    assert after_names[: len(PRESERVE)] == PRESERVE
    assert after_names[-3:] == NEW_SHEETS
    wb.save(MASTER)
    return len(kept), len(after_names)


def main() -> None:
    check()
    mix = mix_rows()
    trends = trend_rows()
    units = unit_rows()
    write_csv("product_mix.csv", mix)
    write_csv("product_trends.csv", trends)
    write_csv("unit_economics.csv", units)

    summary = {
        "q2_internet_share_company": pct_label(Q2_REV["Internet"], Q2_TOTAL),
        "q2_video_share_company": pct_label(Q2_REV["Video"], Q2_TOTAL),
        "q2_mobile_share_company": pct_label(Q2_REV["Mobile service"], Q2_TOTAL),
        "q2_voice_share_company": pct_label(Q2_REV["Voice"], Q2_TOTAL),
        "q2_commercial_share_company": pct_label(Q2_REV["Commercial"], Q2_TOTAL),
        "q2_other_share_company": pct_label(Q2_REV["Other"], Q2_TOTAL),
        "q2_ad_share_company": pct_label(Q2_REV["Advertising sales"], Q2_TOTAL),
        "q2_internet_share_residential": pct_label(Q2_REV["Internet"], Q2_RESIDENTIAL),
        "q2_video_share_residential": pct_label(Q2_REV["Video"], Q2_RESIDENTIAL),
        "q2_mobile_share_residential": pct_label(Q2_REV["Mobile service"], Q2_RESIDENTIAL),
        "q2_voice_share_residential": pct_label(Q2_REV["Voice"], Q2_RESIDENTIAL),
        "q1_internet_share_company": pct_label(Q1_REV["Internet"], Q1_TOTAL),
        "fy_internet_share_company": pct_label(FY_REV["Internet"], FY_TOTAL),
        "fy_video_share_company": pct_label(FY_REV["Video"], FY_TOTAL),
        "fy_mobile_share_company": pct_label(FY_REV["Mobile service"], FY_TOTAL),
        "inet_of_relationships": pct_label(INET_RES_Q2, RES_REL_Q2),
        "video_of_relationships": pct_label(VID_RES_Q2, RES_REL_Q2),
        "voice_of_relationships": pct_label(VOICE_RES_Q2, RES_REL_Q2),
        "relationships_without_internet_thousands": RES_REL_Q2 - INET_RES_Q2,
        "prog_share_fy25": f"{q1(share(PROG_FY25, OPEX_FY25) * 100)}%",
        "prog_share_fy24": f"{q1(share(PROG_FY24, OPEX_FY24) * 100)}%",
        "inet_na": [INET_NA[q] for q in QUARTERS],
        "vid_na": [VID_NA[q] for q in QUARTERS],
        "mob_na": [MOB_NA[q] for q in QUARTERS],
        "voice_na": [voice_na(q) for q in QUARTERS],
        "ten_quarter_internet_sum": sum(INET_NA.values()),
        "q2_internet_na": INET_NA["2026Q2"],
        "q1_internet_na": INET_NA["2026Q1"],
    }
    (OUT / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    before, after = append_master(mix, trends, units)
    summary["master_sheets_before"] = before
    summary["master_sheets_after"] = after
    (OUT / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
