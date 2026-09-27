# -*- coding: utf-8 -*-
"""EH&B executive summary: HTML → Chrome PDF, and a critical-values workbook.

Reads Analysis/ehb-unit-economics/out/summary.json from recompute_rate_volume.py.
Does not rebuild MASTER.xlsx.
"""
from __future__ import annotations

import json
import subprocess
import time
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"c:\Users\Owner\Desktop\MAN6930 Case")
SUMMARY = ROOT / "Analysis" / "ehb-unit-economics" / "out" / "summary.json"
DEST = ROOT / "Decisions" / "EH&B"
PDF_PATH = DEST / "EH&B Executive Summary.pdf"
XLSX_PATH = DEST / "EH&B Executive Summary.xlsx"
HTML_PATH = Path(r"C:\Users\Owner\AppData\Local\Temp\chtr-ehb-executive-summary.html")
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

# Guard the typed typo. Never write "Summaryy".
assert "Summaryy" not in PDF_PATH.name
assert "Summaryy" not in XLSX_PATH.name

with open(SUMMARY, encoding="utf-8") as f:
    S = json.load(f)

P = S["printed_bridge"]
D = S["derived_bridge"]
PL = S["plant"]
FY = S["fy25"]
MOB = S["mobile"]


def esc(s: str) -> str:
    return (
        str(s)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def bar(label: str, width_pct: float, display: str, accent: bool) -> str:
    w = max(0.0, min(float(width_pct), 100.0))
    cls = "accent" if accent else ""
    return (
        f'<div class="bar-row"><div class="bar-label">{esc(label)}</div>'
        f'<div class="bar-track"><div class="bar-fill {cls}" style="width:{w:.2f}%"></div></div>'
        f'<div class="bar-val">{esc(display)}</div></div>'
    )


mix_bars = "\n".join(
    bar(
        r["line"],
        r["share_pct_1dp"],
        f'{r["share_pct_1dp"]:.1f}%',
        r["accent"],
    )
    for r in S["mix"]
)
# Scale the two bridge bars to the larger absolute piece so both start at zero.
bridge_max = max(abs(P["volume_m"]), abs(P["rate_mix_m"]))
vol_w = abs(P["volume_m"]) / bridge_max * 100
rate_w = abs(P["rate_mix_m"]) / bridge_max * 100

css = """
@page { size: letter; margin: 0.55in; }
* { box-sizing: border-box; }
body { font-family: "Segoe UI", Calibri, Arial, sans-serif; color: #1a1a1a; font-size: 10.5pt; line-height: 1.38; margin: 0; }
h1 { font-size: 16pt; margin: 6px 0 8px; font-weight: 650; line-height: 1.25; }
h2 { font-size: 12.5pt; margin: 14px 0 6px; font-weight: 650; page-break-after: avoid; }
p { margin: 4px 0; }
.muted { color: #444; }
.small, .src { font-size: 8pt; color: #444; line-height: 1.35; margin: 4px 0; }
.kicker { font-size: 8.5pt; letter-spacing: 0.04em; text-transform: uppercase; color: #666; margin: 0; }
.q { font-size: 9.5pt; color: #333; margin: 2px 0 8px; }
.callout { border: 1px solid #e0c88a; background: #fff8e8; padding: 8px 10px; margin: 8px 0 10px; page-break-inside: avoid; }
.callout-title { font-weight: 650; margin-bottom: 3px; }
.stats { display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 8px; margin: 8px 0; }
.stat { border: 1px solid #e0e0e0; padding: 8px 10px; page-break-inside: avoid; }
.stat .v { font-size: 16pt; font-weight: 650; }
.stat.accent .v { color: #9b2c2c; }
.stat .l { font-size: 8pt; color: #555; }
.legend { font-size: 8.5pt; color: #555; margin: 2px 0 6px; }
.swatch { display: inline-block; width: 10px; height: 10px; margin-right: 4px; vertical-align: -1px; }
.swatch.accent { background: #9b2c2c; }
.swatch.gray { background: #b0b4b8; }
.bar-row { display: flex; align-items: center; margin: 3px 0; page-break-inside: avoid; }
.bar-label { width: 168px; font-size: 9pt; flex-shrink: 0; }
.bar-track { flex: 1; background: #eee; height: 14px; }
.bar-fill { height: 100%; background: #b0b4b8; }
.bar-fill.accent { background: #9b2c2c; }
.bar-val { width: 72px; text-align: right; font-size: 9pt; margin-left: 6px; }
table { width: 100%; border-collapse: collapse; font-size: 9pt; margin: 6px 0; }
th, td { border-bottom: 1px solid #e5e5e5; padding: 4px 6px; text-align: left; vertical-align: top; }
th { font-weight: 650; border-bottom: 1.5px solid #bbb; }
td.num, th.num { text-align: right; font-variant-numeric: tabular-nums; }
tr.accent td { background: #f8ecec; }
.break { page-break-before: always; }
"""

html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"/>
<title>EH&amp;B Executive Summary</title>
<style>{css}</style></head><body>
<p class="kicker">Employee Health &amp; Benefits · Charter executive talk track · Q1 2026</p>
<h1>Tracker first. One employer-paid Internet test only if the leak is fixed wireless or non-mover housing.</h1>
<div class="callout">
<div class="callout-title">One sentence a CEO can refuse</div>
This quarter, stand up a quarterly residential-Internet net-loss tracker on the core (who left and why), and only if the fill is fixed wireless or housing among non-movers, authorize a two-quarter test of one employer-paid residential Spectrum Internet account on plant already passed; refuse it if the home is off plant, the household already has Spectrum, or the line dies when employment ends.
</div>
<p class="small">A billed account is one residential Spectrum Internet subscription. The test count is new accounts opened, minus homes that already have Spectrum, minus lines that die when employment ends. No take-up rate is in the file.</p>

<h2>Internet is already 43% of Q1 2026 revenue, and it is the line that shrank (−$78 million)</h2>
<p class="q">CEO question: What am I actually trying to fix?</p>
<div class="stats">
  <div class="stat accent"><div class="v">43%</div><div class="l">Internet share of $13,597 million</div></div>
  <div class="stat accent"><div class="v">−$78 million</div><div class="l">Internet revenue vs Q1 2025</div></div>
  <div class="stat"><div class="v">−117,000</div><div class="l">Residential Internet customers</div></div>
</div>
<p class="legend"><span class="swatch accent"></span>Internet<span style="display:inline-block;width:14px"></span><span class="swatch gray"></span>Every other line, gray. Bars start at zero.</p>
{mix_bars}
<p class="src">Shares = line ÷ $13,597 million, one decimal. Video 23.9% (the talk track’s “about 24%”). Mobile service 7.7%. Internet −$78 million is $5,852 million versus $5,930 million (−1.3%). The −117,000 is residential Internet quarterly net additions; total Internet was −120,000. Source: Charter Ex99.1, first quarter 2026.</p>

<h2 class="break">Residential Internet dollars fell $78 million because volume took $87 million and rate put back $9 million</h2>
<p class="q">CEO question: What does this cost, and what do I give up?</p>
<p class="legend"><span class="swatch accent"></span>Volume, −$87 million (printed)<span style="display:inline-block;width:14px"></span><span class="swatch gray"></span>Rate and mix, +$9 million (printed). Length is millions, from zero.</p>
{bar("Smaller average base", vol_w, "−$87 million", True)}
{bar("Rate and product mix", rate_w, "+$9 million", False)}
<p class="small">Full-year 2025 check, printed in the 10-K: rate and mix +$785 million covered volume −$380 million, net +$405 million, on 393,000 fewer residential Internet customers. The first quarter is when the base won.</p>
<p class="small">The −$90 million / +$12 million split in the earlier talk track is derived, not printed. Holding last year’s implied rate (${D["arpu_25"]:.2f}) on this year’s average base costs ${D["volume_m"]:.1f} million; the rate change puts back ${D["rate_mix_m"]:.1f} million. Implied rate ${D["arpu_26"]:.2f} versus ${D["arpu_25"]:.2f}. About ${D["arpu_gap_per_month"]:.2f} more per month (${D["arpu_hold_flat"]:.2f}) would have held the dollars flat on that same identity. Those figures do not replace the printed −$87 / +$9.</p>
<p class="src">Printed bridge: Charter Form 10-Q, quarter ended 31 March 2026, “Internet revenues from our residential customers.” FY2025 bridge: Form 10-K, year ended 31 December 2025. Derived identity recomputed in Analysis/ehb-unit-economics/recompute_rate_volume.py. Quarterly customer loss −117,000 is not the 10-Q’s −455,000 year-over-year change in ending customers.</p>

<h2 class="break">Company take is 54.0% of 58.7 million passings; rural take is 38.1% after $7.7 billion</h2>
<p class="q">CEO question: What do I give up?</p>
<p class="legend"><span class="swatch gray"></span>Company penetration<span style="display:inline-block;width:14px"></span><span class="swatch accent"></span>Rural penetration on already-lit passings. Bars start at zero.</p>
{bar("Company, Q1 2026", PL["q1_penetration_pct"], "54.0%", False)}
{bar("Rural lit plant, Q1 2026", PL["rural_penetration_pct"], "38.1%", True)}
<p class="small">$93.8 billion principal and 4.18× net leverage at 30 June 2026. This density is not a consumer price match, and it is not a reason to build more rural miles.</p>
<p class="small">Later public fact, same story: second-quarter penetration 53.4% on 58.981 million passings (58,981 thousand). Contribution per incremental Internet account: not disclosed. Consumer acquisition cost: not disclosed. Left blank.</p>
<p class="src">Company 54.0% is {PL["q1_customer_relationships_k"]:,} thousand relationships ÷ {PL["q1_passings_k"]:,} thousand estimated passings (Ex99.1, 31 March 2026). “Nearly 59 million” is the release’s prose; the table is 58.661 million. Rural 38.1% is {PL["rural_cr_k"]:,} thousand relationships ÷ {PL["rural_passings_k"]:,} thousand passings (Q1 2026 trending; the talk track rounded this to 38%). $7.7 billion: FY2025 10-K, subsidized rural construction since the start of 2022 through 31 December 2025. Principal: Q2 2026 Ex99.1. Leverage: Q2 2026 10-Q.</p>

<h2 class="break">After the tracker, test one employer-paid account for two quarters only if the leak is fixed wireless or non-mover housing</h2>
<p class="q">CEO question: What do you want me to do, who does it, and by when?</p>
<table>
<thead><tr><th>Loss reason</th><th>Fit</th><th>What the room should hear</th></tr></thead>
<tbody>
<tr class="accent"><td>Fixed wireless</td><td>Best</td><td>Authorize the two-quarter test only if this is the fill</td></tr>
<tr><td>Fiber</td><td>Works</td><td>Team score. An earlier check said no. Do not lead with it.</td></tr>
<tr><td>Satellite</td><td>No</td><td>Refuse the employer test</td></tr>
<tr class="accent"><td>Housing, household did not move</td><td>Best</td><td>Authorize the test only for non-movers, on plant already passed</td></tr>
<tr><td>Non-pay</td><td>Works</td><td>Team score. An earlier check said no. Do not fire it the same quarter.</td></tr>
<tr><td>No dominant reason</td><td>No</td><td>Refuse the employer test</td></tr>
</tbody></table>
<p class="src">Team scores from Decisions/CHURN. Accent only on the two Best rows. Fiber and non-pay stay “Works,” with the footnote that an earlier validation said no.</p>
<table>
<thead><tr><th>Gate</th><th>Line</th></tr></thead>
<tbody>
<tr><td>Capacity</td><td>The miles are already built.</td></tr>
<tr><td>Capability</td><td>Unknown whether a broker can put Internet on a benefits ballot. That is the first gate. The cell is blank.</td></tr>
<tr><td>Calendar</td><td>Tracker question this week. First fill is one quarter. The test is two quarters after that fill, not before.</td></tr>
<tr><td>Commitment</td><td>Say no if the fill is not fixed wireless or non-mover housing.</td></tr>
</tbody></table>
<p>Operations and marketing fill the tracker. Sales operations run one employer test only if the cell is fixed wireless or non-mover housing. It does not close the 117,000. No take-up is named.</p>

<h2 class="break">Do not claim</h2>
<table>
<thead><tr><th>Figure people reach for</th><th>What it is not</th></tr></thead>
<tbody>
<tr><td>94.6% of jobs in Charter-present states</td><td>Not serviceable homes</td></tr>
<tr><td>54.1% of payroll at firms of 500 or more</td><td>Not covered lives</td></tr>
<tr><td>Top 10 brokers’ 44% of Form 5500 compensation</td><td>Not covered lives</td></tr>
<tr><td>13.3% ACS usually works from home</td><td>Not the ATUS 32.5% any-work-at-home rate</td></tr>
<tr><td>Charter leads 19 states, Comcast 27</td><td>Locations, not subscribers</td></tr>
<tr><td>Mobile +368,000 lines and +$138 million</td><td>Did not replace the Internet print of −117,000 customers and −$78 million</td></tr>
</tbody></table>
<p class="src">Mobile verified in Q1 2026 Ex99.1: total lines +368,000; mobile service revenue $1,052 million versus $914 million (+$138 million). The other refusals are locked in Decisions/EH&amp;B/README.md (SUSB 2022, ACS 2024, ATUS 2024, FCC BDC 31 Dec 2025, PlanOptica Form 5500 compensation, M5 jobs-in-footprint).</p>
</body></html>
"""

HTML_PATH.write_text(html, encoding="utf-8")

if PDF_PATH.exists():
    PDF_PATH.unlink()

cmd = [
    CHROME,
    "--headless=new",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={PDF_PATH}",
    HTML_PATH.as_uri(),
]
r = subprocess.run(cmd, capture_output=True, text=True)
print("chrome_rc", r.returncode)
if r.returncode != 0:
    print((r.stderr or "")[:800])
    raise SystemExit(r.returncode)
time.sleep(1.0)

# --- workbook ---
wb = Workbook()
thin = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)
header_fill = PatternFill("solid", fgColor="F3F3F3")
accent_fill = PatternFill("solid", fgColor="F8ECEC")
header_font = Font(bold=True)
wrap = Alignment(wrap_text=True, vertical="top")


def style_header(ws, ncol: int) -> None:
    for col in range(1, ncol + 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
        cell.border = thin
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_rows(ws, headers, rows, highlight=None) -> None:
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    for i, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(i, c, val)
            cell.alignment = wrap
            cell.border = thin
            if highlight and highlight(row):
                cell.fill = accent_fill
    style_header(ws, len(headers))
    for c in range(1, len(headers) + 1):
        letter = get_column_letter(c)
        width = 18
        for row in ws.iter_rows(min_col=c, max_col=c, max_row=min(ws.max_row, 40)):
            val = row[0].value
            if val is not None:
                width = max(width, min(62, len(str(val)) + 2))
        ws.column_dimensions[letter].width = width


notes = wb.active
notes.title = "Notes"
write_rows(
    notes,
    ["Item", "Text"],
    [
        [
            "Channel",
            "Employee Health & Benefits: a corporate or broker door that can open one residential Spectrum Internet account at the employee’s home. Productivity is the pitch, not the KPI. This is not the other use of EH&B in this project (existing homes and businesses), though the account has to land on plant already passed.",
        ],
        [
            "Recommendation",
            "This quarter, stand up a quarterly residential-Internet net-loss tracker on the core (who left and why), and only if the fill is fixed wireless or housing among non-movers, authorize a two-quarter test of one employer-paid residential Spectrum Internet account on plant already passed; refuse it if the home is off plant, the household already has Spectrum, or the line dies when employment ends.",
        ],
        [
            "Retrieved",
            "A figure printed in a filing or official table used here: Ex99.1, the 10-Q, the 10-K, or the Q1 2026 trending schedule.",
        ],
        [
            "Derived",
            "A figure computed from retrieved inputs. The implied monthly rate and the −$90 / +$12 split are derived. They are a cross-check. They do not replace the printed 10-Q bridge of −$87 million volume and +$9 million rate.",
        ],
        [
            "Blank on purpose",
            "Contribution per incremental Internet account and consumer acquisition cost are not in the filings. Those cells are empty. No take-up rate is estimated.",
        ],
        [
            "Arithmetic",
            "Analysis/ehb-unit-economics/recompute_rate_volume.py writes out/summary.json. This workbook reads that file.",
        ],
    ],
)
notes.row_dimensions[2].height = 48
notes.row_dimensions[3].height = 60

critical = [
    ["Company revenue", 13597, "Q1 2026", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Millions of dollars. Not a stock price."],
    ["Internet revenue", 5852, "Q1 2026", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Residential product line inside residential revenue. Not SMB Internet."],
    ["Internet revenue, prior year", 5930, "Q1 2025", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", ""],
    ["Internet revenue change", -78, "Q1 2026 vs Q1 2025", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Millions. Printed −1.3%. Not the derived −$90."],
    ["Internet share of revenue", round(5852 / 13597, 6), "Q1 2026", "Derived", "Ex99.1 dollars; share = 5,852 / 13,597", "Displays as 43%. Not a customer share."],
    ["Residential Internet net additions (000s)", -117, "Q1 2026", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Not total Internet (−120). Not the −455,000 year-over-year ending-customer change."],
    ["Volume, printed bridge", -87, "Q1 2026", "Retrieved", "PrimarySources/readable/CHTR-10-Q-Q1-2026.md", "Headline split. Do not replace with the derived −$90."],
    ["Rate and product mix, printed bridge", 9, "Q1 2026", "Retrieved", "PrimarySources/readable/CHTR-10-Q-Q1-2026.md", "Headline split. Do not replace with the derived +$12."],
    ["Implied Internet rate, this year", D["arpu_26"], "Q1 2026", "Derived", "Analysis/ehb-unit-economics/out/summary.json", "Dollars per month. Not a printed ARPU."],
    ["Implied Internet rate, prior year", D["arpu_25"], "Q1 2025", "Derived", "Analysis/ehb-unit-economics/out/summary.json", "Dollars per month. Not a printed ARPU."],
    ["Extra rate to hold dollars flat", D["arpu_gap_per_month"], "Q1 2026", "Derived", "Analysis/ehb-unit-economics/out/summary.json", "Dollars per month on the derived identity only. Not the 10-Q bridge."],
    ["FY2025 rate and mix", 785, "FY2025", "Retrieved", "PrimarySources/readable/CHTR-10-K-FY2025.md", "Residential Internet. Covered the volume decline that year."],
    ["FY2025 volume", -380, "FY2025", "Retrieved", "PrimarySources/readable/CHTR-10-K-FY2025.md", "Decrease in average residential Internet customers."],
    ["FY2025 residential Internet customers", -393000, "FY2025 vs FY2024", "Retrieved", "PrimarySources/readable/CHTR-10-K-FY2025.md", "Customers, not dollars. Not the Q1 −117,000."],
    ["Company penetration", 0.540, "31 Mar 2026", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Customer relationships / estimated passings. Not Internet-only take."],
    ["Estimated passings (000s)", 58661, "31 Mar 2026", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "58.661 million. The release also says nearly 59 million."],
    ["Rural penetration", 0.381, "Q1 2026", "Retrieved", "PrimarySources/CHTR-Q1-2026-Trending-Schedule.md", "Talk track rounded to 38%. 527 / 1,385 thousand."],
    ["Rural construction spend ($ billions)", 7.7, "Inception 2022 through 31 Dec 2025", "Retrieved", "PrimarySources/readable/CHTR-10-K-FY2025.md", "Cumulative spend. Not Q1 capex."],
    ["Debt principal ($ billions)", 93.8, "30 Jun 2026", "Retrieved", "PrimarySources/readable/CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Later public fact vs the Q1 course anchor."],
    ["Net leverage (times)", 4.18, "30 Jun 2026", "Retrieved", "PrimarySources/readable/CHTR-10-Q-Q2-2026.md", "Net debt / LTM Adjusted EBITDA. Not in the Q2 earnings headline table."],
    ["Q2 penetration", 0.534, "30 Jun 2026", "Retrieved", "PrimarySources/readable/CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Later public fact. 53.4% on 58,981 thousand passings."],
    ["Mobile line net additions (000s)", 368, "Q1 2026", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Total lines. Residential was +344. Did not replace the Internet print."],
    ["Mobile service revenue change", 138, "Q1 2026 vs Q1 2025", "Retrieved", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Millions. $1,052 vs $914. Did not replace −$78 million of Internet."],
    ["Contribution per incremental PSU", None, "—", "Not disclosed", "—", "Left blank. Do not invent."],
    ["Consumer acquisition cost", None, "—", "Not disclosed", "—", "Left blank. Do not invent."],
]
ws_c = wb.create_sheet("Critical values")
write_rows(ws_c, ["Metric", "Value", "Period", "Retrieved or derived", "Source file", "Do not confuse"], critical)
ws_c.column_dimensions["A"].width = 46
ws_c.column_dimensions["F"].width = 62

mix_rows = []
for r in S["mix"]:
    mix_rows.append([
        r["line"],
        r["revenue_m"],
        r["share"],
        "Accent" if r["accent"] else "Gray",
        "Retrieved dollars; share = dollars / 13,597",
    ])
ws_m = wb.create_sheet("Mix")
write_rows(
    ws_m,
    ["Line", "Q1 2026 revenue ($ millions)", "Share of $13,597 million", "Chart role", "Label"],
    mix_rows,
    highlight=lambda row: row[3] == "Accent",
)
ws_m["B11"] = "=SUM(B2:B9)"
ws_m["B11"].font = Font(bold=True)
ws_m["A11"] = "Sum (must equal company revenue)"
ws_m["C11"] = "=SUM(C2:C9)"

ws_r = wb.create_sheet("Rate vs volume")
rate_rows = [
    ["Printed volume", P["volume_m"], "Q1 2026", "Retrieved", "10-Q: decrease in average residential Internet customers", "Use this figure"],
    ["Printed rate and mix", P["rate_mix_m"], "Q1 2026", "Retrieved", "10-Q: increase related to rate and product mix", "Use this figure"],
    ["Printed net", P["net_m"], "Q1 2026", "Retrieved", "10-Q total of the two lines; also Ex99.1 $5,852 − $5,930", "Must equal −78"],
    ["FY2025 rate and mix", FY["rate_mix_m"], "FY2025", "Retrieved", "10-K", "Covered volume that year"],
    ["FY2025 volume", FY["volume_m"], "FY2025", "Retrieved", "10-K", ""],
    ["FY2025 net", FY["net_m"], "FY2025", "Retrieved", "785 + (−380)", ""],
    ["FY2025 customer change", FY["residential_internet_customers"], "FY2025", "Retrieved", "10-K: residential Internet customers decreased by 393,000", "Not Q1 −117,000"],
    ["Derived volume, unrounded $ millions", D["volume_m"], "Q1 2026", "Derived", D["formula_volume"], "Cross-check only. Rounds to −90. Does not override −87."],
    ["Derived rate, unrounded $ millions", D["rate_mix_m"], "Q1 2026", "Derived", D["formula_rate"], "Cross-check only. Rounds to +12. Does not override +9."],
    ["Derived implied rate Q1 2026", D["arpu_26"], "Q1 2026", "Derived", D["formula_arpu"], "Dollars per month"],
    ["Derived implied rate Q1 2025", D["arpu_25"], "Q1 2025", "Derived", D["formula_arpu"], "Dollars per month"],
    ["Derived extra $/month to hold dollars flat", D["arpu_gap_per_month"], "Q1 2026", "Derived", D["formula_hold_flat"], "On the derived identity only"],
    ["Average residential Internet customers, Q1 2026 (000s)", D["avg_customers_26_k"], "Q1 2026", "Derived", D["formula_average"], "Beginning 27,641 thousand; ending 27,524 thousand"],
    ["Average residential Internet customers, Q1 2025 (000s)", D["avg_customers_25_k"], "Q1 2025", "Derived", D["formula_average"], "Beginning 28,034 thousand; ending 27,979 thousand"],
]
write_rows(
    ws_r,
    ["Piece", "Value", "Period", "Retrieved or derived", "Formula or source", "How to use"],
    rate_rows,
    highlight=lambda row: row[0].startswith("Printed"),
)
ws_r.column_dimensions["A"].width = 56
ws_r.column_dimensions["E"].width = 72

ws_p = wb.create_sheet("Plant and leverage")
plant_rows = [
    ["Company customer-relationship penetration", 0.540, "31 Mar 2026", "Retrieved", "Ex99.1 Q1 2026", "31,683 / 58,661 thousand"],
    ["Estimated passings (000s)", 58661, "31 Mar 2026", "Retrieved", "Ex99.1 Q1 2026", "58.661 million, not a silent 59"],
    ["Rural penetration of activated passings", 0.381, "Q1 2026", "Retrieved", "Q1 2026 trending schedule", "Printed 38.1%. Talk track said 38%."],
    ["Rural passings (000s)", 1385, "Q1 2026", "Retrieved", "Q1 2026 trending schedule", ""],
    ["Rural customer relationships (000s)", 527, "Q1 2026", "Retrieved", "Q1 2026 trending schedule", "527 / 1,385 = 38.1%"],
    ["Subsidized rural construction spend ($ billions)", 7.7, "2022 through 31 Dec 2025", "Retrieved", "FY2025 10-K", ""],
    ["Principal amount of debt ($ billions)", 93.8, "30 Jun 2026", "Retrieved", "Ex99.1 Q2 2026", ""],
    ["Net leverage (times Adjusted EBITDA)", 4.18, "30 Jun 2026", "Retrieved", "10-Q Q2 2026", ""],
    ["Q2 penetration", 0.534, "30 Jun 2026", "Retrieved", "Ex99.1 Q2 2026", "Later public fact"],
    ["Q2 estimated passings (000s)", 58981, "30 Jun 2026", "Retrieved", "Ex99.1 Q2 2026", "58.981 million"],
    ["Contribution per incremental Internet PSU", None, "—", "Not disclosed", "—", "Blank on purpose"],
    ["Consumer acquisition cost", None, "—", "Not disclosed", "—", "Blank on purpose"],
]
write_rows(
    ws_p,
    ["Metric", "Value", "Period", "Retrieved or derived", "Source", "Note"],
    plant_rows,
)

plays = [
    ["Fixed wireless", "Best", "Test", "Authorize only if the tracker fill is this reason", "Accent"],
    ["Fiber", "Works", "Do not lead", "Team score. Earlier validation said no.", "Gray"],
    ["Satellite", "No", "Refuse", "Not an employer-ballot play", "Gray"],
    ["Housing, non-mover", "Best", "Test", "Household that did not move; on plant already passed", "Accent"],
    ["Non-pay", "Works", "Do not lead", "Team score. Earlier validation said no.", "Gray"],
    ["No dominant reason", "No", "Refuse", "Null result", "Gray"],
]
ws_f = wb.create_sheet("Play fit")
write_rows(
    ws_f,
    ["Loss reason", "Fit", "Action", "Scope", "Chart role"],
    plays,
    highlight=lambda row: row[4] == "Accent",
)
ws_f["A9"] = "Does not close"
ws_f["B9"] = 117000
ws_f["C9"] = "No share of the Q1 residential Internet customer loss is claimed. No take-up is invented."
ws_f["A10"] = "Capability"
ws_f["B10"] = None
ws_f["C10"] = "Whether a broker can put Internet on a benefits ballot is unknown. Cell left blank."

claims = [
    ["94.6% of jobs in Charter-present states", "Not serviceable homes", "M5 jobs-in-footprint. Workplace state is not a billable home."],
    ["54.1% of payroll at firms of 500 or more", "Not covered lives", "SUSB 2022 employment mix. Not a broker book."],
    ["44% top-10 Form 5500 compensation", "Not covered lives", "PlanOptica compensation share. M2 lives are still blank."],
    ["13.3% ACS usually works from home", "Not ATUS 32.5%", "Different questions. Do not put them on one unlabeled bar."],
    ["FCC 19 vs 27 states", "Locations, not subscribers", "Where each company reports more broadband-serviceable places."],
    ["Mobile +368,000 lines and +$138 million", "Did not replace the Internet print", "Verified in Q1 2026 Ex99.1 before inclusion."],
]
ws_d = wb.create_sheet("Do not claim")
write_rows(ws_d, ["Figure", "What it is not", "Why"], claims)
ws_d.column_dimensions["A"].width = 52

# Headline cells that must match the filing.
assert ws_c["B2"].value == 13597
assert ws_c["B5"].value == -78
assert ws_c["B7"].value == -117
assert ws_c["B8"].value == -87
assert ws_c["B9"].value == 9
assert ws_m["B2"].value == 5852
assert ws_c["B6"].value == round(5852 / 13597, 6)

wb.save(XLSX_PATH)
print("PDF", PDF_PATH, PDF_PATH.stat().st_size)
print("XLSX", XLSX_PATH, XLSX_PATH.stat().st_size)
print("sheets", wb.sheetnames)
