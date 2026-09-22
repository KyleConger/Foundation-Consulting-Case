# -*- coding: utf-8 -*-
"""Export canvas chart/table contents to EH&B Excel workbook. Read-only on CSVs."""
import csv
import json
import statistics
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"c:\Users\Owner\Desktop\MAN6930 Case")
OUT = ROOT / "Analysis" / "small-large-employment" / "out"
DEST_DIR = ROOT / "Decisions" / "EH&B"
DEST = DEST_DIR / "Small Large Employment Shares.xlsx"

DEST_DIR.mkdir(parents=True, exist_ok=True)

with open(OUT / "summary.json", encoding="utf-8") as f:
    summary = json.load(f)


def load_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


us_rows = load_csv(OUT / "us_employment_by_firm_size.csv")
state_rows = load_csv(OUT / "state_employment_by_firm_size.csv")
metro_only = load_csv(OUT / "msa_metro_only_employment_by_firm_size.csv")
msa_all = load_csv(OUT / "msa_employment_by_firm_size.csv")

us = us_rows[0]
state_rows_sorted = sorted(state_rows, key=lambda r: -float(r["share_lt500_pct"]))

metro_sorted = sorted(
    [r for r in metro_only if r.get("share_lt500_pct") not in ("", None)],
    key=lambda r: -float(r["share_lt500_pct"]),
)
metro_high = metro_sorted[:10]
# Canvas lowest table: ascending share (Rochester first = most large-firm)
metro_low = list(reversed(metro_sorted[-10:]))

metro_shares = [float(r["share_lt500_pct"]) for r in metro_only if r.get("share_lt500_pct")]
metro_median_full = statistics.median(metro_shares)
metro_median_1dp = round(metro_median_full, 1)

header_font = Font(bold=True)
header_fill = PatternFill("solid", fgColor="F2F2F2")
wrap = Alignment(wrap_text=True, vertical="top")
pct_1dp = "0.0"
int_fmt = "#,##0"


def style_header(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def round1(x):
    return round(float(x), 1)


def as_int(x):
    if x is None or x == "":
        return None
    return int(float(x))


def short_metro_name(name):
    if name.endswith(" Metro Area"):
        return name[: -len(" Metro Area")]
    return name


wb = Workbook()

# ----- 1. Notes -----
ws = wb.active
ws.title = "Notes"
notes = [
    ("Title", "Where people work: small vs large firms — chart and table contents"),
    ("Canvas", "Chtr Small Large Employment Shares (chtr-small-large-employment-shares.canvas.tsx)"),
    ("Year", str(summary["year"])),
    ("Origin", summary["origin"]),
    ("Primary cutoff", summary["primary_cutoff"]),
    ("Definition", summary["definition"]),
    ("Measure", summary["measure"]),
    ("Vintage note", summary["vintage_note"]),
    ("Agency", summary["agency"]),
    ("Program", summary["program"]),
    ("Landing page", summary["source_page"]),
    ("US/State file", "us_state_naics_detailedsizes_2022.txt"),
    ("US/State URL", summary["state_table_url"]),
    ("US/State release", summary["state_release_note"]),
    ("MSA file", "msa_3digitnaics_2022.txt"),
    ("MSA URL", summary["msa_table_url"]),
    ("MSA release", summary["msa_release_note"]),
    ("Enterprise codes", summary["enterprise_codes_url"]),
    ("Full citation (US/State)", summary["citation_us_state"]),
    ("Full citation (MSA)", summary["citation_msa"]),
    (
        "Exclusions",
        "Excludes nonemployers, most government, and most of agriculture. "
        "Classified by enterprise (firm) size, not establishment size.",
    ),
    (
        "Metro median note",
        "Canvas metro median 49.8% = median of share_lt500 across 387 Metropolitan "
        "Statistical Areas (Micro Areas excluded). Reproduce on sheet 'Metros Full (387)' "
        "by taking the median of column 'Share <500 (%)'.",
    ),
    (
        "Sheet guide",
        "United States = national binary + size-class chart data. "
        "States = all 50+DC sorted by small-firm share (canvas table). "
        "Metros Exhibit High Low = canvas high/low tables only. "
        "Metros Full (387) = all Metro Areas (median universe). "
        "MSA Full File = metros + micros from msa_employment_by_firm_size.csv. "
        "Cutoff Sensitivity = <100 vs <500 vs 500+ chart.",
    ),
    (
        "Source CSVs",
        "Analysis/small-large-employment/out/*.csv · "
        "Script: Analysis/small-large-employment/build_susb_employment_shares.py",
    ),
    ("Charter footprint", summary["charter_footprint"]["reason"]),
]
ws.append(["Field", "Content"])
for field, content in notes:
    ws.append([field, content])
style_header(ws, 2)
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 110
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
    for cell in row:
        cell.alignment = wrap
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 45 if r >= 19 else 18

# ----- 2. United States -----
ws = wb.create_sheet("United States")
ws.append(
    [
        "Geography",
        "Employment <500",
        "Employment 500+",
        "Total employment",
        "Share <500 (%)",
        "Share 500+ (%)",
        "Share <500 (full precision)",
        "Share 500+ (full precision)",
        "Source URL",
    ]
)
ws.append(
    [
        us["geo_name"],
        as_int(us["employment_lt500"]),
        as_int(us["employment_ge500"]),
        as_int(us["employment_total"]),
        round1(us["share_lt500_pct"]),
        round1(us["share_ge500_pct"]),
        float(us["share_lt500_pct"]),
        float(us["share_ge500_pct"]),
        us["source_url"],
    ]
)
for col in (2, 3, 4):
    ws.cell(2, col).number_format = int_fmt
for col in (5, 6):
    ws.cell(2, col).number_format = pct_1dp
for col in (7, 8):
    ws.cell(2, col).number_format = "0.000000000000000"

ws.append([])
ws.append(
    [
        "Enterprise size class",
        "Employment",
        "Share of U.S. employment (%)",
        "Share (full precision)",
        "Notes",
    ]
)
size_classes = [
    ("<5", "employment_lt5", "share_lt5_pct", "Other size classes (canvas)"),
    ("5–9", "employment_5_9", "share_5_9_pct", "Other size classes (canvas)"),
    ("10–19", "employment_10_19", "share_10_19_pct", "Other size classes (canvas)"),
    ("20–99", "employment_20_99", "share_20_99_pct", "Other size classes (canvas)"),
    ("100–499", "employment_100_499", "share_100_499_pct", "Other size classes (canvas)"),
    ("500+", "employment_ge500", "share_ge500_pct", "Firms with 500+ employees (canvas focal)"),
]
start = 5
for i, (label, emp_key, share_key, note) in enumerate(size_classes):
    r = start + i
    ws.append([label, as_int(us[emp_key]), round1(us[share_key]), float(us[share_key]), note])
    ws.cell(r, 2).number_format = int_fmt
    ws.cell(r, 3).number_format = pct_1dp
    ws.cell(r, 4).number_format = "0.000000000000000"

ws.append([])
ws.append(["Citation", us["citation"]])
for col in range(1, 10):
    ws.cell(1, col).font = header_font
    ws.cell(1, col).fill = header_fill
for col in range(1, 6):
    ws.cell(4, col).font = header_font
    ws.cell(4, col).fill = header_fill
ws.freeze_panes = "A2"
set_widths(ws, [22, 18, 28, 22, 16, 16, 24, 24, 55])
ws.merge_cells("B12:I12")
ws["B12"].alignment = wrap
ws.row_dimensions[12].height = 60

# ----- 3. States -----
ws = wb.create_sheet("States")
headers = [
    "Rank (by small-firm share)",
    "State / DC",
    "Employment <500",
    "Employment 500+",
    "Total employment (base)",
    "Share <500 (%)",
    "Share 500+ (%)",
    "Share <500 (full precision)",
    "Share 500+ (full precision)",
    "Source URL",
]
ws.append(headers)
for i, r in enumerate(state_rows_sorted, 1):
    name = r["geo_name"]
    display = "D.C." if name == "District of Columbia" else name
    ws.append(
        [
            i,
            display,
            as_int(r["employment_lt500"]),
            as_int(r["employment_ge500"]),
            as_int(r["employment_total"]),
            round1(r["share_lt500_pct"]),
            round1(r["share_ge500_pct"]),
            float(r["share_lt500_pct"]),
            float(r["share_ge500_pct"]),
            r["source_url"],
        ]
    )
    row = ws.max_row
    for col in (3, 4, 5):
        ws.cell(row, col).number_format = int_fmt
    for col in (6, 7):
        ws.cell(row, col).number_format = pct_1dp
style_header(ws, len(headers))
set_widths(ws, [14, 22, 16, 16, 22, 14, 14, 24, 24, 55])

# ----- 4. Metros Exhibit -----
ws = wb.create_sheet("Metros Exhibit High Low")
ws.append(
    [
        "Exhibit group",
        "Metro Area",
        "Share <500 (%)",
        "Share <500 (full precision)",
        "Total employment",
        "Employment <500",
        "Employment 500+",
        "Share 500+ (%)",
        "Source URL",
    ]
)

for group, rows in (
    ("Highest small-firm shares", metro_high),
    ("Lowest small-firm shares", metro_low),
):
    for r in rows:
        ws.append(
            [
                group,
                short_metro_name(r["geo_name"]),
                round1(r["share_lt500_pct"]),
                float(r["share_lt500_pct"]),
                as_int(r["employment_total"]),
                as_int(r["employment_lt500"]),
                as_int(r["employment_ge500"]),
                round1(r["share_ge500_pct"]),
                r["source_url"],
            ]
        )
        row = ws.max_row
        ws.cell(row, 3).number_format = pct_1dp
        ws.cell(row, 8).number_format = pct_1dp
        for col in (5, 6, 7):
            ws.cell(row, col).number_format = int_fmt

style_header(ws, 9)
set_widths(ws, [28, 42, 14, 24, 16, 16, 16, 14, 55])

# ----- 5. Metros Full 387 -----
ws = wb.create_sheet("Metros Full (387)")
ws.append(
    [
        "Rank (by small-firm share)",
        "Metro Area",
        "CBSA code",
        "Share <500 (%)",
        "Share <500 (full precision)",
        "Share 500+ (%)",
        "Total employment",
        "Employment <500",
        "Employment 500+",
        "Source URL",
    ]
)
for i, r in enumerate(metro_sorted, 1):
    ws.append(
        [
            i,
            r["geo_name"],
            r["geo_id"],
            round1(r["share_lt500_pct"]),
            float(r["share_lt500_pct"]),
            round1(r["share_ge500_pct"]),
            as_int(r["employment_total"]),
            as_int(r["employment_lt500"]),
            as_int(r["employment_ge500"]),
            r["source_url"],
        ]
    )
    row = ws.max_row
    for col in (4, 6):
        ws.cell(row, col).number_format = pct_1dp
    for col in (7, 8, 9):
        ws.cell(row, col).number_format = int_fmt

note_row = ws.max_row + 2
ws.cell(note_row, 1, "Median note")
ws.cell(
    note_row,
    2,
    f"Median share <500 across these {len(metro_sorted)} Metro Areas = {metro_median_1dp}% "
    f"(full precision {metro_median_full}). Matches canvas Stat 'Metro median <500 share' 49.8%. "
    f"Filter: geo_type=metro from msa_metro_only_employment_by_firm_size.csv; Micro Areas excluded.",
)
ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=10)
ws.cell(note_row, 2).alignment = wrap
ws.row_dimensions[note_row].height = 48

style_header(ws, 10)
set_widths(ws, [12, 48, 12, 14, 24, 14, 16, 16, 16, 55])

# ----- 6. MSA Full File -----
ws = wb.create_sheet("MSA Full File")
ws.append(
    [
        "Geography type",
        "Area name",
        "CBSA code",
        "Share <500 (%)",
        "Share <500 (full precision)",
        "Share 500+ (%)",
        "Total employment",
        "Employment <500",
        "Employment 500+",
        "Source URL",
    ]
)
msa_sorted = sorted(
    msa_all,
    key=lambda r: (
        0 if r["geo_type"] == "metro" else 1,
        -float(r["share_lt500_pct"]) if r.get("share_lt500_pct") else 0,
    ),
)
for r in msa_sorted:
    share = r.get("share_lt500_pct")
    if share in ("", None):
        continue
    ws.append(
        [
            r["geo_type"],
            r["geo_name"],
            r["geo_id"],
            round1(share),
            float(share),
            round1(r["share_ge500_pct"]) if r.get("share_ge500_pct") else None,
            as_int(r["employment_total"]),
            as_int(r["employment_lt500"]),
            as_int(r["employment_ge500"]),
            r["source_url"],
        ]
    )
    row = ws.max_row
    for col in (4, 6):
        if ws.cell(row, col).value is not None:
            ws.cell(row, col).number_format = pct_1dp
    for col in (7, 8, 9):
        ws.cell(row, col).number_format = int_fmt
style_header(ws, 10)
set_widths(ws, [14, 48, 12, 14, 24, 14, 16, 16, 16, 55])

# ----- 7. Cutoff Sensitivity -----
ws = wb.create_sheet("Cutoff Sensitivity")
ws.append(
    [
        "Cutoff definition",
        "Employment",
        "Share of U.S. employment (%)",
        "Share (full precision)",
        "How derived",
    ]
)
for rowdata in (
    (
        "<100",
        as_int(us["employment_lt100"]),
        round1(us["share_lt100_pct"]),
        float(us["share_lt100_pct"]),
        "Sum of enterprise classes <5 + 5–9 + 10–19 + 20–99 ÷ U.S. total",
    ),
    (
        "<500 (SBA)",
        as_int(us["employment_lt500"]),
        round1(us["share_lt500_pct"]),
        float(us["share_lt500_pct"]),
        "Published cumulative class firms <500 ÷ U.S. total",
    ),
    (
        "500+",
        as_int(us["employment_ge500"]),
        round1(us["share_ge500_pct"]),
        float(us["share_ge500_pct"]),
        "Published / residual class firms 500+ ÷ U.S. total",
    ),
):
    ws.append(list(rowdata))
    row = ws.max_row
    ws.cell(row, 2).number_format = int_fmt
    ws.cell(row, 3).number_format = pct_1dp
ws.append([])
ws.append(["U.S. total employment", as_int(us["employment_total"])])
ws.cell(ws.max_row, 2).number_format = int_fmt
ws.append(["Citation", us["citation"]])
ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=5)
ws.cell(ws.max_row, 2).alignment = wrap
ws.row_dimensions[ws.max_row].height = 60
style_header(ws, 5)
set_widths(ws, [18, 16, 28, 24, 55])

wb.save(DEST)
print("Saved:", DEST)
print("Sheets:", wb.sheetnames)

# Verification
checks = []
checks.append(("US <500 %", round1(us["share_lt500_pct"]), 45.9))
checks.append(("US <500 emp", as_int(us["employment_lt500"]), 62250556))
checks.append(("US 500+ %", round1(us["share_ge500_pct"]), 54.1))
checks.append(("US 500+ emp", as_int(us["employment_ge500"]), 73497851))
mt = next(r for r in state_rows_sorted if r["geo_name"] == "Montana")
fl = next(r for r in state_rows_sorted if r["geo_name"] == "Florida")
checks.append(("Montana", round1(mt["share_lt500_pct"]), 66.3))
checks.append(("Florida", round1(fl["share_lt500_pct"]), 39.6))
checks.append(("Metro median", metro_median_1dp, 49.8))
checks.append(("Metro n", len(metro_only), 387))

with open(DEST, "rb") as f:
    magic = f.read(2)
checks.append(("PK zip header", magic, b"PK"))

wb2 = load_workbook(DEST, data_only=True)
us_sheet = wb2["United States"]
checks.append(("xlsx US B2 emp lt500", us_sheet["B2"].value, 62250556))
checks.append(("xlsx US C2 emp ge500", us_sheet["C2"].value, 73497851))
checks.append(("xlsx US E2 share", us_sheet["E2"].value, 45.9))
checks.append(("xlsx US F2 share", us_sheet["F2"].value, 54.1))
st = wb2["States"]
checks.append(("xlsx Montana F2", st["F2"].value, 66.3))
checks.append(("xlsx Florida F52", st["F52"].value, 39.6))
checks.append(("xlsx Montana B2", st["B2"].value, "Montana"))
checks.append(("xlsx Florida B52", st["B52"].value, "Florida"))

print("\nVerification:")
all_ok = True
for name, got, expected in checks:
    ok = got == expected
    if not ok:
        all_ok = False
    print(f"  [{'OK' if ok else 'FAIL'}] {name}: got={got!r} expected={expected!r}")
print("ALL PASS" if all_ok else "SOME FAILED")
print("File size:", DEST.stat().st_size)
print("Sheet names:", wb2.sheetnames)
