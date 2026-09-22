# -*- coding: utf-8 -*-
"""Export five-metric validation digest to Decisions/EH&B.

Writes:
  Decisions/EH&B/Brokerage Metrics Breakdown.pdf
  Decisions/EH&B/Brokerage Metrics Breakdown.xlsx

Does not modify Brokerage Metrics.pdf/xlsx or other EH&B exhibits.
Does not invent Charter take-up, serviceable %, broker lives by state,
or M13 channel counts — those stay blank / not retrieved.
"""
from __future__ import annotations

import csv
import json
import subprocess
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"c:\Users\Owner\Desktop\MAN6930 Case")
BASE = ROOT / "Analysis" / "ehb-broker-metrics"
DEST_DIR = ROOT / "Decisions" / "EH&B"
PDF_PATH = DEST_DIR / "Brokerage Metrics Breakdown.pdf"
XLSX_PATH = DEST_DIR / "Brokerage Metrics Breakdown.xlsx"
HTML_PATH = Path(r"C:\Users\Owner\AppData\Local\Temp\chtr-ehb-brokerage-metrics-breakdown.html")
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

DATE = "21 September 2026"

# --- retrieved / derived figures (must match canvases + analysis JSON/CSV) ---
N_PRESENT = 42
N_LEAD = 19
ESI_USA = 178.2
ESI_STATE_SUM = 178.4
ESI_PRESENT = 168.2
ESI_LEAD = 80.7
ESI_PRESENT_NOT_LEAD = 87.5
ESI_ABSENT = 10.2
ESI_PRESENT_PCT = 94.4
ESI_LEAD_PCT = 45.3

PO_TOP5 = 28.1
PO_TOP10 = 44.0
PO_TAIL = 56.0
PO_COUNT = "43,129"
BI_TOP5 = 46.7
BI_TOP10 = 67.2
BI_TOP100_B = 79.8
BI_TOP10_B = 53.6

US_EMP = 135_748_407
EMP_PRESENT = 128_435_628
EMP_ABSENT = 7_312_779
SHARE_PRESENT = 94.6
FABRIC_SHARE = 33.6

PLANT = [
    ("Hawaii", 92.3),
    ("N. Carolina", 79.0),
    ("New York", 60.2),
    ("Texas", 48.8),
    ("California", 47.0),
    ("Florida", 36.2),
    ("Nevada", 21.2),
    ("Arizona", 2.6),
    ("Maryland", 0.2),
    ("Rhode Island", 0.02),
]

# M9 analogue offer rates used on the PDF chart (all labeled analogues)
M9_OFFER_BARS = [
    ("Pet insurance (SHRM 2025)", 22.0, False),
    ("Transit subsidy (SHRM 2025)", 12.0, False),
    ("Parking subsidy (SHRM 2025)", 10.0, False),
    ("Home internet subsidy (SHRM 2017)", 5.0, True),
    ("QSLP 401(k) match (SHRM 2025)", 4.0, False),
]


def esc(s) -> str:
    return (
        str(s)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def load_csv(path: Path):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def load_json(path: Path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


m1_sum = load_json(BASE / "m1" / "summary.json")
m2_sum = load_json(BASE / "m2" / "out" / "summary.json")
m5_sum = load_json(BASE / "m5" / "out" / "summary.json")
m9_sum = load_json(BASE / "m9" / "summary.json")
m1_reg = load_csv(BASE / "m1" / "assumption_register.csv")
m2_reg = load_csv(BASE / "m2" / "out" / "assumption_register.csv")
m5_reg = load_csv(BASE / "m5" / "out" / "assumption_register.csv")
m13_reg = load_csv(BASE / "m13" / "assumption_register.csv")
analogues = load_csv(BASE / "m9" / "analogue_register.csv")
po_top10 = load_csv(BASE / "m2" / "out" / "planoptica_top10_ehb_compensation.csv")
bi_top10 = load_csv(BASE / "m2" / "out" / "bi_top10_us_brokerage_revenue_2024.csv")

assert m1_sum["m1_numerator"] is None
assert m1_sum["verdict"] == "valid-with-caveats"
assert m2_sum["lives_share_pct"] is None
assert abs(m2_sum["headline_proxy"]["value_pct"] - 44.0) < 1e-9
assert m5_sum["m5_value"] is None
assert round(m5_sum["share_employment_in_present_rounded_pct"], 1) == 94.6
assert m9_sum["invented_charter_takeup"] is False
assert m9_sum["verdict"] == "needs_split"

# ---------------------------------------------------------------------------
# PDF (HTML → Chrome print-to-pdf)
# ---------------------------------------------------------------------------

css = """
@page { size: letter; margin: 0.40in 0.46in; }
* { box-sizing: border-box; }
body { font-family: "Segoe UI", Calibri, Arial, sans-serif; color: #1a1a1a; font-size: 9.4pt; line-height: 1.32; margin: 0; }
h1 { font-size: 15pt; margin: 4px 0 6px; font-weight: 650; line-height: 1.22; }
h2 { font-size: 11.5pt; margin: 0 0 5px; font-weight: 650; page-break-after: avoid; }
h3 { font-size: 10pt; margin: 7px 0 3px; font-weight: 600; page-break-after: avoid; }
p { margin: 3px 0; }
.muted { color: #555; }
.small, .src { font-size: 7.1pt; color: #555; line-height: 1.28; margin: 2px 0 4px; word-break: break-word; }
.src { color: #444; }
.pills { display: flex; gap: 6px; flex-wrap: wrap; margin-bottom: 4px; }
.pill { border: 1px solid #c8c8c8; border-radius: 999px; padding: 2px 9px; font-size: 8pt; color: #444; background: #f5f5f5; }
.pill.info { background: #e8f1fb; border-color: #9bb8d9; color: #1a4a7a; }
.pill.warn { background: #fff4d6; border-color: #e0c88a; color: #6a4b00; }
.callout { border: 1px solid #d0d0d0; background: #fafafa; padding: 6px 9px; margin: 5px 0; border-radius: 4px; page-break-inside: avoid; }
.callout.warn { background: #fff8e8; border-color: #e0c88a; }
.callout.info { background: #f0f6fc; border-color: #b7cce3; }
.callout-title { font-weight: 650; margin-bottom: 2px; }
.stats { display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 7px; margin: 5px 0; }
.stat { border: 1px solid #e0e0e0; padding: 6px 8px; border-radius: 4px; page-break-inside: avoid; }
.stat .v { font-size: 13.5pt; font-weight: 650; }
.stat .l { font-size: 8pt; color: #666; }
.stat.info .v { color: #1a5fb4; }
.stat.warn .v { color: #8a5a00; }
.stat.danger .v { color: #a11; }
hr { border: none; border-top: 1px solid #ddd; margin: 10px 0; }
.page { page-break-before: always; }
table { width: 100%; border-collapse: collapse; font-size: 7.5pt; margin: 4px 0; }
th, td { border-bottom: 1px solid #e5e5e5; padding: 2px 4px; text-align: left; vertical-align: top; }
th { font-weight: 650; border-bottom: 1.5px solid #bbb; background: #f7f7f7; }
td.num, th.num { text-align: right; font-variant-numeric: tabular-nums; }
tr.accent td { background: #eef4fb; }
tr.warn td { background: #fff8e8; }
tr.blank td { background: #f3f3f3; color: #555; }
.bar-row { display: flex; align-items: center; margin: 1px 0; page-break-inside: avoid; }
.bar-label { width: 200px; font-size: 8pt; flex-shrink: 0; }
.bar-track { flex: 1; background: #eee; height: 11px; }
.bar-fill { height: 100%; background: #9aa0a6; }
.bar-fill.accent { background: #1a5fb4; }
.bar-val { width: 52px; text-align: right; font-size: 8.5pt; margin-left: 5px; font-variant-numeric: tabular-nums; }
.legend { display: flex; gap: 14px; font-size: 8pt; margin: 3px 0 5px; color: #555; }
.swatch { display: inline-block; width: 10px; height: 10px; margin-right: 3px; vertical-align: -1px; }
.swatch.gray { background: #9aa0a6; }
.swatch.blue { background: #1a5fb4; }
.two { display: grid; grid-template-columns: 1fr 1fr; gap: 10px; }
.card { border: 1px solid #ddd; padding: 7px 9px; border-radius: 4px; page-break-inside: avoid; }
.card-h { font-weight: 650; margin-bottom: 3px; }
.kicker { font-size: 8pt; color: #666; margin: 0 0 2px; }
"""


def bar_row(label: str, pct: float, accent: bool = False, scale: float = 100.0) -> str:
    width = max((pct / scale) * 100.0, 0.35)
    cls = "accent" if accent else ""
    val = f"{pct:.2f}%" if pct < 1 else f"{pct:.1f}%"
    return (
        f'<div class="bar-row"><div class="bar-label">{esc(label)}</div>'
        f'<div class="bar-track"><div class="bar-fill {cls}" style="width:{width:.2f}%"></div></div>'
        f'<div class="bar-val">{val}</div></div>'
    )


m1_bars = "\n".join(
    [
        bar_row(f"Charter-lead ({N_LEAD} states) — ESI persons", ESI_LEAD, False, 180),
        bar_row("Present, not lead (23 states) — ESI persons", ESI_PRESENT_NOT_LEAD, False, 180),
        bar_row("Charter absent (8 states + DC) — ESI persons", ESI_ABSENT, True, 180),
    ]
)

po_short = [
    ("MMA", 7.1),
    ("Gallagher Benefit", 5.9),
    ("Lockton", 5.8),
    ("Mercer Health", 5.0),
    ("USI", 4.3),
    ("WTW US", 4.3),
    ("Hub Midwest", 4.1),
    ("Aon Consulting", 2.8),
    ("Alliant", 2.5),
    ("Brown & Brown", 2.2),
]
m2_bars = "\n".join(bar_row(n, p, False, 60) for n, p in po_short)
m2_bars += "\n" + bar_row("All other 43,119 names", PO_TAIL, True, 60)

m5_bars = "\n".join(
    bar_row(n, p, accent=(n in ("Hawaii", "Rhode Island")), scale=100)
    for n, p in PLANT
)

m9_bars = "\n".join(bar_row(n, p, accent=acc, scale=30) for n, p, acc in M9_OFFER_BARS)

html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"/>
<title>Brokerage Metrics Breakdown — five validated gates</title>
<style>{css}</style></head><body>

<div class="pills">
  <span class="pill info">Validation digest</span>
  <span class="pill warn">Not a TAM</span>
  <span class="pill">Not a launch recommendation</span>
  <span class="pill">Blanks are intentional</span>
</div>
<p class="kicker">Employee Health &amp; Benefits broker channel · Charter / Spectrum · {DATE}</p>
<h1>Five gates are specified; none of them yet has a Charter channel number</h1>
<p class="muted">M1, M2, M5, and M13 are valid measurement contracts with the specified Charter figure still blank. M9 is a valid gate that two different rates were glued into — split who pays (M9a) from who enrolls (M9b) before anyone quotes take-up. This file synthesizes the five completed validations. It does not size a TAM and it does not invent take-up, serviceable-home percent, broker lives by state, or channel net adds.</p>

<div class="callout warn"><div class="callout-title">What this file will not fill</div>
Charter take-up of internet-as-benefit; serviceable-home percent of any book; broker-placed lives by state; M13 channel account counts. Those cells stay blank / not retrieved. Do not paste an analogue, a state employment share, or company Internet net adds into the blank.
</div>

<h2>Scorecard — what each validation lets you use</h2>
<table>
<thead><tr><th>Metric</th><th>Verdict</th><th>What you can use</th></tr></thead>
<tbody>
<tr class="warn"><td>M1 · Broker-placed covered lives in Spectrum states</td><td>Valid — numerator blank</td><td>National ESI scale (KFF / MEPS / EBSA) and Charter’s 42-state presence. The 168.2M / 178.2M ESI overlay is not M1. SUSB 45.9% is not M1.</td></tr>
<tr class="warn"><td>M2 · Large-broker concentration of that book</td><td>Valid — wrong unit</td><td>PlanOptica 44% top-10 of Form 5500 compensation; Business Insurance 67% top-10 of Top 100 revenue. Neither is lives.</td></tr>
<tr class="warn"><td>M5 · Serviceable-home share of enrolled employees</td><td>Valid — no number</td><td>The hard footprint gate. 94.6% of SUSB jobs in 42 Charter-present states is an ESTIMATED ceiling, not M5. Plant share runs Hawaii 92% to Rhode Island 0.02%.</td></tr>
<tr class="warn"><td>M9 · Payor model and enrollment take-up</td><td>Needs split M9a / M9b</td><td>Analogues only. Closest internet-as-benefit offer analogue is SHRM 2017 5% of employers. KFF 76% medical take-up is the wrong analogue. Charter take-up is blank.</td></tr>
<tr class="warn"><td>M13 · Net new residential Internet accounts</td><td>Valid — tighten to P−E−D</td><td>M13 = path-tagged provisions − existing/payer-only − same-window job-exit disconnects. Q1 2026 residential Internet −117k is company context, not this channel.</td></tr>
</tbody></table>
<p class="src">Sources: validation canvases chtr-ehb-m1-broker-covered-lives, chtr-ehb-m2-large-broker-concentration, chtr-ehb-m5-serviceable-home-share, chtr-ehb-m9-payor-takeup, chtr-ehb-m13-net-new-residential · parent register chtr-ehb-broker-channel-core-metrics · arithmetic in Analysis/ehb-broker-metrics/m1, m2, m5, m9, m13. Original register also in Decisions/EH&amp;B/Brokerage Metrics.xlsx (Core metrics). This digest does not overwrite that register.</p>

<div class="two">
  <div class="card"><div class="card-h">Naming collision</div>
  Repo folder Decisions/EH&amp;B means existing homes and businesses on plant already built. This digest means Employee Health &amp; Benefits (insurance). A broker channel that provisions residential Internet would feed existing-plant Internet net adds — it does not redefine the folder.</div>
  <div class="card"><div class="card-h">Cheapest kill still sits outside these five</div>
  The parent register’s cheapest stop is M10: in the pilot states, may the brokerage firm place or enroll Spectrum residential Internet as an employer benefit, or only refer? One counsel memo. These five metrics assume that question is still open.</div>
</div>

<!-- ===================== M1 ===================== -->
<div class="page">
<p class="kicker">M1 · Broker reach</p>
<h2>M1 is the right reach gate — the specified number is blank, and a “Spectrum state” worksite screen would barely fail anyone</h2>
<div class="stats">
  <div class="stat warn"><div class="v">Valid</div><div class="l">Verdict — numerator blank</div></div>
  <div class="stat warn"><div class="v">Not retrieved</div><div class="l">Broker-placed lives in Charter states</div></div>
  <div class="stat"><div class="v">94.4%</div><div class="l">FLAGGED: ESI persons living in Charter-present states — not M1</div></div>
</div>
<p>Definition: employees enrolled in broker-placed medical (or full benefits) at employers with at least one worksite in a Charter-present state, over that named broker’s US medical book. Unit: percent and absolute employees. Report dependents separately. Spectrum state = Charter-present on FCC BDC D25 (provider_id 130235): {N_PRESENT} states. Charter-lead = {N_LEAD} states. Absent = AK, IA, ND, OK, SD, AR, DE, UT + DC.</p>
<p class="muted">M1 gates whether a broker book even reaches employers Charter could put a residential Internet account in front of. It does not gate take-up, commissions, or net adds. The “any worksite in a Spectrum state” rule is a weak kill for a national broker: Charter reports some plant in {N_PRESENT} of 50 states.</p>

<div class="callout warn"><div class="callout-title">Do not substitute SUSB 45.9% or the 168.2M / 178.2M overlay for this metric</div>
Census SUSB 2022 payroll employment at firms under 500 is 45.9% of US employment — the labor market, not who buys medical through a broker. The ESI overlay below is all employer-sponsored insurance by residence, not a broker book.
</div>

<h3>Only {ESI_ABSENT} million ESI persons live outside Charter-present states — that overlay is not a kill, and it is not M1</h3>
<div class="legend"><span><span class="swatch gray"></span>ESI in Charter-present states (not M1)</span><span><span class="swatch blue"></span>ESI outside Charter-present states (not M1)</span></div>
{m1_bars}
<p class="src">Source: EBSA Health Insurance Coverage Bulletin Table 1A, CY2023 ESI persons (millions) × FCC BDC D25 Charter presence. Present {ESI_PRESENT}M ({ESI_PRESENT_PCT}% of printed USA {ESI_USA}M); lead {ESI_LEAD}M ({ESI_LEAD_PCT}%); absent {ESI_ABSENT}M. State cells sum to {ESI_STATE_SUM}M vs printed USA {ESI_USA}M (bulletin rounding). Residence is not worksite. ESI is not broker-placed. Arithmetic: Analysis/ehb-broker-metrics/m1/summary.json. EBSA: https://www.dol.gov/sites/dolgov/files/EBSA/researchers/data/health-and-welfare/health-insurance-coverage-bulletin-2024.pdf · FCC: https://broadbandmap.fcc.gov/data-download</p>

<h3>Closest retrieved figures — none of these is M1</h3>
<table>
<thead><tr><th>ID</th><th>Figure</th><th class="num">Value</th><th>Year</th><th>Origin</th><th>Why it is not M1</th></tr></thead>
<tbody>
<tr><td>C5</td><td>Private-sector workers enrolled at own employer</td><td class="num">65.6M</td><td>2024</td><td>RETRIEVED · AHRQ MEPS-IC</td><td>Own-employer enrollment, not broker-placed, not by state</td></tr>
<tr><td>C4</td><td>Workers covered through a current employer</td><td class="num">82.8M</td><td>CY2023</td><td>RETRIEVED · EBSA CPS bulletin</td><td>Employees, still all ESI, not broker-placed</td></tr>
<tr><td>C3</td><td>ESI from own job, under 65</td><td class="num">85.7M</td><td>Mar 2025</td><td>RETRIEVED · KFF / CPS ASEC</td><td>People, not a broker book</td></tr>
<tr><td>B1</td><td>Persons with any ESI</td><td class="num">178.2M</td><td>CY2023</td><td>RETRIEVED · EBSA Table 1A</td><td>Includes dependents, retiree, COBRA</td></tr>
<tr class="blank"><td>F1</td><td>Broker-placed lives in Charter-present states</td><td class="num">blank</td><td>—</td><td>NOT RETRIEVED</td><td>This is M1. Leave blank.</td></tr>
</tbody></table>
<p class="src">Also retrieved, still not M1 (full table in the Excel M1 sheet): KFF ACS nonelderly ESI 154M (2025); AHIP &gt;180M (AHIP-funded); Form 5500 ~88M participants (2023); Aon Health Solutions $3.335B FY2024. MEPS-IC: https://meps.ahrq.gov/data_files/publications/rf54/rf54.pdf · KFF 2025 EHBS: https://www.kff.org/health-costs/2025-employer-health-benefits-survey/ (Kaiser Family Foundation, not Kaiser Permanente; fielded by Davis Research) · Aon 10-K: https://www.sec.gov/Archives/edgar/data/315293/000162828025006093/aon-20241231.htm · Line register: Analysis/ehb-broker-metrics/m1/assumption_register.csv</p>

<h3>Unknown / kill implication</h3>
<p>Unknown: broker vs direct split; those lives by worksite or home state; which broker’s book is in scope; employee vs dependent. Public sources (KFF EHBS, Marsh / Aon / WTW / Gallagher 10-Ks, Form 5500) do not fill F1. Kill implication: if a named book has no employers Charter can put a residential Internet account in front of, the channel cannot win Spectrum residential from EH&amp;B. A national book will almost never fail the weak “any worksite in a present state” screen — that is why the specified number staying blank is not a hidden yes.</p>
</div>

<!-- ===================== M2 ===================== -->
<div class="page">
<p class="kicker">M2 · Broker concentration</p>
<h2>Top 10 EH&amp;B brokers hold 44% of Form 5500 compensation — not covered lives, and not few enough doors to cover the book</h2>
<div class="stats">
  <div class="stat warn"><div class="v">Valid</div><div class="l">Verdict — wrong unit</div></div>
  <div class="stat info"><div class="v">44%</div><div class="l">Top 10 share of EH&amp;B Form 5500 compensation</div></div>
  <div class="stat warn"><div class="v">Not retrieved</div><div class="l">Specified lives share (top N ÷ broker-placed lives)</div></div>
</div>
<p>Definition: covered lives at the top N broker firms (register example: top 5 / top 10) ÷ total broker-placed covered lives in the same geography as M1. Unit: percent. N is not locked. Public rankings retrieve revenue (Business Insurance) and compensation (Form 5500). Lives share is blank on purpose.</p>

<div class="callout warn"><div class="callout-title">Do not treat 44% or 67% as the specified metric</div>
Business Insurance ranks U.S. brokerage revenue. PlanOptica ranks Form 5500 health-and-welfare compensation. Neither is lives. The two rankings do not share a unit or a denominator — their agreement would not validate M2.
</div>

<h3>Top 10 EH&amp;B filing names hold 44% of Form 5500 compensation; 56% is a long tail</h3>
<div class="legend"><span><span class="swatch gray"></span>Top 10 filing names (compensation share)</span><span><span class="swatch blue"></span>All other 43,119 names</span></div>
{m2_bars}
<p class="src">Source: PlanOptica health-and-welfare broker table, ranked by Form 5500 compensation (not lives). Public page lists {PO_COUNT} names. MMA 7.1% + Gallagher Benefit 5.9 + Lockton 5.8 + Mercer Health 5.0 + USI 4.3 + WTW US 4.3 + Hub Midwest 4.1 + Aon Consulting 2.8 + Alliant 2.5 + Brown &amp; Brown 2.2 = 44.0%. Remainder 56.0%. URL: https://planoptica.com/welfare/broker · Underlying filings: U.S. DOL / EBSA Form 5500, paid for by plan sponsors. Indexer: PlanOptica (for-profit). Entity names are filing names — MMA and Mercer are separate; Hub is Midwest only. Arithmetic: Analysis/ehb-broker-metrics/m2/out/summary.json</p>

<table>
<thead><tr><th>ID</th><th>Figure</th><th class="num">Value</th><th>Year</th><th>Origin</th><th>Unit caveat</th></tr></thead>
<tbody>
<tr class="accent"><td>B2</td><td>PlanOptica top 10 compensation share</td><td class="num">44.0%</td><td>—</td><td>DERIVED · PlanOptica / Form 5500</td><td>Compensation, not lives</td></tr>
<tr class="accent"><td>A5</td><td>BI top 10 / Top 100 U.S. brokerage revenue</td><td class="num">67%</td><td>2024</td><td>DERIVED · Business Insurance</td><td>Revenue, P&amp;C + benefits; of Top 100, not the market. Top 5 = 46.7%; Top 100 sum = ${BI_TOP100_B}B</td></tr>
<tr><td>A7</td><td>MarshBerry citing BI next vintage</td><td class="num">70%</td><td>2025</td><td>RETRIEVED · MarshBerry / BI</td><td>Different year — do not mix with 67%</td></tr>
<tr><td>D1/D2</td><td>KFF broker fees PMPM, small vs large group</td><td class="num">$29.79 / $10.00</td><td>2024</td><td>RETRIEVED · KFF / Mark Farrah</td><td>Why compensation understates jumbo lives</td></tr>
<tr class="blank"><td>X1</td><td>Covered-lives share at top N</td><td class="num">blank</td><td>—</td><td>NOT RETRIEVED</td><td>This is specified M2. Leave blank.</td></tr>
</tbody></table>
<p class="src">BI PDF: https://communityprod.blob.core.windows.net/public/top_100_2025_bi.pdf · Directory: https://www.businessinsurance.com/biresources/2025-agents-and-brokers-directory/ · Funded by Business Insurance / Crain Communications; brokers self-report. Firms with more than 49% of gross revenue from personal lines are excluded. KFF PMPM: https://www.kff.org/health-costs/state-indicator/health-insurance-broker-compensation/ · MarshBerry 2025 vintage: https://www.marshberry.com/resource/the-next-era-of-insurance-brokerage-the-top-100-defined-by-growth-consolidation-and-ai/ (broker M&amp;A advisory thought leadership). Named 10-Ks (MMC, Aon, WTW, Gallagher) disclose revenue lines, not covered lives.</p>

<h3>Unknown / kill implication</h3>
<p>Unknown: specified lives share; brokers below the BI Top 100; Spectrum-state cut; welfare plans and jumbo consulting fees off Form 5500 Schedule A; sister-entity undercount (Hub Midwest, WTW US LLC). Channel-architecture implication: the closest EH&amp;B proxy says a handful of national partnerships does not cover the book (56% of observed compensation sits in a 43,119-name tail). That changes how many doors to open. It is not a product go-live. Lives share — the specified number — remains unknown, so this gate is directional, not closed.</p>
</div>

<!-- ===================== M5 ===================== -->
<div class="page">
<p class="kicker">M5 · Footprint</p>
<h2>M5 is the right hard gate; we cannot claim a serviceable-home share — 42-state presence is not employee-home serviceability</h2>
<div class="stats">
  <div class="stat info"><div class="v">Valid</div><div class="l">Verdict on the metric</div></div>
  <div class="stat warn"><div class="v">None</div><div class="l">Verdict on a number</div></div>
  <div class="stat warn"><div class="v">94.6%</div><div class="l">ESTIMATED ceiling — SUSB jobs in Charter-present states, not M5</div></div>
</div>
<p>Definition: enrolled (or broker-covered) employees whose residential address passes a Spectrum residential serviceability check ÷ enrolled / covered employees in the pilot book who supply a home address. Unit: percent. Hard fail for out-of-footprint homes. Serviceability is installable residential Internet at that address — not “lives in a Charter state,” not FCC advertised availability alone, and not already a subscriber (that cut is a later metric).</p>

<div class="callout warn"><div class="callout-title">No employee-level serviceable percent exists in this repo</div>
Do not read 94.6% as the share of enrolled employees Charter can serve at home. Do not read a state’s Charter location share as that rate either. M5 needs a consented home address on a named book, geocoded and matched to Spectrum residential serviceability. That file is not here.
</div>

<p>The 94.6% ceiling is 128,435,628 ÷ 135,748,407: share of 2022 SUSB payroll employment in the 42 jurisdictions with any Charter residential filing as of 31 Dec 2025. The remaining 7,312,779 jobs (5.4%) are in Alaska, Arkansas, Delaware, the District of Columbia, Iowa, North Dakota, Oklahoma, South Dakota, and Utah. Four biases, all high: workplace not home; any presence not plant (Rhode Island and Maryland count as “in”); book ≠ nation; 2022 jobs nailed to 2025-12-31 plant.</p>
<p class="src">SUSB 2022: https://www2.census.gov/programs-surveys/susb/tables/2022/us_state_naics_detailedsizes_2022.txt · Landing: https://www.census.gov/data/datasets/2022/econ/susb/2022-susb.html · Funded by U.S. Census Bureau. Ceiling C4 recomputed in Analysis/ehb-broker-metrics/m5/build_m5_proxy.py · must match Analysis/ehb-broker-metrics/m5/out/summary.json</p>

<h3>Charter’s residential location share inside “present” states runs from 92% in Hawaii to 0.02% in Rhode Island</h3>
<div class="legend"><span><span class="swatch blue"></span>Range ends (Hawaii / Rhode Island)</span><span><span class="swatch gray"></span>Other selected present states</span></div>
{m5_bars}
<p class="src">Selected states, sorted descending. Unit: percent of residential fabric units (res_st_pct × 100). Florida 36.2 is 0.3615 to one decimal; Rhode Island 0.02 is 0.0002; Maryland 0.2 is 0.0016. Across all 42 present jurisdictions, Charter’s estimated units are {FABRIC_SHARE}% of that group’s residential fabric (50,939,445 ÷ 151,719,024) — a plant statistic, still not M5. Unweighted median among the 42 is 20.3%. Source: FCC Broadband Data Collection / National Broadband Map, vintage 31 Dec 2025 (D25), revision 15 Sep 2026. Holding company Charter Communications, provider ID 130235. https://broadbandmap.fcc.gov/data-download · Arithmetic: Analysis/cable-share-small-firm/</p>

<table>
<thead><tr><th>ID</th><th>Figure</th><th class="num">Value</th><th>Origin</th><th>Why it is not M5</th></tr></thead>
<tbody>
<tr class="blank"><td>A2</td><td>Serviceable-home share of enrolled employees</td><td class="num">none</td><td>BLOCKING</td><td>No roster × address match. This is M5.</td></tr>
<tr><td>B1</td><td>Charter-present jurisdictions</td><td class="num">42</td><td>RETRIEVED · FCC D25</td><td>Any filing, including Rhode Island at 106 estimated residential units</td></tr>
<tr class="warn"><td>C4</td><td>SUSB jobs in those 42 jurisdictions</td><td class="num">94.6%</td><td>ESTIMATED ceiling</td><td>Workplace employment, not homes; book ≠ nation</td></tr>
<tr><td>D3</td><td>Charter units ÷ fabric in the 42</td><td class="num">33.6%</td><td>DERIVED · plant</td><td>BSLs in present states, not enrolled employees</td></tr>
<tr><td>X1</td><td>Employment-weighted state location share</td><td class="num">refused</td><td>REFUSED</td><td>Two wrong substitutions. Not quoted as a rate.</td></tr>
</tbody></table>

<h3>Unknown / kill implication</h3>
<p>Unknown: the single join that is M5 — consented employee home address × Spectrum residential serviceability on a named book. Write the denominator before the pilot (eligible-with-address for a go/no-go on the book; enrolled-with-address for fulfillment). Treat indeterminate (PO box, campus, stale HR file, MDU rights pending) as a third bucket, not a silent drop. Kill: if a consented address match on the candidate book shows Spectrum cannot fulfill the benefit as written for the homes the contract covers, and the design has no multi-ISP fallback, do not scale a national “included Spectrum Internet” benefit. The public 94.6% ceiling cannot clear or kill this. Only the roster match can. Addresses are for serviceability only — not usage, health, or productivity monitoring.</p>
</div>

<!-- ===================== M9 ===================== -->
<div class="page">
<p class="kicker">M9 · Take-up</p>
<h2>M9 is a valid gate glued from two metrics — split who pays from who enrolls; there is no Charter take-up number to cite</h2>
<div class="stats">
  <div class="stat warn"><div class="v">Needs split</div><div class="l">Verdict — valid gate, two metrics</div></div>
  <div class="stat warn"><div class="v">None</div><div class="l">Charter take-up rate in this build</div></div>
  <div class="stat info"><div class="v">5%</div><div class="l">Closest analogue: 2017 SHRM employer offer of subsidized home internet — not take-up</div></div>
</div>
<p>The register already named two rates under one ID. Keep both tests; stop reporting them as one number.</p>
<table>
<thead><tr><th>Split ID</th><th>Unit</th><th>Numerator / denominator</th><th>What a yes/no means</th></tr></thead>
<tbody>
<tr class="accent"><td>M9a · Payor mix</td><td>Employer (and lives)</td><td>Employers choosing employer-paid vs employee-paid voluntary vs hybrid ÷ pilot employers offered a design</td><td>Someone other than the employee is willing to fund the line — or not</td></tr>
<tr class="warn"><td>M9b · Enrollment take-up</td><td>Eligible employee</td><td>Enrolled internet accounts ÷ eligible employees offered the benefit, inside each M9a cell</td><td>Eligible people actually elect the line once it is on the ballot</td></tr>
</tbody></table>
<p class="muted">A blended enrolled/eligible headline across payor models is forbidden: 100% of a tiny employer-paid cell plus 4% of a large voluntary cell can look like a healthy mid-teens rate. Time-box M9b at close of open enrollment plus 90 days. Transfers of existing Spectrum accounts are enrollment, not net adds (M13 owns that).</p>

<div class="callout warn"><div class="callout-title">No Charter internet-as-benefit SKU in this repo</div>
Filings document Spectrum Internet, Mobile, TV, Voice, Advanced/Invincible WiFi, and rural. Brokerage Metrics.xlsx and the core-metrics canvas leave M9 blank on purpose. Do not fill the blank with an analogue. KFF 76% medical take-up among eligible workers is heavily subsidized core medical — the wrong analogue for home internet as a benefit.
</div>

<h3>Lifestyle-adjacent benefits are offered by a minority of employers — and those are offer rates, not take-up</h3>
<div class="legend"><span><span class="swatch blue"></span>Closest internet-as-benefit offer analogue (2017)</span><span><span class="swatch gray"></span>Other SHRM employer offer rates (analogues)</span></div>
{m9_bars}
<p class="src">Pet 22%, transit 12%, parking 10%, QSLP match 4%: SHRM 2025 Employee Benefits Survey (plan year 2025; n=3,969 U.S. HR professionals; unweighted; fielded 21 Jan–10 Mar 2025; funded by SHRM). Home internet 5%: SHRM 2017 customized benefits-prevalence (n≈2,717) — free, discounted, or subsidized home Internet service; offer, not enrollment; dated. SHRM 2025/2026 executive summaries publish at-home equipment subsidy (55% / 53%), not an internet-specific stipend. Do not treat equipment subsidy as internet. Chart scale 0–30 so the 5% bar is readable; bars still start at zero. Full analogue register: Analysis/ehb-broker-metrics/m9/analogue_register.csv</p>

<table>
<thead><tr><th>Role</th><th>Analogue</th><th class="num">Rate</th><th>Year</th><th>Funder</th><th>Why it is not Charter take-up</th></tr></thead>
<tbody>
<tr class="warn"><td>Take-up (wrong analogue)</td><td>KFF medical take-up among eligible workers</td><td class="num">76%</td><td>2025</td><td>KFF (nonprofit; not Kaiser Permanente). Davis Research. n=1,862 employers 10+</td><td>Heavily employer-subsidized core medical. Workers still pay 16% single / 26% family. Eligibility 80%; enrollment among workers at offering firms 61%.</td></tr>
<tr class="accent"><td>Offer (closest)</td><td>SHRM subsidized home Internet</td><td class="num">5%</td><td>2017</td><td>SHRM</td><td>Offer, not take-up. Dated. No Spectrum SKU.</td></tr>
<tr><td>Offer</td><td>SHRM pet insurance</td><td class="num">22%</td><td>2025</td><td>SHRM</td><td>Offer, not enrollment. Pet ≠ internet.</td></tr>
<tr><td>Offer — not internet</td><td>SHRM at-home equipment subsidy</td><td class="num">55% / 53%</td><td>2025 / 2026</td><td>SHRM</td><td>Equipment, not an internet stipend. LIMRA/EY interest rates are also not take-up (full analogue sheet in Excel).</td></tr>
<tr class="blank"><td>Excluded</td><td>Stealth Agents “SHRM internet stipend 38%”</td><td class="num">not used</td><td>2025</td><td>Secondary aggregator</td><td>Not retrieved from SHRM primary.</td></tr>
<tr class="blank"><td>Charter M9</td><td>Internet-as-benefit take-up</td><td class="num">blank</td><td>—</td><td>NOT RETRIEVED</td><td>Pilot is the source.</td></tr>
</tbody></table>
<p class="src">KFF 2025 EHBS: https://www.kff.org/health-costs/2025-employer-health-benefits-survey/ · SHRM 2025 exec summary: https://www.shrm.org/content/dam/en/shrm/topics-tools/research/employee-benefits/2025_annual_benefits_survey_executive_summary.pdf · SHRM 2017 customized: https://www.shrm.org/content/dam/en/shrm/topics-tools/news/organizational-employee-development/Benefits-Prevalence-Report-All-Industries-All-FTEs.pdf · SHRM 2026 exec summary: https://www.shrm.org/content/dam/en/shrm/topics-tools/research/employee-benefits/2026_employee-benefits_executive-summary.pdf</p>

<h3>Unknown / kill implication</h3>
<p>Unknown: any Charter internet-as-benefit take-up; whether a SKU exists; 2026 prevalence of an internet stipend; payor mix on a real ballot. Write kill floors before the first employee sees an offer: if no pilot employer will fund employer-paid and none will put a voluntary SKU on the ballot — stop. If a design is on the ballot and voluntary take-up among eligibles in that cell, after OE plus 90 days, lands in the single digits with no path to change price or payor — do not scale on the hope that “benefits take-up is usually high.” Analogues bound the plausible range. Analogues are not the floor.</p>
</div>

<!-- ===================== M13 ===================== -->
<div class="page">
<p class="kicker">M13 · Charter economics</p>
<h2>M13 is the channel’s Internet-print metric — but only as an account count, not as enrollments</h2>
<div class="stats">
  <div class="stat warn"><div class="v">Valid</div><div class="l">Verdict — tighten to P−E−D</div></div>
  <div class="stat warn"><div class="v">Not in repo</div><div class="l">Channel M13 (accounts)</div></div>
  <div class="stat danger"><div class="v">−117k</div><div class="l">Company res. Internet net adds, Q1 2026 — context only</div></div>
</div>
<p>The earlier benefits analogy transfers only when the employer / broker path opens a household Spectrum Internet customer relationship. Enrollments, passings, and “employees reached” can all rise while the Internet print stays unchanged. The register definition mixed an account count with a percent of enrollments and omitted same-window job-exit disconnects. Tighten to one unit — residential Internet accounts — aligned to Charter’s Ex99.1 operating-statistic quarter:</p>

<div class="card"><div class="card-h">M13<sub>t</sub> = P<sub>t</sub> − E<sub>t</sub> − D<sub>t</sub></div>
<p><strong>P · path-tagged provisions.</strong> Residential Spectrum Internet customer relationships (same PSU definition as Ex99.1 residential Internet) whose order in t is tagged to the broker / employer path. Count the billed household Internet PSU. Not a benefits enrollment row, not a mobile line, not a Spectrum Business circuit.</p>
<p><strong>E · existing / payer-only.</strong> Subset of P where that household already had an active Spectrum residential Internet PSU at t−1. Employer becomes the payer, or a discount is applied, and no new customer relationship opens. Retention or mix, not net adds.</p>
<p><strong>D · same-window job-exit disconnects.</strong> Subset of (P − E) that disconnect from Spectrum residential Internet in t because employment eligibility ended. Do not subtract households that stay Spectrum as self-pay. Job-exit after t is M14, not this period’s M13.</p>
</div>
<p class="muted">P, E, D, and M13 are blank. Do not invent a conversion of enrollments into accounts. Optional diagnostic, not M13: (P − E) / benefit enrollments — leave blank until a pilot produces both counts.</p>

<h3>Company already reports this unit — not this channel</h3>
<table>
<thead><tr><th>Period</th><th class="num">Residential Internet net adds</th><th>What it is</th><th>File</th></tr></thead>
<tbody>
<tr><td>FY2025</td><td class="num">−393,000</td><td>10-K: residential Internet customers decreased vs 2024</td><td>PrimarySources/readable/CHTR-10-K-FY2025.md</td></tr>
<tr class="warn"><td>Q1 2026</td><td class="num">−117,000</td><td>Ex99.1 operating statistics, residential Internet quarterly net additions (117)</td><td>PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md</td></tr>
<tr><td>Q2 2026</td><td class="num">−166,000</td><td>Ex99.1 operating statistics, residential Internet quarterly net additions (166)</td><td>PrimarySources/readable/CHTR-Q2-2026-Earnings-Release-Ex99.1.md</td></tr>
<tr class="blank"><td>Channel M13</td><td class="num">blank</td><td>P − E − D for the broker / employer path</td><td>Not in repo. Do not scale company print by an invented take-up.</td></tr>
</tbody></table>
<p class="src">Q1 total Internet was −120k (residential −117k + small business −3k). M13 is residential only — do not use the −120k headline as the analog. These figures prove Charter tracks residential Internet net adds as accounts. They are not M13.</p>

<table>
<thead><tr><th>ID</th><th>Field required to compute M13</th><th>System</th><th>Status</th></tr></thead>
<tbody>
<tr><td>D1</td><td>Channel tag: employer ID + broker / path code on the residential Internet PSU</td><td>Order entry / billing</td><td>MISSING</td></tr>
<tr><td>D2</td><td>Prior-account flag: active Spectrum residential Internet at t−1</td><td>Billing snapshot</td><td>MISSING</td></tr>
<tr><td>D3</td><td>Payer type: employer-billed vs employee-billed vs self-pay</td><td>Billing</td><td>MISSING</td></tr>
<tr><td>D4</td><td>Eligibility end date</td><td>Benefits-admin / employer roster</td><td>MISSING</td></tr>
<tr><td>D5</td><td>Disconnect date and whether the PSU actually ended</td><td>Billing</td><td>MISSING</td></tr>
<tr><td>D6</td><td>Window = Ex99.1 calendar quarter, locked before the pull</td><td>Reporting calendar</td><td>MISSING</td></tr>
</tbody></table>

<h3>Unknown / kill implication</h3>
<p>Unknown: any channel net-add count; any conversion rate; any share of the −117k / −393k company print this path would close; contribution, subsidy, or broker cost (M14; inputs blank). Tagged is not incremental — a path tag does not prove the household would not have signed up anyway. Kill the channel as an Internet-print play if a time-boxed pilot with D1–D5 filled shows M13 ≈ 0 because almost all of P is E or D — even if enrollments are large. Do not substitute enrollments, passings, or reach.</p>

<hr/>
<p class="src">Parent register: Decisions/EH&amp;B/Brokerage Metrics.xlsx (Core metrics) and chtr-ehb-broker-channel-core-metrics.canvas.tsx. Validations: chtr-ehb-m1-broker-covered-lives, chtr-ehb-m2-large-broker-concentration, chtr-ehb-m5-serviceable-home-share, chtr-ehb-m9-payor-takeup, chtr-ehb-m13-net-new-residential. Arithmetic: Analysis/ehb-broker-metrics/m1, m2, m5, m9, m13. This digest dated {DATE}. Unwilling to claim: a filled M1, M2 lives share, M5 percent, Charter take-up, or M13 channel count; a TAM; a launch recommendation.</p>
</div>

</body></html>"""

HTML_PATH.write_text(html, encoding="utf-8")
DEST_DIR.mkdir(parents=True, exist_ok=True)

if PDF_PATH.exists():
    try:
        PDF_PATH.unlink()
    except OSError as e:
        print("WARN unlink PDF", e)

cmd = [
    CHROME,
    "--headless=new",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={PDF_PATH}",
    HTML_PATH.as_uri(),
]
r = subprocess.run(cmd, capture_output=True, text=True)
print("chrome_rc", r.returncode)
if r.stderr:
    print("chrome_err", r.stderr[:800])
time.sleep(1.5)

assert PDF_PATH.exists(), "PDF missing"
pdf_bytes = PDF_PATH.read_bytes()
print("PDF_BYTES", len(pdf_bytes))
print("PDF_HEADER", pdf_bytes[:8])
assert pdf_bytes[:5] == b"%PDF-", "PDF is not a real PDF"

# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

header_font = Font(bold=True)
header_fill = PatternFill("solid", fgColor="F2F2F2")
wrap = Alignment(wrap_text=True, vertical="top")
warn_fill = PatternFill("solid", fgColor="FFF4D6")
blank_fill = PatternFill("solid", fgColor="F3F3F3")
info_fill = PatternFill("solid", fgColor="EEF4FB")


def style_header(ws, ncols, freeze="A2"):
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = freeze
    last_row = max(ws.max_row, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last_row}"


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def wrap_all(ws, min_row=2):
    for row in ws.iter_rows(min_row=min_row, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.alignment = wrap


wb = Workbook()

# ----- Notes -----
ws = wb.active
ws.title = "Notes"
ws.append(["Field", "Content"])
notes = [
    ("Title", "Brokerage Metrics Breakdown — validation of five Employee Health & Benefits broker-channel metrics (M1, M2, M5, M9, M13)"),
    ("Date", DATE),
    ("What this file is", "A synthesis of five completed metric validations. It records the verdict, the retrieved public figures (with sources), what remains unknown, and any kill implication the validation stated. It is not a TAM, not a launch recommendation, and not a rewrite of the original register."),
    ("What this file is not", "Not Decisions/EH&B/Brokerage Metrics.xlsx (that file is the original 16-metric register and is left unchanged). Not a filled M1, M2 lives share, M5 percent, Charter take-up, or M13 channel count."),
    ("Blanks are intentional", "Charter take-up, serviceable-home percent of any book, broker-placed lives by state, and M13 channel account counts are blank / not retrieved on purpose. Do not invent them. Do not paste an analogue, a state employment share, or company Internet net adds into those cells."),
    ("Naming collision", "Repo folder Decisions/EH&B means existing homes and businesses on plant already built. This file means Employee Health & Benefits (insurance)."),
    ("Offering under test", "Insurance brokerage firms include employee home connectivity in the benefits offering they already sell to employers. The motion only counts when a residential Spectrum Internet account opens at the employee’s home."),
    ("Headline verdicts", "M1 valid — numerator blank. M2 valid — wrong unit (compensation / revenue, not lives). M5 valid — no number. M9 needs split M9a/M9b. M13 valid — tighten to P−E−D."),
    ("Do not substitute", "M1: 168.2M/178.2M ESI overlay is not M1; SUSB 45.9% is not M1. M2: PlanOptica 44% and BI 67% are not lives. M5: 94.6% jobs in 42 Charter-present states is an ESTIMATED ceiling, not M5. M9: KFF 76% medical take-up is the wrong analogue; SHRM 2017 5% is an employer offer, not take-up. M13: Q1 2026 residential Internet −117k is company context, not this channel."),
    ("Sheet guide", "Notes = this framing. Scorecard = five-row verdict table. M1 / M2 / M5 / M9 / M13 = one sheet per metric (id, name, verdict, definition, retrieved figures with year and source URL, unknowns, kill rule). Analogues = M9 public analogue rates only, labeled as analogues, not Charter take-up."),
    ("Parent register (unchanged)", "Decisions/EH&B/Brokerage Metrics.xlsx · Core metrics; canvases/chtr-ehb-broker-channel-core-metrics.canvas.tsx"),
    ("Validation canvases", "chtr-ehb-m1-broker-covered-lives.canvas.tsx; chtr-ehb-m2-large-broker-concentration.canvas.tsx; chtr-ehb-m5-serviceable-home-share.canvas.tsx; chtr-ehb-m9-payor-takeup.canvas.tsx; chtr-ehb-m13-net-new-residential.canvas.tsx"),
    ("Arithmetic (recomputed in code)", "Analysis/ehb-broker-metrics/m1/; m2/out/; m5/out/; m9/; m13/"),
    ("Companion PDF", "Decisions/EH&B/Brokerage Metrics Breakdown.pdf"),
]
for row in notes:
    ws.append(list(row))
style_header(ws, 2)
set_widths(ws, [28, 110])
wrap_all(ws)
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 48

# ----- Scorecard -----
ws = wb.create_sheet("Scorecard")
score_headers = [
    "Metric ID",
    "Metric name",
    "Verdict",
    "What you can use",
    "Specified number",
    "Closest retrieved (labeled — not the metric)",
    "Source URL / path",
]
ws.append(score_headers)
score_rows = [
    (
        "M1",
        "Broker-placed covered lives in Spectrum states",
        "Valid — numerator blank",
        "National ESI scale (KFF / MEPS / EBSA) and Charter 42-state presence. Use as a reach-screen specification, not a filled percent.",
        "blank / not retrieved",
        "168.2M of 178.2M ESI persons reside in Charter-present states (94.4%) — FLAGGED overlay, not M1. SUSB 45.9% is not M1.",
        "https://www.dol.gov/sites/dolgov/files/EBSA/researchers/data/health-and-welfare/health-insurance-coverage-bulletin-2024.pdf",
    ),
    (
        "M2",
        "Large-broker concentration of that book",
        "Valid — wrong unit",
        "PlanOptica 44% top-10 of Form 5500 compensation; Business Insurance 67% top-10 of Top 100 revenue. Directional on how many doors to open. Not a lives share.",
        "blank / not retrieved (lives)",
        "44.0% PlanOptica top 10 compensation; 67.2% BI top 10 / Top 100 2024 U.S. brokerage revenue. Neither is lives.",
        "https://planoptica.com/welfare/broker",
    ),
    (
        "M5",
        "Serviceable-home share of enrolled employees",
        "Valid — no number",
        "The hard footprint gate. Specification and kill rule only. Public plant/employment figures are the wrong grain.",
        "none — no roster × address match",
        "94.6% of SUSB 2022 jobs sit in 42 Charter-present jurisdictions — ESTIMATED ceiling, not M5. Plant share Hawaii 92.3% to Rhode Island 0.02%.",
        "https://broadbandmap.fcc.gov/data-download",
    ),
    (
        "M9",
        "Payor model and enrollment take-up",
        "Needs split M9a / M9b",
        "Keep both tests. Analogues bound a plausible range. Do not headline a blended enrolled/eligible rate. Pilot is the source.",
        "blank / not retrieved (Charter take-up)",
        "SHRM 2017 5% of employers offered subsidized home Internet (offer, not take-up). KFF 2025 76% medical take-up is the wrong analogue.",
        "https://www.shrm.org/content/dam/en/shrm/topics-tools/news/organizational-employee-development/Benefits-Prevalence-Report-All-Industries-All-FTEs.pdf",
    ),
    (
        "M13",
        "Net new residential Internet accounts from the channel",
        "Valid — tighten to P−E−D",
        "Account count: path-tagged provisions minus existing/payer-only minus same-window job-exit disconnects. Requires billing tags. Not enrollments.",
        "blank / not retrieved (channel accounts)",
        "Company residential Internet net adds Q1 2026 = −117,000 (Ex99.1). Context that Charter already reports this unit. Not this channel.",
        "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md",
    ),
]
for row in score_rows:
    ws.append(list(row))
    for c in range(1, 8):
        ws.cell(ws.max_row, c).fill = warn_fill
style_header(ws, 7)
set_widths(ws, [12, 42, 26, 55, 32, 62, 55])
wrap_all(ws)
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 72


def add_metric_sheet(title, identity_rows, figure_rows, unknown_rows, kill_text):
    """One sheet: identity block, then retrieved figures with source URL, then unknowns + kill."""
    sh = wb.create_sheet(title)
    headers = [
        "Block",
        "ID",
        "Item",
        "Value",
        "Year",
        "Origin",
        "Source URL",
        "Notes",
    ]
    sh.append(headers)
    for row in identity_rows:
        sh.append(list(row))
    for row in figure_rows:
        sh.append(list(row))
        if str(row[3]).lower() in ("blank", "none", "not retrieved", "missing", "refused"):
            for c in range(1, 9):
                sh.cell(sh.max_row, c).fill = blank_fill
    for row in unknown_rows:
        sh.append(list(row))
    sh.append(["Kill rule", "", kill_text, "", "", "", "", "Written by the validation, not invented here."])
    style_header(sh, 8)
    set_widths(sh, [16, 10, 42, 22, 12, 28, 62, 55])
    wrap_all(sh)
    for r in range(2, sh.max_row + 1):
        sh.row_dimensions[r].height = 36
    return sh


# ----- M1 -----
add_metric_sheet(
    "M1",
    [
        ("Identity", "M1", "Broker-placed covered lives in Spectrum states", "", "", "DEFINITION", "chtr-ehb-broker-channel-core-metrics.canvas.tsx", "Parent register: Decisions/EH&B/Brokerage Metrics.xlsx Core metrics M1"),
        ("Verdict", "V1", "Verdict on M1 as specified", "valid-with-caveats — numerator blank", DATE, "ESTIMATED", "chtr-ehb-m1-broker-covered-lives.canvas.tsx", "Must match Analysis/ehb-broker-metrics/m1/summary.json"),
        ("Definition", "S1", "Numerator", "Employees whose employer has a worksite in a Charter-present state (FCC BDC D25, provider_id 130235)", "", "DEFINITION", "https://broadbandmap.fcc.gov/data-download", "Also report the stricter employee-home-state cut when an address file exists. Presence is not home serviceability (that is M5)."),
        ("Definition", "S2", "Denominator", "All employees on that named broker’s US medical book", "", "DEFINITION", "Decisions/EH&B/Brokerage Metrics.xlsx", "US ESI is not the denominator. Industry-level M1 needs a named book."),
        ("Definition", "S3", "Unit", "Covered employees enrolled in broker-placed medical; dependents reported separately as lives", "", "DEFINITION", "chtr-ehb-m1-broker-covered-lives.canvas.tsx", "Do not mix employees and dependents."),
        ("Definition", "A3", "Charter-present states", "42", "2025-12-31", "DERIVED from RETRIEVED", "https://broadbandmap.fcc.gov/data-download", "Any reported residential availability. Absent: AK, IA, ND, OK, SD, AR, DE, UT + DC."),
        ("Definition", "A4", "Charter-lead states", "19", "2025-12-31", "DERIVED from RETRIEVED", "https://broadbandmap.fcc.gov/data-download", "Strictly larger res_st_pct than Comcast. Lead is availability share, not serviceable homes."),
    ],
    [
        ("Retrieved — not M1", "C5", "Private-sector workers enrolled at own employer", "65.6 million workers", "2024", "RETRIEVED", "https://meps.ahrq.gov/data_files/publications/rf54/rf54.pdf", "AHRQ MEPS-IC Research Findings 54. Own-employer enrollment, not broker-placed, not by state."),
        ("Retrieved — not M1", "C4", "Workers covered through a current employer", "82.8 million workers", "CY2023", "RETRIEVED", "https://www.dol.gov/sites/dolgov/files/EBSA/researchers/data/health-and-welfare/health-insurance-coverage-bulletin-2024.pdf", "EBSA CPS bulletin. Employees, still all ESI."),
        ("Retrieved — not M1", "C3", "ESI from own job, under 65", "85.7 million people", "Mar 2025", "RETRIEVED", "https://www.kff.org/health-costs/health-policy-101-employer-sponsored-health-insurance/", "KFF / CPS ASEC. People, not a broker book."),
        ("Retrieved — not M1", "B1", "Persons with any ESI", "178.2 million persons", "CY2023", "RETRIEVED", "https://www.dol.gov/sites/dolgov/files/EBSA/researchers/data/health-and-welfare/health-insurance-coverage-bulletin-2024.pdf", "EBSA Table 1A. Includes dependents, retiree, COBRA."),
        ("Retrieved — not M1", "B2", "Sum of Table 1A state ESI cells", "178.4 million persons", "CY2023", "DERIVED", "https://www.dol.gov/sites/dolgov/files/EBSA/researchers/data/health-and-welfare/health-insurance-coverage-bulletin-2024.pdf", "Bulletin rounding vs printed USA 178.2."),
        ("Retrieved — not M1", "B3", "ESI persons residing in Charter-present states", "168.2 million persons", "CY2023 × D25", "DERIVED — FLAGGED", "https://www.dol.gov/sites/dolgov/files/EBSA/researchers/data/health-and-welfare/health-insurance-coverage-bulletin-2024.pdf", "NOT M1. Residence, not worksite. All ESI, not broker-placed."),
        ("Retrieved — not M1", "B4", "ESI persons residing in Charter-absent jurisdictions", "10.2 million persons", "CY2023 × D25", "DERIVED — FLAGGED", "https://www.dol.gov/sites/dolgov/files/EBSA/researchers/data/health-and-welfare/health-insurance-coverage-bulletin-2024.pdf", "NOT M1. Same caveats as B3."),
        ("Retrieved — not M1", "B5", "ESI persons residing in Charter-lead states", "80.7 million persons", "CY2023 × D25", "DERIVED — FLAGGED", "https://www.dol.gov/sites/dolgov/files/EBSA/researchers/data/health-and-welfare/health-insurance-coverage-bulletin-2024.pdf", "NOT M1. 19 lead states."),
        ("Retrieved — not M1", "B6", "B3 / B1 present share of US ESI persons", "94.4 percent", "CY2023 × D25", "DERIVED — FLAGGED", "Analysis/ehb-broker-metrics/m1/summary.json", "NOT M1. Uses printed USA 178.2 as denominator."),
        ("Retrieved — not M1", "C1", "Nonelderly with ESI (ACS hierarchy)", "154 million people under 65", "2025", "RETRIEVED", "https://www.kff.org/health-costs/2025-employer-health-benefits-survey/", "KFF EHBS (ACS). No broker vs direct split in 2024 or 2025 EHBS. KFF-funded; fielded by Davis Research."),
        ("Retrieved — not M1", "C6", "AHIP employer-provided coverage", ">180 million Americans", "2023/24", "RETRIEVED — AHIP-funded", "https://www.ahip.org/news/press-releases/new-ahip-report-highlights-indispensability-of-employer-provided-coverage-state-by-state", "Rounded advocacy total including families."),
        ("Retrieved — not M1", "D1", "Form 5500-filing group health participants", "~88 million participants", "2023", "RETRIEVED", "https://www.dol.gov/agencies/ebsa/researchers/statistics/retirement-bulletins", "Filing subset; most small fully insured plans exempt; not broker of record."),
        ("Retrieved — not M1", "E1", "Aon Health Solutions revenue", "$3.335 billion", "FY2024", "RETRIEVED", "https://www.sec.gov/Archives/edgar/data/315293/000162828025006093/aon-20241231.htm", "Global. Not US covered lives. Not state-split."),
        ("Retrieved — not M1", "E2–E4", "Marsh / WTW / Gallagher covered lives in 10-K", "blank", "FY2024", "NOT DISCLOSED", "https://www.sec.gov/Archives/edgar/data/62709/000006270925000015/mmc-20241231.htm", "Health businesses described; lives not filed."),
        ("Specified M1", "F1", "Broker-placed lives in Charter-present states", "blank", "", "NOT RETRIEVED", "", "This is M1. Leave blank. Do not allocate national lives by population."),
        ("Specified M1", "F2", "Share of a named broker book in Charter-present states", "blank", "", "NOT RETRIEVED", "", "Denominator is that broker book, not US ESI."),
        ("Excluded", "X1", "SUSB small-firm employment share", "45.9 percent of US payroll employment at firms <500", "2022", "RETRIEVED elsewhere — DO NOT USE AS M1", "Analysis/small-large-employment", "Labor market, not broker-placed covered lives."),
    ],
    [
        ("Unknown", "U1", "Broker vs direct (or consultant) split of medical lives", "not retrieved", "", "NOT RETRIEVED", "", "Every national ESI figure includes carrier-direct and self-administered books."),
        ("Unknown", "U2", "Those lives by state of worksite or home", "not retrieved", "", "NOT RETRIEVED", "", "Cannot overlap the book with Charter-present / lead / absent. Do not allocate national lives by population."),
        ("Unknown", "U3", "Which broker’s book is in scope", "not retrieved", "", "NOT RETRIEVED", "", "M1’s denominator is that book, not the US."),
        ("Unknown", "U4", "Employee vs dependent", "not retrieved", "", "NOT RETRIEVED", "", "AHIP / ACS / Table 1A are people; the offering is an employee home account."),
    ],
    "If a named book has no employers Charter can put a residential Internet account in front of, the channel cannot win Spectrum residential from EH&B. A national book will almost never fail the weak “any worksite in a present state” screen (Charter files in 42 of 50 states). The specified number staying blank is not a hidden yes.",
)

# ----- M2 -----
add_metric_sheet(
    "M2",
    [
        ("Identity", "M2", "Large-broker concentration of that book", "", "", "DEFINITION", "chtr-ehb-broker-channel-core-metrics.canvas.tsx", "Parent register: Decisions/EH&B/Brokerage Metrics.xlsx Core metrics M2"),
        ("Verdict", "V1", "Verdict on M2 as specified", "valid-with-caveats — wrong unit", DATE, "ESTIMATED", "chtr-ehb-m2-large-broker-concentration.canvas.tsx", "Must match Analysis/ehb-broker-metrics/m2/out/summary.json. Specified unit is lives; public rankings are revenue or compensation."),
        ("Definition", "S1", "Numerator", "Covered lives at the top N broker firms (register example: top 5 / top 10)", "", "DEFINITION", "Decisions/EH&B/Brokerage Metrics.xlsx", "N is illustrated, not locked."),
        ("Definition", "S2", "Denominator", "Total broker-placed covered lives in the same geography as M1 (Spectrum-state book)", "", "DEFINITION", "Decisions/EH&B/Brokerage Metrics.xlsx", "Public tables are U.S. No Spectrum-state conversion invented."),
        ("Definition", "S3", "Unit", "Percent of lives", "", "DEFINITION", "Decisions/EH&B/Brokerage Metrics.xlsx", "Public proxies are percent of compensation or of Top 100 revenue."),
        ("Definition", "S4", "Gate", "Whether a few national brokers can move volume or the channel is fragmented", "", "DEFINITION", "Decisions/EH&B/Brokerage Metrics.xlsx", "Changes how many doors to open. Not a product go-live."),
    ],
    [
        ("Retrieved — not lives", "B2", "PlanOptica top 10 share of EH&B Form 5500 compensation", "44.0 percent", "", "DERIVED", "https://planoptica.com/welfare/broker", "Headline EH&B proxy. Compensation, not lives. MMA 7.1 + Gallagher Benefit 5.9 + Lockton 5.8 + Mercer Health 5.0 + USI 4.3 + WTW US 4.3 + Hub Midwest 4.1 + Aon Consulting 2.8 + Alliant 2.5 + Brown & Brown 2.2."),
        ("Retrieved — not lives", "B1", "PlanOptica top 5 compensation share", "28.1 percent", "", "DERIVED", "https://planoptica.com/welfare/broker", "MMA + Gallagher Benefit + Lockton + Mercer Health + USI."),
        ("Retrieved — not lives", "B3", "PlanOptica remainder after top 10", "56.0 percent", "", "DERIVED", "https://planoptica.com/welfare/broker", "100 − B2. Table lists 43,129 broker names. Long tail."),
        ("Retrieved — not lives", "B4", "MMA + Mercer published shares (Marsh parent)", "12.1 percent", "", "DERIVED", "https://planoptica.com/welfare/broker", "Filing names left uncombined on the ranking. Still compensation."),
        ("Retrieved — not lives", "B5", "Named list Marsh+Gallagher+Lockton+WTW+Aon+NFP", "32.5 percent", "", "DERIVED", "https://planoptica.com/welfare/broker", "Uses B4 plus Gallagher 5.9, Lockton 5.8, WTW US 4.3, Aon Consulting 2.8, NFP 1.6."),
        ("Retrieved — not lives", "A5", "BI top 10 / Top 100 U.S. brokerage revenue", "67.2 percent (headline 67%)", "2024", "DERIVED", "https://communityprod.blob.core.windows.net/public/top_100_2025_bi.pdf", "A2/A3 = 53.611B / 79.802B = 67.18%. Revenue, not lives. P&C + benefits. Of Top 100, not the market."),
        ("Retrieved — not lives", "A4", "BI top 5 / Top 100 U.S. brokerage revenue", "46.7 percent", "2024", "DERIVED", "https://communityprod.blob.core.windows.net/public/top_100_2025_bi.pdf", "Marsh, Aon, Gallagher, WTW, Alliant."),
        ("Retrieved — not lives", "A3", "BI Top 100 2024 U.S. brokerage revenue", "$79.8 billion", "2024", "RETRIEVED", "https://communityprod.blob.core.windows.net/public/top_100_2025_bi.pdf", "Business Insurance / Crain; brokers self-report. Excludes firms with >49% personal-lines gross revenue."),
        ("Retrieved — not lives", "A2", "BI Top 10 2024 U.S. brokerage revenue", "$53.6 billion", "2024", "RETRIEVED", "https://communityprod.blob.core.windows.net/public/top_100_2025_bi.pdf", "Adds Brown & Brown, Hub, Lockton, AssuredPartners, USI."),
        ("Retrieved — not lives", "A6", "Marsh+Aon+Gallagher+WTW+Lockton / Top 100", "44.1 percent", "2024", "DERIVED", "https://communityprod.blob.core.windows.net/public/top_100_2025_bi.pdf", "NFP inside Aon pro forma."),
        ("Retrieved — not lives", "A7", "MarshBerry citing BI: top 10 of Top 100", "70 percent", "2025", "RETRIEVED", "https://www.marshberry.com/resource/the-next-era-of-insurance-brokerage-the-top-100-defined-by-growth-consolidation-and-ai/", "Top 100 $83.3B, top 10 $58.5B. Different year — do not mix with 67%."),
        ("Retrieved — caveat", "D1", "U.S. small-group broker fees PMPM", "$29.79", "2024", "RETRIEVED", "https://www.kff.org/health-costs/state-indicator/health-insurance-broker-compensation/", "KFF (nonprofit, not Kaiser Permanente) analysis of Mark Farrah Associates Health Coverage Portal. Why compensation understates jumbo lives."),
        ("Retrieved — caveat", "D2", "U.S. large-group broker fees PMPM", "$10.00", "2024", "RETRIEVED", "https://www.kff.org/health-costs/state-indicator/health-insurance-broker-compensation/", "Large-group PMPM is about one-third of small-group."),
        ("Retrieved — not lives", "C1", "Marsh McLennan 2024 GAAP revenue", "$24.458 billion", "FY2024", "RETRIEVED", "https://www.corporate.marsh.com/web-assets/files-for-download/investors/2025/pdf-2024-marsh-mclennan-investors-annual-report.pdf", "Global. No US EH&B covered-lives figure."),
        ("Retrieved — not lives", "C3", "Aon Health Solutions 2024 revenue", "$3.335 billion", "FY2024", "RETRIEVED", "https://aon.mediaroom.com/2025-01-31-Aon-Reports-Fourth-Quarter-and-Full-Year-2024-Results", "Global. Includes NFP. Not US lives."),
        ("Specified M2", "X1", "Share of broker-placed EH&B covered lives at largest N brokers", "blank", "", "NOT RETRIEVED", "", "This is specified M2. 10-Ks and rankings do not publish it. Leave blank."),
    ],
    [
        ("Unknown", "X2", "Brokers below BI Top 100 and personal-lines-majority firms", "not retrieved", "", "NOT RETRIEVED", "https://communityprod.blob.core.windows.net/public/top_100_2025_bi.pdf", "67% is of the Top 100, not of U.S. brokerage."),
        ("Unknown", "X3", "Spectrum-state cut of any concentration figure", "not retrieved", "", "NOT RETRIEVED", "", "M2 geography is M1’s book. Public tables are U.S. No conversion invented."),
        ("Unknown", "X4", "Welfare plans / fees outside Form 5500 Schedule A as used here", "not retrieved", "", "NOT RETRIEVED", "https://planoptica.com/freeerisa", "Small welfare plans and jumbo consulting paid off Schedule A are missing."),
        ("Unknown", "—", "Sister-entity undercount", "not restated", "", "NOT RESTATED", "https://planoptica.com/welfare/broker", "Hub is Midwest only; WTW is US LLC only. Parent share ≥ filing-name share."),
    ],
    "The closest EH&B proxy says a handful of national partnerships does not cover the book (56% of observed compensation sits in a 43,119-name tail). That changes how many doors to open. It is not a product go-live. Lives share remains unknown, so this gate is directional, not closed.",
)

# ----- M5 -----
add_metric_sheet(
    "M5",
    [
        ("Identity", "M5", "Serviceable-home share of enrolled employees", "", "", "DEFINITION", "chtr-ehb-broker-channel-core-metrics.canvas.tsx", "Parent register: Decisions/EH&B/Brokerage Metrics.xlsx Core metrics M5"),
        ("Verdict", "V1", "Verdict on the metric", "valid — footprint hard gate", DATE, "ESTIMATED", "chtr-ehb-m5-serviceable-home-share.canvas.tsx", "Must match Analysis/ehb-broker-metrics/m5/out/summary.json"),
        ("Verdict", "V2", "Verdict on a number", "none — cannot claim", DATE, "ESTIMATED", "chtr-ehb-m5-serviceable-home-share.canvas.tsx", "No roster × BSL/addressability match."),
        ("Definition", "A1", "Numerator / denominator", "Enrolled (or broker-covered) employees whose residential address passes a Spectrum residential serviceability check ÷ enrolled/covered employees who supply a home address", "", "RETRIEVED", "Decisions/EH&B/Brokerage Metrics.xlsx", "Hard fail for out-of-footprint homes. Write one denominator before the pilot."),
        ("Definition", "S4", "What it gates", "Whether Spectrum can actually fulfill the benefit at the employee’s address", "", "DEFINITION", "Decisions/EH&B/Brokerage Metrics.xlsx", "If M5 is low and there is no multi-ISP fallback, the offering cannot be sold as a national included benefit."),
    ],
    [
        ("Specified M5", "A2", "M5 number for any broker book", "none", "", "BLOCKING", "", "No roster, no address match. This is M5. Leave blank."),
        ("Retrieved — not M5", "B1", "Charter-present jurisdictions (50 states + DC)", "42", "2025-12-31", "RETRIEVED", "https://broadbandmap.fcc.gov/data-download", "FCC BDC D25 revision 15 Sep 2026. Presence = any residential fixed-broadband filing. Includes Rhode Island at 106 estimated residential units."),
        ("Retrieved — not M5", "B2", "Charter-absent jurisdictions", "9", "2025-12-31", "RETRIEVED", "https://broadbandmap.fcc.gov/data-download", "AK, AR, DC, DE, IA, ND, OK, SD, UT."),
        ("Retrieved — not M5", "C1", "US SUSB 2022 payroll employment", "135,748,407", "2022", "RETRIEVED", "https://www2.census.gov/programs-surveys/susb/tables/2022/us_state_naics_detailedsizes_2022.txt", "U.S. Census Bureau SUSB 2022. Workplace establishment employment, not homes."),
        ("Retrieved — not M5", "C2", "SUSB employment in Charter-present jurisdictions", "128,435,628", "2022 × D25", "DERIVED", "Analysis/ehb-broker-metrics/m5/out/summary.json", "Workplace, not home."),
        ("Retrieved — not M5", "C3", "SUSB employment in Charter-absent jurisdictions", "7,312,779", "2022 × D25", "DERIVED", "Analysis/ehb-broker-metrics/m5/out/summary.json", "5.4% of US SUSB employment."),
        ("Retrieved — not M5", "C4", "C2 ÷ C1 (crude ceiling)", "94.6 percent", "2022 × D25", "ESTIMATED — not M5", "Analysis/ehb-broker-metrics/m5/out/summary.json", "Four biases, all high: workplace not home; any presence not plant; book ≠ nation; vintage join."),
        ("Retrieved — not M5", "D3", "Charter units ÷ fabric in the 42", "33.6 percent", "2025-12-31", "DERIVED — plant, not M5", "https://broadbandmap.fcc.gov/data-download", "50,939,445 ÷ 151,719,024. Homes Charter can serve among BSLs in states where it files, not employee homes."),
        ("Retrieved — not M5", "D-HI", "Hawaii Charter share of residential fabric", "92.3 percent", "2025-12-31", "RETRIEVED", "https://broadbandmap.fcc.gov/data-download", "Highest present-state location share. Random housing units are not enrolled employees."),
        ("Retrieved — not M5", "D-RI", "Rhode Island Charter share of residential fabric", "0.02 percent", "2025-12-31", "RETRIEVED", "https://broadbandmap.fcc.gov/data-download", "Lowest present-state location share (0.0002). Presence is not plant."),
        ("Retrieved — not M5", "D-med", "Unweighted median location share among the 42", "20.3 percent", "2025-12-31", "DERIVED", "Analysis/ehb-broker-metrics/m5/out/summary.json", "Plant statistic, still not M5."),
        ("Refused", "X1", "Employment-weighted state location share", "refused", "", "REFUSED", "Analysis/ehb-broker-metrics/m5/out/assumption_register.csv", "Two wrong substitutions: workplace for home, and a typical BSL for this employee. Not quoted as a rate."),
    ],
    [
        ("Unknown", "U1", "Employee home address × Spectrum residential serviceability on a named book", "not retrieved", "", "NOT RETRIEVED", "", "The blocking input. Consent, geocode, BSL match, ops serviceability (drop-ready / MDU / hold), then rate."),
        ("Unknown", "U2", "Denominator rule written before the pilot", "not retrieved", "", "MUST WRITE", "", "Eligible-with-address for a go/no-go on the book; enrolled-with-address for fulfillment. Indeterminate is a third bucket."),
    ],
    "If a consented address match on the candidate book shows that Spectrum cannot fulfill the benefit as written for the homes the contract covers, and the design has no multi-ISP fallback for the complement, do not scale a national “included Spectrum Internet” benefit. The public 94.6% ceiling cannot clear or kill this. Only the roster match can.",
)

# ----- M9 -----
add_metric_sheet(
    "M9",
    [
        ("Identity", "M9", "Payor model and enrollment take-up", "", "", "DEFINITION", "chtr-ehb-broker-channel-core-metrics.canvas.tsx", "Parent register already named two rates under one ID."),
        ("Verdict", "V1", "Verdict", "needs split — valid gate, two metrics", DATE, "ESTIMATED", "chtr-ehb-m9-payor-takeup.canvas.tsx", "Must match Analysis/ehb-broker-metrics/m9/summary.json. No Charter take-up invented."),
        ("Split", "M9a", "Payor mix", "Share of pilot employers (and of covered lives) on employer-paid vs employee-paid voluntary vs hybrid", "", "DEFINITION", "Analysis/ehb-broker-metrics/m9/recommended_split.csv", "Employer unit. Freeze at contracting, before open enrollment."),
        ("Split", "M9b", "Enrollment take-up by payor", "Enrolled internet accounts ÷ eligible employees offered the benefit, reported separately inside each M9a cell", "", "DEFINITION", "Analysis/ehb-broker-metrics/m9/recommended_split.csv", "Employee unit. Time-box: OE close + 90 days. Never headline a blend."),
        ("Split", "M9_blended", "Do not report a blended enrolled/eligible headline", "forbidden", "", "DEFINITION", "Analysis/ehb-broker-metrics/m9/recommended_split.csv", "100% of a tiny paid cell plus 4% of a large voluntary cell can look mid-teens."),
    ],
    [
        ("Specified M9", "X2", "Charter internet-as-benefit take-up", "blank", "", "NOT RETRIEVED", "", "No SKU, price sheet, or enrollment file in this repo. Pilot is the source."),
        ("Analogue — not Charter", "D1", "Employers offering free / discounted / subsidized home Internet", "5 percent of employers", "2017", "RETRIEVED", "https://www.shrm.org/content/dam/en/shrm/topics-tools/news/organizational-employee-development/Benefits-Prevalence-Report-All-Industries-All-FTEs.pdf", "SHRM customized benefits-prevalence, n≈2,717. Closest internet-as-benefit OFFER analogue. Not take-up. Dated."),
        ("Analogue — wrong unit", "A1", "KFF medical take-up among eligible workers", "76 percent of eligible workers", "2025", "RETRIEVED", "https://www.kff.org/health-costs/2025-employer-health-benefits-survey/", "WRONG ANALOGUE for home internet. Heavily employer-subsidized core medical. KFF (nonprofit, not Kaiser Permanente); Davis Research; n=1,862 employers 10+."),
        ("Analogue — not Charter", "A2", "Workers eligible at firms offering health benefits", "80 percent of workers at offering firms", "2025", "RETRIEVED", "https://www.kff.org/health-costs/2025-employer-health-benefits-survey/", "Eligibility, not take-up, not internet."),
        ("Analogue — not Charter", "A3", "Workers enrolled among firms that offer health benefits", "61 percent of workers at offering firms", "2025", "RETRIEVED", "https://www.kff.org/health-costs/2025-employer-health-benefits-survey/", "Product of eligibility × take-up."),
        ("Analogue — not Charter", "B1", "Employers offering pet insurance", "22 percent of employers", "2025", "RETRIEVED", "https://www.shrm.org/content/dam/en/shrm/topics-tools/research/employee-benefits/2025_annual_benefits_survey_executive_summary.pdf", "SHRM 2025 EBS, n=3,969. Offer, not take-up."),
        ("Analogue — not Charter", "E1", "Employers offering a transit subsidy", "12 percent of employers", "2025", "RETRIEVED", "https://www.shrm.org/topics-tools/news/benefits-compensation/irs-boosts-commuter-benefit-limits-2026", "Offer, not take-up."),
        ("Analogue — not Charter", "E2", "Employers offering a parking subsidy", "10 percent of employers", "2025", "RETRIEVED", "https://www.shrm.org/topics-tools/news/benefits-compensation/irs-boosts-commuter-benefit-limits-2026", "Offer, not take-up."),
        ("Analogue — not Charter", "C1", "Employers offering QSLP 401(k)/403(b) match", "4 percent of employers", "2025", "RETRIEVED", "https://www.shrm.org/content/dam/en/shrm/topics-tools/research/employee-benefits/2025_annual_benefits_survey_executive_summary.pdf", "SECURE 2.0 QSLP. Offer, not take-up."),
        ("Analogue — not internet", "D2", "Employers offering at-home equipment subsidy", "55 percent of employers", "2025", "RETRIEVED", "https://www.shrm.org/content/dam/en/shrm/topics-tools/research/employee-benefits/2025_annual_benefits_survey_executive_summary.pdf", "Equipment, not an internet stipend. Do not treat as D1."),
        ("Analogue — not internet", "D3", "Employers offering at-home equipment subsidy", "53 percent of employers", "2026", "RETRIEVED", "https://www.shrm.org/content/dam/en/shrm/topics-tools/research/employee-benefits/2026_employee-benefits_executive-summary.pdf", "Same caveat as D2. Cellphone service 35% of 2026 equipment programs. Internet not broken out."),
        ("Excluded", "X1", "Stealth Agents “SHRM internet-specific stipend”", "38 percent — NOT USED", "2025", "NOT RETRIEVED FROM PRIMARY", "https://stealthagents.com/research/remote-work-internet-stipend-statistics-2026", "Secondary restatement. SHRM primary publishes equipment subsidy, not this 38%."),
    ],
    [
        ("Unknown", "U1", "Any Charter internet-as-benefit take-up rate", "not retrieved", "", "NOT RETRIEVED", "", "No SKU in repo. Analogues are not a forecast."),
        ("Unknown", "U2", "2026 prevalence of an employer internet stipend", "not retrieved", "", "NOT RETRIEVED", "", "2017 5% is dated. 2025/2026 SHRM exec summaries do not publish an internet-specific rate."),
        ("Unknown", "U3", "Payor mix on a real ballot (M9a)", "not retrieved", "", "MUST PULL", "Analysis/ehb-broker-metrics/m9/pilot_must_measure.csv", "Measure at contracting, before OE."),
    ],
    "If no pilot employer will fund employer-paid and none will put a voluntary SKU on the ballot — stop. If a design is on the ballot and voluntary take-up among eligibles in that cell, after OE plus 90 days, lands in the single digits with no path to change price or payor — do not scale on the hope that “benefits take-up is usually high.” Analogues bound the plausible range. Analogues are not the floor. Write the floor before the first employee sees the offer.",
)

# ----- M13 -----
add_metric_sheet(
    "M13",
    [
        ("Identity", "M13", "Net new residential Internet accounts from the channel", "", "", "DEFINITION", "chtr-ehb-broker-channel-core-metrics.canvas.tsx", "Parent register: Decisions/EH&B/Brokerage Metrics.xlsx Core metrics M13"),
        ("Verdict", "V1", "Verdict on M13 as specified", "valid-with-caveats — tighten to P−E−D", DATE, "ESTIMATED", "chtr-ehb-m13-net-new-residential.canvas.tsx", "Right commercial outcome (Internet print in accounts). Excel mixed enrollments into the unit and omitted same-window job-exit D."),
        ("Definition", "F0", "Tightened formula", "M13_t = P_t − E_t − D_t", "", "ESTIMATED", "Analysis/ehb-broker-metrics/m13/assumption_register.csv", "Primary KPI is the account count. Window t = Ex99.1 calendar quarter locked before the pull. Do not report M13 as a percent of enrollments."),
        ("Definition", "P", "Path-tagged provisions P_t", "blank", "", "MUST PULL", "", "Gross broker/employer-path billing events on a residential Internet PSU in t. Not enrollments; not mobile; not Spectrum Business."),
        ("Definition", "E", "Existing / payer-only E_t", "blank", "", "MUST PULL", "", "Subset of P already active Spectrum residential Internet at t−1. Retention/transfer — not a net add."),
        ("Definition", "D", "Same-window job-exit disconnects D_t", "blank", "", "MUST PULL", "", "Subset of (P−E) that disconnect in t because employment eligibility ended. Do not subtract self-pay stayers."),
    ],
    [
        ("Specified M13", "M13", "Channel net new residential Internet accounts", "blank", "", "MUST PULL / NOT RETRIEVED", "", "P − E − D. No number exists for this path. Do not invent conversion rates."),
        ("Context only — not M13", "X1", "Company residential Internet net adds Q1 2026", "-117000", "Q1 2026", "RETRIEVED", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Ex99.1: Residential quarterly net additions (117). NOT channel M13. Do not scale."),
        ("Context only — not M13", "X2", "Company residential Internet customers decreased FY2025", "-393000", "FY2025", "RETRIEVED", "PrimarySources/readable/CHTR-10-K-FY2025.md", "10-K: decreased by 393,000 in 2025 vs 2024. Context only."),
        ("Context only — not M13", "X3", "Company residential Internet net adds Q2 2026", "-166000", "Q2 2026", "RETRIEVED", "PrimarySources/readable/CHTR-Q2-2026-Earnings-Release-Ex99.1.md", "Ex99.1: Residential quarterly net additions (166). Context only."),
        ("Context only — not M13", "X4", "Company total Internet net adds Q1 2026", "-120000", "Q1 2026", "RETRIEVED", "PrimarySources/readable/CHTR-Q1-2026-Earnings-Release-Ex99.1.md", "Headline −120k = residential −117k + SMB −3k. Do not use as the M13 analog; M13 is residential only."),
        ("Data needed", "D1", "Channel tag (employer ID + broker/path code)", "missing", "", "MUST PULL", "", "Without it M13 cannot be isolated from ordinary adds."),
        ("Data needed", "D2", "Prior-account flag (active Spectrum res. Internet at t−1)", "missing", "", "MUST PULL", "", "Computes E. Customer ID preferred; address match is fallback."),
        ("Data needed", "D3", "Payer type (employer / employee / self-pay)", "missing", "", "MUST PULL", "", "Separates discounting an existing line from a new relationship."),
        ("Data needed", "D4", "Eligibility termination date", "missing", "", "MUST PULL", "", "Join to disconnects to compute D. Benefit unenroll ≠ Spectrum disconnect."),
        ("Data needed", "D5", "Disconnect date and whether PSU ended", "missing", "", "MUST PULL", "", "D is Spectrum disconnect in t; self-pay stayers stay in the print."),
        ("Data needed", "D6", "Measurement window locked to Ex99.1 quarter", "missing", "", "MUST PULL", "", "So M13 is in the same unit as the company print."),
    ],
    [
        ("Unknown", "U1", "Channel net-add count", "not retrieved", "", "NOT RETRIEVED", "", "Do not invent."),
        ("Unknown", "U2", "Conversion of enrollments into accounts", "not retrieved", "", "NOT RETRIEVED", "", "Optional diagnostic (P−E)/enrollments — leave blank until a pilot produces both counts."),
        ("Unknown", "U3", "Share of company print this path would close", "not retrieved", "", "NOT RETRIEVED", "", "Do not scale −117k / −393k by an invented take-up."),
        ("Unknown", "U4", "Incrementality (tagged ≠ would-not-have-signed-up)", "not estimated", "", "NOT ESTIMATED", "", "A holdout test is a later design, not required to compute the tagged count."),
    ],
    "Kill the channel as an Internet-print play if a time-boxed pilot with D1–D5 filled shows M13 ≈ 0 because almost all of P is E or D — even if enrollments are large. Do not substitute enrollments, passings, or reach.",
)

# ----- Analogues (M9 only) -----
ws = wb.create_sheet("Analogues")
an_headers = [
    "Label",
    "Line ID",
    "Family",
    "Metric role",
    "What it measures",
    "Value",
    "Unit",
    "Year",
    "Origin",
    "Funder",
    "Source URL",
    "Why it is not Charter take-up",
]
ws.append(an_headers)
for r in analogues:
    role = (r.get("metric_role") or "").strip()
    if role == "do_not_use":
        label = "EXCLUDED — not retrieved from named primary"
    elif (r.get("analogue_family") or "") == "charter_sku":
        label = "ABSENCE — no Charter SKU in repo"
    else:
        label = "ANALOGUE — not Charter take-up"
    ws.append(
        [
            label,
            r.get("line_id"),
            r.get("analogue_family"),
            r.get("metric_role"),
            r.get("what_it_measures"),
            r.get("value"),
            r.get("unit"),
            r.get("year"),
            r.get("origin"),
            r.get("funder"),
            r.get("source_url"),
            r.get("not_charter_because"),
        ]
    )
    if label.startswith("EXCLUDED") or label.startswith("ABSENCE"):
        for c in range(1, 13):
            ws.cell(ws.max_row, c).fill = blank_fill
style_header(ws, 12)
set_widths(ws, [36, 10, 22, 26, 48, 14, 28, 10, 16, 36, 62, 55])
wrap_all(ws)
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 32

if XLSX_PATH.exists():
    try:
        XLSX_PATH.unlink()
    except OSError as e:
        print("WARN unlink XLSX", e)

wb.save(XLSX_PATH)
print("XLSX_SAVED", XLSX_PATH)
print("SHEETS", wb.sheetnames)

# ---------------------------------------------------------------------------
# Verify
# ---------------------------------------------------------------------------
xlsx_bytes = XLSX_PATH.read_bytes()
print("XLSX_BYTES", len(xlsx_bytes))
print("XLSX_HEADER", xlsx_bytes[:2])
assert xlsx_bytes[:2] == b"PK", "xlsx is not a zip"

import pypdf

reader = pypdf.PdfReader(str(PDF_PATH))
print("PDF_PAGES", len(reader.pages))
raw = "\n".join(p.extract_text() or "" for p in reader.pages)
text = (
    raw.replace("\u2014", "-")
    .replace("\u2013", "-")
    .replace("\u2212", "-")
    .replace("\xa0", " ")
)
print("PDF_TEXT_LEN", len(raw))

checks = [
    "Five gates are specified",
    "numerator blank",
    "wrong unit",
    "no number",
    "Needs split",
    "M9a",
    "P-E-D",
    "168.2",
    "178.2",
    "45.9%",
    "44%",
    "67%",
    "94.6%",
    "92%",
    "0.02%",
    "5%",
    "76%",
    "117",
    "not M1",
    "not retrieved",
    "kff.org",
    "planoptica.com",
    "broadbandmap.fcc.gov",
    "shrm.org",
]
for c in checks:
    print(f"FOUND[{c}]={c in text}")

wb2 = load_workbook(XLSX_PATH, data_only=True)
print("XLSX_SHEETS", wb2.sheetnames)
expected_sheets = ["Notes", "Scorecard", "M1", "M2", "M5", "M9", "M13", "Analogues"]
assert wb2.sheetnames == expected_sheets, wb2.sheetnames

# Scorecard five rows
sc = wb2["Scorecard"]
assert sc.max_row == 6, sc.max_row
assert sc["A2"].value == "M1"
assert sc["A6"].value == "M13"
assert "blank" in str(sc["E2"].value).lower()
assert "blank" in str(sc["E6"].value).lower()

# Specified blanks exist
m1 = wb2["M1"]
m1_vals = {(row[1], row[3]) for row in m1.iter_rows(min_row=2, values_only=True)}
assert any(i == "F1" and str(v).lower() == "blank" for i, v in m1_vals), m1_vals
assert any(i == "X1" and "45.9" in str(v) for i, v in m1_vals)

m2 = wb2["M2"]
m2_map = {row[1]: row[3] for row in m2.iter_rows(min_row=2, values_only=True)}
assert "44.0" in str(m2_map.get("B2"))
assert str(m2_map.get("X1")).lower() == "blank"

m5 = wb2["M5"]
m5_map = {row[1]: row[3] for row in m5.iter_rows(min_row=2, values_only=True)}
assert str(m5_map.get("A2")).lower() == "none"
assert "94.6" in str(m5_map.get("C4"))

m9 = wb2["M9"]
m9_map = {row[1]: row[3] for row in m9.iter_rows(min_row=2, values_only=True)}
assert str(m9_map.get("X2")).lower() == "blank"
assert "5" in str(m9_map.get("D1"))
assert "76" in str(m9_map.get("A1"))

m13 = wb2["M13"]
m13_map = {row[1]: row[3] for row in m13.iter_rows(min_row=2, values_only=True)}
assert str(m13_map.get("M13")).lower() == "blank"
assert str(m13_map.get("P")).lower() == "blank"
assert "-117000" in str(m13_map.get("X1"))

an = wb2["Analogues"]
labels = {row[0] for row in an.iter_rows(min_row=2, values_only=True)}
assert any(str(x).startswith("ANALOGUE") for x in labels)
assert any("EXCLUDED" in str(x) for x in labels)

# Notes says blanks intentional
notes_map = {row[0]: row[1] for row in wb2["Notes"].iter_rows(min_row=2, values_only=True)}
assert "Blanks are intentional" in notes_map
assert "on purpose" in notes_map["Blanks are intentional"].lower()

print("VERIFY_OK")
print("PDF_PATH", PDF_PATH)
print("XLSX_PATH", XLSX_PATH)
