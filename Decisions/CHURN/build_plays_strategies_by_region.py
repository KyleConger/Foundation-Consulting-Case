"""
Build Decisions/CHURN/Plays Strategies by Region.xlsx

Locked team hypothesis: play × strategy scores from the USER (not the
earlier agent verdict that EH&B is no on plays 2 and 5).

Regional sheets are plant / competitive CONTEXT from already-retrieved
FCC BDC state leads. They do not diagnose which churn play is live in
a division. No TAM. No launch rec.

Sources:
- Plays 1–6: CHTR_Rec1_AttributionPlaybook.pptx
- State leads: Analysis/cable-share-small-firm/out/chtr_vs_cmcsa_state_leads.csv
  (FCC BDC D25, as of 2025-12-31, revision 15 Sep 2026)
- Census divisions: U.S. Census Bureau (9 divisions)
- Charter RDOF locations: Analysis/chtr-rural-geography.canvas.tsx
  (FCC Auction 904, CCO Holdings) — ESTIMATED rural context only
"""
from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "Plays Strategies by Region.xlsx"
LEADS = ROOT / "Analysis" / "cable-share-small-firm" / "out" / "chtr_vs_cmcsa_state_leads.csv"

GEORGIA = "Georgia"
NAVY = "1F4E79"
INK = "1A1A1A"
GRAY = "595959"

thin = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)

font = Font(name=GEORGIA, size=10, color=INK)
font_b = Font(name=GEORGIA, size=10, bold=True, color=INK)
font_title = Font(name=GEORGIA, size=16, bold=True, color=NAVY)
font_h2 = Font(name=GEORGIA, size=12, bold=True, color=NAVY)
font_h = Font(name=GEORGIA, size=10, bold=True, color="FFFFFF")
font_note = Font(name=GEORGIA, size=9, italic=True, color=GRAY)
font_small = Font(name=GEORGIA, size=8, color=GRAY)

fill_head = PatternFill("solid", fgColor=NAVY)
fill_best = PatternFill("solid", fgColor="C6EFCE")
fill_esp = PatternFill("solid", fgColor="D6E3F0")
fill_works = PatternFill("solid", fgColor="F2F2F2")
fill_no = PatternFill("solid", fgColor="F8CBAD")
fill_unspec = PatternFill("solid", fgColor="EEEEEE")
fill_chtr = PatternFill("solid", fgColor="D6E3F0")
fill_cmcsa = PatternFill("solid", fgColor="F2F2F2")
fill_neither = PatternFill("solid", fgColor="FFF2CC")
fill_est = PatternFill("solid", fgColor="DDEBF7")
fill_white = PatternFill("solid", fgColor="FFFFFF")

wrap = Alignment(wrap_text=True, vertical="center")
left = Alignment(wrap_text=True, vertical="center", horizontal="left")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")

# Census Bureau divisions (9). DC is in South Atlantic.
DIVISION: dict[str, tuple[str, str]] = {
    "CT": ("New England", "Northeast"),
    "ME": ("New England", "Northeast"),
    "MA": ("New England", "Northeast"),
    "NH": ("New England", "Northeast"),
    "RI": ("New England", "Northeast"),
    "VT": ("New England", "Northeast"),
    "NJ": ("Middle Atlantic", "Northeast"),
    "NY": ("Middle Atlantic", "Northeast"),
    "PA": ("Middle Atlantic", "Northeast"),
    "IL": ("East North Central", "Midwest"),
    "IN": ("East North Central", "Midwest"),
    "MI": ("East North Central", "Midwest"),
    "OH": ("East North Central", "Midwest"),
    "WI": ("East North Central", "Midwest"),
    "IA": ("West North Central", "Midwest"),
    "KS": ("West North Central", "Midwest"),
    "MN": ("West North Central", "Midwest"),
    "MO": ("West North Central", "Midwest"),
    "NE": ("West North Central", "Midwest"),
    "ND": ("West North Central", "Midwest"),
    "SD": ("West North Central", "Midwest"),
    "DE": ("South Atlantic", "South"),
    "DC": ("South Atlantic", "South"),
    "FL": ("South Atlantic", "South"),
    "GA": ("South Atlantic", "South"),
    "MD": ("South Atlantic", "South"),
    "NC": ("South Atlantic", "South"),
    "SC": ("South Atlantic", "South"),
    "VA": ("South Atlantic", "South"),
    "WV": ("South Atlantic", "South"),
    "AL": ("East South Central", "South"),
    "KY": ("East South Central", "South"),
    "MS": ("East South Central", "South"),
    "TN": ("East South Central", "South"),
    "AR": ("West South Central", "South"),
    "LA": ("West South Central", "South"),
    "OK": ("West South Central", "South"),
    "TX": ("West South Central", "South"),
    "AZ": ("Mountain", "West"),
    "CO": ("Mountain", "West"),
    "ID": ("Mountain", "West"),
    "MT": ("Mountain", "West"),
    "NV": ("Mountain", "West"),
    "NM": ("Mountain", "West"),
    "UT": ("Mountain", "West"),
    "WY": ("Mountain", "West"),
    "AK": ("Pacific", "West"),
    "CA": ("Pacific", "West"),
    "HI": ("Pacific", "West"),
    "OR": ("Pacific", "West"),
    "WA": ("Pacific", "West"),
}

DIV_ORDER = [
    "New England",
    "Middle Atlantic",
    "East North Central",
    "West North Central",
    "South Atlantic",
    "East South Central",
    "West South Central",
    "Mountain",
    "Pacific",
]

# FCC Auction 904 RDOF locations awarded to CCO Holdings.
# Same figures as Analysis/chtr-rural-geography.canvas.tsx. Context only.
RDOF: dict[str, int] = {
    "WI": 143269,
    "TX": 133993,
    "NC": 128502,
    "OH": 112777,
    "SC": 98670,
    "TN": 79193,
    "MO": 61524,
    "AL": 56451,
    "IN": 54541,
    "MI": 35944,
    "KY": 31747,
    "LA": 25389,
    "GA": 23854,
    "FL": 17869,
    "OR": 15139,
    "MA": 14344,
    "VA": 11369,
    "PA": 5328,
    "WA": 4625,
    "CA": 1045,
    "NH": 1044,
    "IL": 501,
    "NM": 485,
    "VT": 85,
}

PLAYS = [
    (1, "FWA-driven losses"),
    (2, "Fiber-driven losses"),
    (3, "LEO-satellite-driven losses"),
    (4, "Housing-driven losses"),
    (5, "Economic / non-pay churn"),
    (6, "Null result — no dominant driver"),
]

STRATS = [
    "EH&B",
    "Partner spend",
    "Reallocating rural",
    "Existing capex",
]

# Locked USER scores. Do not replace EH&B works-on-2-and-5 with the older no.
# Partner 2,3,5,6: unspecified — Decisions/ has no partner-spend score.
# Existing capex 2–5: unspecified — IMP TO CAPEX / six-areas do not score
# those plays as "works"; Area 3 reserves incremental P1 rebuild for fiber
# + speed-named sample (spend-mix constraint, not a user "works" score).
MATRIX: dict[int, dict[str, str]] = {
    1: {
        "EH&B": "Best",
        "Partner spend": "Best",
        "Reallocating rural": "Works",
        "Existing capex": "Best",
    },
    2: {
        "EH&B": "Works",
        "Partner spend": "Unspecified",
        "Reallocating rural": "Especially",
        "Existing capex": "Unspecified",
    },
    3: {
        "EH&B": "No",
        "Partner spend": "Unspecified",
        "Reallocating rural": "Especially",
        "Existing capex": "Unspecified",
    },
    4: {
        "EH&B": "Best",
        "Partner spend": "Best",
        "Reallocating rural": "Works",
        "Existing capex": "Unspecified",
    },
    5: {
        "EH&B": "Works",
        "Partner spend": "Unspecified",
        "Reallocating rural": "Especially",
        "Existing capex": "Unspecified",
    },
    6: {
        "EH&B": "No",
        "Partner spend": "Unspecified",
        "Reallocating rural": "Works",
        "Existing capex": "Best",
    },
}

FIT_FILL = {
    "Best": fill_best,
    "Especially": fill_esp,
    "Works": fill_works,
    "No": fill_no,
    "Unspecified": fill_unspec,
}


def fnum(s: str) -> float:
    if s is None or str(s).strip() == "":
        return 0.0
    return float(s)


def truthy(s: str) -> bool:
    return str(s).strip().lower() == "true"


def style_range(ws: Worksheet, row: int, cols: int, fill=None, font_=None, align=None):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.border = thin
        cell.font = font_ or font
        cell.alignment = align or wrap
        if fill is not None:
            cell.fill = fill


def widths(ws: Worksheet, widths_list: list[float]) -> None:
    for i, w in enumerate(widths_list, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def load_states() -> list[dict]:
    rows: list[dict] = []
    with LEADS.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            abbr = r["abbr"]
            div, region = DIVISION[abbr]
            leader = r["leader"]
            if leader == "neither present":
                leader_n = "Neither present"
            else:
                leader_n = leader
            rows.append(
                {
                    "fips": r["state_fips"],
                    "state": r["state"],
                    "abbr": abbr,
                    "division": div,
                    "region": region,
                    "leader": leader_n,
                    "charter_present": truthy(r["charter_present"]),
                    "comcast_present": truthy(r["comcast_present"]),
                    "charter_pct": fnum(r["charter_res_st_pct"]),
                    "comcast_pct": fnum(r["comcast_res_st_pct"]),
                    "charter_units": fnum(r["charter_est_units"]),
                    "comcast_units": fnum(r["comcast_est_units"]),
                    "margin_units": fnum(r["margin_units_charter_minus_comcast"]),
                    "fabric": fnum(r["total_residential_units"]),
                    "rdof": RDOF.get(abbr, 0),
                    "origin": "RETRIEVED",
                }
            )
    rows.sort(key=lambda x: (DIV_ORDER.index(x["division"]), x["state"]))
    return rows


def rollup(states: list[dict]) -> list[dict]:
    buckets: dict[str, dict] = {}
    for d in DIV_ORDER:
        buckets[d] = {
            "division": d,
            "region": next(s["region"] for s in states if s["division"] == d),
            "n": 0,
            "chtr_lead": 0,
            "cmcsa_lead": 0,
            "neither": 0,
            "chtr_present": 0,
            "margin_units": 0.0,
            "rdof": 0,
        }
    for s in states:
        b = buckets[s["division"]]
        b["n"] += 1
        if s["leader"] == "Charter":
            b["chtr_lead"] += 1
        elif s["leader"] == "Comcast":
            b["cmcsa_lead"] += 1
        else:
            b["neither"] += 1
        if s["charter_present"]:
            b["chtr_present"] += 1
        b["margin_units"] += s["margin_units"]
        b["rdof"] += s["rdof"]
    return [buckets[d] for d in DIV_ORDER]


def eligibility(row: dict) -> str:
    plant = (
        "Existing-plant strategies (EH&B, partner spend, existing capex) "
        f"are even eligible in the {row['chtr_present']} Charter-present "
        "state(s) — plant already reported in FCC BDC, not a play diagnosis."
    )
    if row["rdof"] >= 100_000:
        rural = (
            f"Rural reallocation is more plausible here: {row['rdof']:,} "
            "Charter RDOF locations (Auction 904) cluster in this division."
        )
    elif row["rdof"] > 0:
        rural = (
            f"Rural reallocation is thinner: {row['rdof']:,} Charter RDOF "
            "locations (Auction 904)."
        )
    else:
        rural = "No Charter RDOF locations in the Auction 904 CCO Holdings file for this division."
    return plant + " " + rural


def write_notes(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False
    widths(ws, [28, 110])
    ws["A1"] = (
        "Team hypothesis: EH&B is best on FWA and housing; "
        "only rural reallocation is scored as working on every play"
    )
    ws["A1"].font = font_title
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 28

    notes = [
        (
            "What this is",
            "Play × strategy fit scores locked to the USER. "
            "Regional sheets are ESTIMATED plant / competitive CONTEXT "
            "(Charter vs Comcast location lead; Charter RDOF awards). "
            "They do not say which churn play is live in a division. "
            "No privatized churn mix. No TAM. No launch recommendation.",
        ),
        (
            "Hypothesis, not prior verdict",
            "THIS is the team's hypothesis. An earlier validation said "
            "EH&B is no on plays 2 (fiber) and 5 (economic / non-pay). "
            "This workbook keeps the team's scores: EH&B works on 2 and 5. "
            "Do not silently revert.",
        ),
        (
            "Plays (pptx)",
            "1 FWA-driven losses; 2 Fiber-driven losses; "
            "3 LEO-satellite-driven losses; 4 Housing-driven losses; "
            "5 Economic / non-pay churn; 6 Null result — no dominant driver. "
            "Source: Decisions/CHURN/CHTR_Rec1_AttributionPlaybook.pptx.",
        ),
        (
            "EH&B (user)",
            "Best in 1 and 4; works in 2 and 5; does not work in 3 and 6. "
            "EH&B = existing homes and businesses on plant already built "
            "(Decisions/decisions.md; Decisions/EH&B/).",
        ),
        (
            "Partner spend (user)",
            "Works best in 1 and 4. Plays 2, 3, 5, 6 left Unspecified: "
            "the user did not score them, and no Decisions/ file scores "
            "partner spend. Do not invent 'works'.",
        ),
        (
            "Reallocating rural (user)",
            "Works everywhere; especially 2, 3, and 5. "
            "Folder: Decisions/IMP TO CAPEX/ (rural take-not-miles, not more miles).",
        ),
        (
            "Existing capex (user)",
            "Works best in 1 and 6. Plays 2–5 left Unspecified. "
            "Decisions/IMP TO CAPEX/ and six-recommendation-areas.md do not "
            "score those plays as 'works'. Area 3 reserves incremental P1 "
            "(+$280M upgrade/rebuild) for fiber + a speed-named sample on "
            "overlap — a spend-mix constraint, not a user 'works' score on play 2.",
        ),
        (
            "How regions are used",
            "Census divisions (9). Each FCC BDC state is mapped to a division. "
            "Rollup = counts of Charter-lead vs Comcast-lead states, plus "
            "Charter-present count and Charter-minus-Comcast unit margin "
            "(same CSV). RDOF locations are rural-award CONTEXT. "
            "No regional 'which play is live' number exists; none is invented.",
        ),
        (
            "State-lead source",
            "Analysis/cable-share-small-firm/out/chtr_vs_cmcsa_state_leads.csv. "
            "FCC Broadband Data Collection, National Broadband Map, as of "
            "31 Dec 2025 (D25), file revision 15 Sep 2026. Availability / "
            "locations served — not subscriber share. Beating = strictly "
            "larger residential share. Missing provider row = absent.",
        ),
        (
            "RDOF source",
            "FCC Auction 904 locations awarded to CCO Holdings, as compiled "
            "in Analysis/chtr-rural-geography.canvas.tsx (~1.06M locations, "
            "24 states). ESTIMATED rural context. Not every ARPA/BEAD/state "
            "grant location.",
        ),
        (
            "Census geography",
            "U.S. Census Bureau divisions and regions. District of Columbia "
            "is in South Atlantic. Not a Charter operating-region file.",
        ),
        (
            "Unwilling to claim",
            "Which play dominates any division. A filled disconnect mix. "
            "That partner spend or existing capex 'works' on unscored plays. "
            "That EH&B is no on 2 and 5 (the older agent verdict). "
            "A TAM, SOM, or launch recommendation.",
        ),
    ]
    ws["A3"] = "Notes"
    ws["A3"].font = font_h2
    r = 4
    for label, body in notes:
        ws.cell(r, 1, label).font = font_b
        ws.cell(r, 1).alignment = wrap
        ws.cell(r, 2, body).font = font
        ws.cell(r, 2).alignment = wrap
        ws.row_dimensions[r].height = 48
        r += 1
    ws.freeze_panes = "A4"


def write_matrix(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False
    widths(ws, [8, 32, 18, 18, 22, 18, 56])
    ws["A1"] = (
        "Locked team scores — Best / Works / Especially / No / Unspecified"
    )
    ws["A1"].font = font_title
    ws.merge_cells("A1:G1")
    ws.row_dimensions[1].height = 24
    ws["A2"] = (
        "User scoring. EH&B works on 2 and 5 (not the earlier no). "
        "Blank-looking Unspecified cells are intentional."
    )
    ws["A2"].font = font_note
    ws.merge_cells("A2:G2")

    headers = [
        "Play",
        "Name",
        "EH&B",
        "Partner spend",
        "Reallocating rural",
        "Existing capex",
        "Source of score",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.font = font_h
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
    sources = {
        1: "User: EH&B Best; partner Best; rural Works; capex Best",
        2: "User: EH&B Works; rural Especially; partner + capex Unspecified",
        3: "User: EH&B No; rural Especially; partner + capex Unspecified",
        4: "User: EH&B Best; partner Best; rural Works; capex Unspecified",
        5: "User: EH&B Works; rural Especially; partner + capex Unspecified",
        6: "User: EH&B No; rural Works; capex Best; partner Unspecified",
    }
    for i, (n, name) in enumerate(PLAYS):
        r = 5 + i
        ws.cell(r, 1, n).font = font_b
        ws.cell(r, 2, name).font = font
        for c, strat in enumerate(STRATS, 3):
            val = MATRIX[n][strat]
            cell = ws.cell(r, c, val)
            cell.font = font_b if val == "Best" else font
            cell.fill = FIT_FILL[val]
            cell.alignment = center
            cell.border = thin
        ws.cell(r, 7, sources[n]).font = font_small
        ws.cell(r, 7).alignment = wrap
        for c in (1, 2, 7):
            ws.cell(r, c).border = thin
            ws.cell(r, c).alignment = wrap if c != 1 else center
        ws.row_dimensions[r].height = 28

    ws["A12"] = "Legend"
    ws["A12"].font = font_h2
    legend = [
        ("Best", "User: works best on this play", fill_best),
        ("Especially", "User: works everywhere, especially this play (rural only)", fill_esp),
        ("Works", "User: works on this play", fill_works),
        ("No", "User: does not work on this play", fill_no),
        (
            "Unspecified",
            "User did not score; Decisions/ does not fill the cell. Not a 'works'.",
            fill_unspec,
        ),
    ]
    for i, (lab, desc, fill) in enumerate(legend):
        r = 13 + i
        ws.cell(r, 1, lab).fill = fill
        ws.cell(r, 1).font = font_b
        ws.cell(r, 1).border = thin
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2, desc).font = font
        ws.cell(r, 2).border = thin

    ws["A19"] = "Counts of plays scored as working (Best + Especially + Works)"
    ws["A19"].font = font_h2
    ws.merge_cells("A19:C19")
    for c, h in enumerate(["Strategy", "Working (n of 6)", "Of which Best / Especially"], 1):
        cell = ws.cell(20, c, h)
        cell.font = font_h
        cell.fill = fill_head
        cell.border = thin
    counts = []
    for strat in STRATS:
        working = sum(
            1
            for n in range(1, 7)
            if MATRIX[n][strat] in ("Best", "Especially", "Works")
        )
        strong = sum(
            1
            for n in range(1, 7)
            if MATRIX[n][strat] in ("Best", "Especially")
        )
        counts.append((strat, working, strong))
    for i, (strat, working, strong) in enumerate(counts):
        r = 21 + i
        ws.cell(r, 1, strat).font = font
        ws.cell(r, 2, working).font = font_b
        ws.cell(r, 3, strong).font = font
        for c in range(1, 4):
            ws.cell(r, c).border = thin
            ws.cell(r, c).alignment = center if c > 1 else left
    ws["A26"] = (
        "Rural = 6/6 working. EH&B = 4/6 (Best on 1 and 4). "
        "Partner = 2/6 (Best on 1 and 4). Existing capex = 2/6 (Best on 1 and 6)."
    )
    ws["A26"].font = font_note
    ws.merge_cells("A26:F26")
    ws.freeze_panes = "A5"


def write_states(ws: Worksheet, states: list[dict]) -> None:
    ws.sheet_view.showGridLines = False
    headers = [
        "State",
        "Abbr",
        "Census division",
        "Census region",
        "Location lead",
        "Charter present",
        "Comcast present",
        "Charter est. units",
        "Comcast est. units",
        "Margin (CHTR − CMCSA)",
        "State fabric units",
        "Charter RDOF locs",
        "Tag",
    ]
    widths(ws, [22, 8, 22, 14, 16, 16, 16, 18, 18, 20, 18, 16, 14])
    ws["A1"] = (
        "States → Census division + Charter vs Comcast location lead "
        "(FCC BDC D25). Context only — not which play is live."
    )
    ws["A1"].font = font_title
    ws.merge_cells("A1:M1")
    ws.row_dimensions[1].height = 24
    ws["A2"] = (
        "Source: Analysis/cable-share-small-firm/out/chtr_vs_cmcsa_state_leads.csv. "
        "Availability, not subscribers. RDOF = Auction 904 CCO Holdings (0 if none)."
    )
    ws["A2"].font = font_note
    ws.merge_cells("A2:M2")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.font = font_h
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
    for i, s in enumerate(states):
        r = 5 + i
        vals = [
            s["state"],
            s["abbr"],
            s["division"],
            s["region"],
            s["leader"],
            "Yes" if s["charter_present"] else "No",
            "Yes" if s["comcast_present"] else "No",
            int(round(s["charter_units"])),
            int(round(s["comcast_units"])),
            int(round(s["margin_units"])),
            int(round(s["fabric"])),
            s["rdof"],
            "RETRIEVED",
        ]
        lead_fill = {
            "Charter": fill_chtr,
            "Comcast": fill_cmcsa,
            "Neither present": fill_neither,
        }[s["leader"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = font_b if c == 5 else font
            cell.alignment = center if c != 1 else left
            cell.border = thin
            if c == 5:
                cell.fill = lead_fill
            if c in (8, 9, 10, 11, 12):
                cell.number_format = "#,##0;(#,##0);—"
        if s["rdof"]:
            ws.cell(r, 12).fill = fill_est
    ws.auto_filter.ref = f"A4:M{4 + len(states)}"
    ws.freeze_panes = "A5"
    ws["A57"] = (
        "19 Charter-lead, 27 Comcast-lead, 5 neither present, 0 ties. "
        "RDOF column is ESTIMATED rural context (blue cells > 0)."
    )
    ws["A57"].font = font_note
    ws.merge_cells("A57:M57")


def write_divisions(ws: Worksheet, divisions: list[dict]) -> None:
    ws.sheet_view.showGridLines = False
    widths(ws, [22, 12, 10, 16, 16, 12, 16, 18, 16, 88])
    ws["A1"] = (
        "Division rollup: Charter-lead vs Comcast-lead state counts "
        "(context only — not a play-share)"
    )
    ws["A1"].font = font_title
    ws.merge_cells("A1:J1")
    ws.row_dimensions[1].height = 24
    ws["A2"] = (
        "Comcast leads South Atlantic 7–2 and New England 5–1 on state count. "
        "West South Central is the largest Charter unit-margin (+3.55M). "
        "Eligibility is existing-plant vs rural-reallocation — ESTIMATED."
    )
    ws["A2"].font = font_note
    ws.merge_cells("A2:J2")
    headers = [
        "Census division",
        "Census region",
        "States",
        "Charter-lead",
        "Comcast-lead",
        "Neither",
        "Charter-present",
        "Unit margin (CHTR − CMCSA)",
        "Charter RDOF locs",
        "Strategy eligibility (ESTIMATED context, not a diagnosis)",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.font = font_h
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
        ws.row_dimensions[4].height = 32
    for i, d in enumerate(divisions):
        r = 5 + i
        vals = [
            d["division"],
            d["region"],
            d["n"],
            d["chtr_lead"],
            d["cmcsa_lead"],
            d["neither"],
            d["chtr_present"],
            int(round(d["margin_units"])),
            d["rdof"],
            eligibility(d),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = font_b if c in (4, 5) else font
            cell.alignment = wrap if c == 10 else (center if c != 1 else left)
            cell.border = thin
            if c == 4:
                cell.fill = fill_chtr
            if c == 5:
                cell.fill = fill_cmcsa
            if c in (8, 9):
                cell.number_format = "#,##0;(#,##0);—"
            if c == 9 and d["rdof"]:
                cell.fill = fill_est
        ws.row_dimensions[r].height = 48
    tot_chtr = sum(d["chtr_lead"] for d in divisions)
    tot_cmcsa = sum(d["cmcsa_lead"] for d in divisions)
    tot_n = sum(d["neither"] for d in divisions)
    tot_margin = sum(d["margin_units"] for d in divisions)
    tot_rdof = sum(d["rdof"] for d in divisions)
    tot_present = sum(d["chtr_present"] for d in divisions)
    r = 14
    ws.cell(r, 1, "Total (50 states + DC)").font = font_b
    for c, v in enumerate(
        ["All", 51, tot_chtr, tot_cmcsa, tot_n, tot_present, int(round(tot_margin)), tot_rdof, ""],
        2,
    ):
        cell = ws.cell(r, c, v)
        cell.font = font_b
        cell.border = thin
        if c in (8, 9):
            cell.number_format = "#,##0;(#,##0);—"
    ws["A16"] = (
        "Do not read a division as 'the FWA play' or 'the fiber play'. "
        "No regional play-share exists in the files. Charter-lead / "
        "Charter-present = existing-plant strategies are even on the table. "
        "RDOF mass = rural reallocation is more plausible. Both are context."
    )
    ws["A16"].font = font_note
    ws.merge_cells("A16:J16")
    ws.row_dimensions[16].height = 36
    ws.freeze_panes = "A5"


def main() -> None:
    states = load_states()
    divisions = rollup(states)
    assert len(states) == 51, len(states)
    assert sum(d["chtr_lead"] for d in divisions) == 19
    assert sum(d["cmcsa_lead"] for d in divisions) == 27
    assert sum(d["neither"] for d in divisions) == 5
    assert sum(d["rdof"] for d in divisions) == sum(RDOF.values())

    wb = Workbook()
    notes = wb.active
    notes.title = "Notes"
    write_notes(notes)
    write_matrix(wb.create_sheet("Play x strategy"))
    write_states(wb.create_sheet("States"), states)
    write_divisions(wb.create_sheet("Division rollup"), divisions)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print("Division rollup:")
    for d in divisions:
        print(
            f"  {d['division']}: CHTR-lead {d['chtr_lead']}  "
            f"CMCSA-lead {d['cmcsa_lead']}  neither {d['neither']}  "
            f"present {d['chtr_present']}  margin {d['margin_units']:,.0f}  "
            f"RDOF {d['rdof']:,}"
        )


if __name__ == "__main__":
    main()
