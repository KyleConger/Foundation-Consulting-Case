"""Read-only critic inspector. Does not write the xlsx."""
from __future__ import annotations

import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "CHTR Rec1 Attribution Dashboard.xlsx"
OUT = ROOT / "_critic_inspect.txt"

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
}


def balance(s: str) -> int:
    n = 0
    in_str = False
    for ch in s:
        if ch == '"':
            in_str = not in_str
        elif not in_str:
            if ch == "(":
                n += 1
            elif ch == ")":
                n -= 1
    return n


def main() -> None:
    lines: list[str] = []
    wb = load_workbook(XLSX, data_only=False)
    lines.append(f"FILE {XLSX} size={XLSX.stat().st_size}")
    lines.append(f"SHEETS {wb.sheetnames}")
    lines.append(f"CALC calcMode={wb.calculation.calcMode} fullCalcOnLoad={wb.calculation.fullCalcOnLoad}")
    lines.append("DEFINED NAMES")
    for name in wb.defined_names:
        dn = wb.defined_names[name]
        lines.append(f"  {name} = {dn.attr_text}")

    yellow = "FFFF99"
    green = "E2EFDA"
    for name in wb.sheetnames:
        ws = wb[name]
        lines.append("")
        lines.append(f"===== {name} =====")
        lines.append(
            f"freeze={ws.freeze_panes} print={ws.print_area} "
            f"fitW={ws.page_setup.fitToWidth} fitH={ws.page_setup.fitToHeight} "
            f"orient={ws.page_setup.orientation} paper={ws.page_setup.paperSize}"
        )
        lines.append(
            f"protect={ws.protection.sheet} password={ws.protection.password!r} "
            f"selectLocked={ws.protection.selectLockedCells} "
            f"autoFilter={ws.protection.autoFilter} sort={ws.protection.sort}"
        )
        lines.append(f"merges={len(list(ws.merged_cells.ranges))} charts={len(ws._charts)}")
        hidden = [r for r, dim in ws.row_dimensions.items() if dim.hidden]
        lines.append(f"hidden_rows={hidden}")
        n_in = n_retr = n_in_empty = n_form = n_unbal = 0
        unbal = []
        lits = []
        unlocked = []
        for row in ws.iter_rows():
            for cell in row:
                fill = (cell.fill.fgColor.rgb[-6:] if cell.fill.fgColor and cell.fill.fgColor.rgb else "") if cell.fill.patternType else ""
                if fill.upper() == yellow:
                    n_in += 1
                    if cell.value in (None, ""):
                        n_in_empty += 1
                if fill.upper() == green:
                    n_retr += 1
                if cell.protection and not cell.protection.locked:
                    unlocked.append(cell.coordinate)
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    n_form += 1
                    b = balance(v)
                    if b != 0:
                        n_unbal += 1
                        unbal.append(f"{cell.coordinate} bal={b} {v[:300]}")
                    for tok in ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#NULL!"):
                        if tok in v:
                            lits.append(f"{cell.coordinate} {tok}")
        lines.append(
            f"yellow={n_in} yellow_empty={n_in_empty} green={n_retr} "
            f"formulas={n_form} unbal={n_unbal} unlocked={len(unlocked)}"
        )
        for u in unbal:
            lines.append(f"  UNBAL {u}")
        for lit in lits:
            lines.append(f"  LIT {lit}")
        for i, ch in enumerate(ws._charts):
            a = ch.anchor
            lines.append(
                f"  chart{i} {type(ch).__name__} type={getattr(ch,'type',None)} "
                f"grouping={getattr(ch,'grouping',None)} "
                f"anchor=c{a._from.col}r{a._from.row} "
                f"cx={getattr(a.ext,'cx',None)} cy={getattr(a.ext,'cy',None)}"
            )
            lines.append(
                f"    y_min={ch.y_axis.scaling.min} x_min={ch.x_axis.scaling.min} "
                f"y_fmt={ch.y_axis.numFmt} x_fmt={ch.x_axis.numFmt}"
            )
            try:
                sref = getattr(getattr(ch.title.tx, "strRef", None), "f", None)
                rich = ch.title.tx.rich is not None
            except Exception as e:
                sref, rich = f"err {e}", None
            lines.append(f"    title_strRef={sref} title_rich={rich} nseries={len(ch.series)}")
            for si, ser in enumerate(ch.series):
                lines.append(
                    f"    series{si} title={ser.title} val={ser.val} cat={ser.cat} "
                    f"tx={getattr(ser, 'tx', None)}"
                )

    # Dashboard specific cells
    ws = wb["Dashboard"]
    lines.append("")
    lines.append("===== DASH KEY CELLS =====")
    for addr in [
        "B2", "B3", "B6", "E6", "H6", "K6", "B7", "C11", "C12", "C13", "C14",
        "C15", "C16", "B18", "B29", "K51", "C8",
    ]:
        cell = ws[addr]
        lines.append(f"{addr}\t{cell.value}")

    # Tracker retrieved prints
    tr = wb["Net-loss tracker"]
    lines.append("")
    lines.append("===== TRACKER PRINTS row10 / used row11 =====")
    for col in range(5, 17):
        L = get_column_letter(col)
        lines.append(
            f"{L}5={tr.cell(5,col).value} T5={tr.cell(10,col).value} "
            f"T6f={tr.cell(11,col).value} begin={tr.cell(6,col).value} "
            f"gross={tr.cell(7,col).value} disc={tr.cell(8,col).value} "
            f"cmcsa={tr.cell(14,col).value}"
        )

    # Attribution inputs
    att = wb["Attribution"]
    lines.append("")
    lines.append("===== ATTRIBUTION A1-A6 Q1 2026 (col M = 13) =====")
    for r in range(6, 16):
        lines.append(f"r{r} B={att.cell(r,2).value} M={att.cell(r,13).value}")

    # Parameters
    par = wb["Parameters"]
    lines.append("")
    lines.append("===== PARAMETERS C5-C22 =====")
    for r in range(5, 23):
        lines.append(f"C{r}={par.cell(r,3).value!r} B={par.cell(r,2).value} D={par.cell(r,4).value}")

    # zip internals
    lines.append("")
    lines.append("===== ZIP XML =====")
    with zipfile.ZipFile(XLSX) as z:
        lines.append("files: " + ", ".join(z.namelist()[:80]))
        if "xl/workbook.xml" in z.namelist():
            root = ET.fromstring(z.read("xl/workbook.xml"))
            for dn in root.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}definedName"):
                lines.append(f"  xml name {dn.get('name')} hidden={dn.get('hidden')} {dn.text}")
        for name in z.namelist():
            if name.startswith("xl/charts/"):
                xml = z.read(name).decode("utf-8", "replace")
                lines.append(f"--- {name} len={len(xml)}")
                if "#REF!" in xml:
                    lines.append("  HAS #REF!")
                if "pie" in xml.lower():
                    lines.append("  HAS PIE")
                # title
                if "<c:strRef>" in xml:
                    lines.append("  has strRef title/cache")
                if "<c:rich>" in xml:
                    lines.append("  has rich title")

    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("wrote", OUT, "lines", len(lines))


if __name__ == "__main__":
    main()
