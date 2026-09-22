from pathlib import Path
from openpyxl import load_workbook

p = Path(r"c:\Users\Owner\Desktop\MAN6930 Case\Decisions\CHURN\CHTR Rec1 Attribution Dashboard.xlsx")
wb = load_workbook(p, data_only=False)
ws = wb["Dashboard"]
print("freeze", ws.freeze_panes, "print", ws.print_area, "titles", ws.print_title_rows)
print("protect", ws.protection.sheet, "charts", len(ws._charts))
print("hidden", [r for r, d in ws.row_dimensions.items() if d.hidden])
print("MERGES", ", ".join(str(m) for m in ws.merged_cells.ranges))
for i, ch in enumerate(ws._charts):
    a = ch.anchor
    print(
        f"chart{i} {type(ch).__name__} type={getattr(ch,'type',None)} grouping={getattr(ch,'grouping',None)} "
        f"from_c={a._from.col} from_r={a._from.row} cx={a.ext.cx} cy={a.ext.cy}"
    )
    print("  ymin", ch.y_axis.scaling.min, "xmin", ch.x_axis.scaling.min)
print("--- cells ---")
for row in ws.iter_rows(min_row=1, max_row=62, max_col=14):
    for cell in row:
        if cell.value is None:
            continue
        print(f"{cell.coordinate}\t{cell.value}")
