"""
Build Decisions/CHURN/CHTR Rec1 Attribution Dashboard.xlsx

Diagnostic TEMPLATE for Rec 1 (attribution engine + G1 net-loss tracker).
Yellow INPUT cells are privatized / internal. Everything else is formula or
RETRIEVED public Ex99.1. Do not treat empty reason cells as a measured mix.

Sources: CHTR_Rec1_AttributionPlaybook.pptx · Decisions/six-recommendation-areas.md
· Decisions/decisions.md (G1) · PrimarySources Ex99.1 / trending.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.text import Text
from openpyxl.chart.title import Title
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.legend import Legend
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import Marker

OUT = Path(__file__).resolve().parent / "CHTR Rec1 Attribution Dashboard.xlsx"

GEORGIA = "Georgia"
NAVY = "1F4E79"
ACCENT = "2E75B6"
GRAY = "A6A6A6"
INK = "1A1A1A"

thin = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
thick_bottom = Border(bottom=Side(style="medium", color=NAVY))

font = Font(name=GEORGIA, size=10, color=INK)
font_b = Font(name=GEORGIA, size=10, bold=True, color=INK)
font_title = Font(name=GEORGIA, size=18, bold=True, color=NAVY)
font_h2 = Font(name=GEORGIA, size=13, bold=True, color=NAVY)
font_h = Font(name=GEORGIA, size=10, bold=True, color="FFFFFF")
font_sec = Font(name=GEORGIA, size=11, bold=True, color="FFFFFF")
font_input = Font(name=GEORGIA, size=10, color="0000FF")
font_input_b = Font(name=GEORGIA, size=10, bold=True, color="0000FF")
font_note = Font(name=GEORGIA, size=8, italic=True, color="595959")
font_kpi = Font(name=GEORGIA, size=16, bold=True, color=NAVY)
font_small = Font(name=GEORGIA, size=8, color="595959")

fill_head = PatternFill("solid", fgColor=NAVY)
fill_sec = PatternFill("solid", fgColor=ACCENT)
fill_input = PatternFill("solid", fgColor="FFFF99")  # yellow INPUT
fill_retr = PatternFill("solid", fgColor="E2EFDA")  # retrieved public
fill_calc = PatternFill("solid", fgColor="F2F2F2")
fill_est = PatternFill("solid", fgColor="DDEBF7")  # ESTIMATED template default
fill_white = PatternFill("solid", fgColor="FFFFFF")
fill_navy_lite = PatternFill("solid", fgColor="D6E3F0")
fill_pass = PatternFill("solid", fgColor="C6EFCE")
fill_warn = PatternFill("solid", fgColor="FFF2CC")
fill_fail = PatternFill("solid", fgColor="F8CBAD")
fill_dash = PatternFill("solid", fgColor="F7F9FC")
fill_card = PatternFill("solid", fgColor="FFFFFF")

wrap = Alignment(wrap_text=True, vertical="center")
left = Alignment(wrap_text=True, vertical="center", horizontal="left")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")

NUM = "#,##0;(#,##0);—"
NUM_K = "#,##0"
PCT = "0.0%"
PCT1 = "0.00%"
MULT = "0.00×"

QUARTERS = [
    "Q1 2024",
    "Q2 2024",
    "Q3 2024",
    "Q4 2024",
    "Q1 2025",
    "Q2 2025",
    "Q3 2025",
    "Q4 2025",
    "Q1 2026",
    "Q2 2026",
    "Q3 2026",
    "Q4 2026",
]
Q0 = 5  # column E
N_Q = len(QUARTERS)

# RETRIEVED · Ex99.1 / trending (customers, not 000s)
RES_NET = {
    "Q1 2024": -72000,
    "Q2 2024": -154000,
    "Q3 2024": -113000,
    "Q4 2024": -171000,
    "Q1 2025": -55000,
    "Q2 2025": -111000,
    "Q3 2025": -108000,
    "Q4 2025": -119000,
    "Q1 2026": -117000,
    "Q2 2026": -166000,
}
RES_EOP = {
    "Q1 2024": 28472000,
    "Q2 2024": 28318000,
    "Q3 2024": 28205000,
    "Q4 2024": 28034000,
    "Q1 2025": 27979000,
    "Q2 2025": 27868000,
    "Q3 2025": 27760000,
    "Q4 2025": 27641000,
    "Q1 2026": 27524000,
}
RES_BEGIN = {
    "Q1 2024": 28544000,  # EOP + |net| identity from trending
    "Q2 2024": 28472000,
    "Q3 2024": 28318000,
    "Q4 2024": 28205000,
    "Q1 2025": 28034000,
    "Q2 2025": 27979000,
    "Q3 2025": 27868000,
    "Q4 2025": 27760000,
    "Q1 2026": 27641000,
    "Q2 2026": 27524000,
}
CO_INET = {
    "Q1 2024": -72000,
    "Q2 2024": -148000,
    "Q3 2024": -110000,
    "Q4 2024": -177000,
    "Q1 2025": -59000,
    "Q2 2025": -116000,
    "Q3 2025": -109000,
    "Q4 2025": -119000,
    "Q1 2026": -120000,
    "Q2 2026": -172000,
}
RURAL_CR = {"Q1 2026": 41000, "Q2 2026": 47000}
CMCSA_NET = {"Q1 2026": -65000}
CMCSA_BEGIN = {"Q1 2026": 28719000}  # 28,654k EOP − (−65k)

REASONS = [
    "Competitive FWA",
    "Competitive fiber (overbuild)",
    "LEO satellite",
    "Move-out / housing",
    "Non-pay / economic",
    "Service / other",
]
EXPOSURE = [
    "Fiber-overbuilt",
    "FWA-covered",
    "Satellite-only",
    "Uncontested (control)",
]
WHERE = ["Core", "Rural", "AT&T overlap", "Verizon overlap", "Rest of footprint"]
TENURE = ["0–12 months", "12–36 months", "36+ months"]
PRODUCT = ["Internet-only", "Internet + mobile", "Internet + video", "Triple-play"]
ARPU = ["Low ARPU", "Mid ARPU", "High ARPU"]
EXIT = ["Price", "Speed / reliability / upload", "Move-out", "Non-pay", "Other / unknown"]


def qcol(i: int) -> int:
    return Q0 + i


def qletter(i: int) -> str:
    return get_column_letter(qcol(i))


def last_q() -> str:
    return get_column_letter(qcol(N_Q - 1))


_FORMULA_PREFIXES = (
    "IF(", "IFERROR(", "INDEX(", "OR(", "AND(", "ABS(", "SUM(", "COUNT",
    "LEFT(", "TEXT(", "LARGE(", "MAX(", "MIN(", "RANK(", "N(", "NA(",
    "CONCAT", "MATCH(", "PARAMETERS!",
)

def put(ws, r, c, value, *, kind="calc", fmt=None, bold=False, align=None, lock=None):
    if (
        kind in ("calc", "kpi", "h2")
        and isinstance(value, str)
        and value
        and not value.startswith("=")
        and (value.upper().startswith(_FORMULA_PREFIXES) or "!" in value)
    ):
        value = "=" + value
    cell = ws.cell(r, c, value)
    if kind == "input":
        cell.font = font_input_b if bold else font_input
        cell.fill = fill_input
        cell.protection = Protection(locked=False)
    elif kind == "retr":
        cell.font = font_b if bold else font
        cell.fill = fill_retr
        cell.protection = Protection(locked=True)
    elif kind == "est":
        cell.font = font_input if not bold else font_input_b
        cell.fill = fill_est
        cell.protection = Protection(locked=False)
    elif kind == "head":
        cell.font = font_h
        cell.fill = fill_head
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.protection = Protection(locked=True)
    elif kind == "sec":
        cell.font = font_sec
        cell.fill = fill_sec
        cell.alignment = left
        cell.protection = Protection(locked=True)
    elif kind == "label":
        cell.font = font_b if bold else font
        cell.fill = fill_white
        cell.alignment = left
        cell.protection = Protection(locked=True)
    elif kind == "note":
        cell.font = font_note
        cell.fill = fill_white
        cell.alignment = left
        cell.protection = Protection(locked=True)
    elif kind == "kpi":
        cell.font = font_kpi
        cell.fill = fill_card
        cell.protection = Protection(locked=True)
    elif kind == "h2":
        cell.font = font_h2
        cell.fill = fill_white
        cell.alignment = left
        cell.protection = Protection(locked=True)
    elif kind == "title":
        cell.font = font_title
        cell.fill = fill_white
        cell.protection = Protection(locked=True)
    else:
        cell.font = font_b if bold else font
        cell.fill = fill_calc
        cell.protection = Protection(locked=True)
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    elif kind not in ("head", "sec", "note", "label", "h2", "title"):
        cell.alignment = Alignment(vertical="center")
    if lock is False:
        cell.protection = Protection(locked=False)
    cell.border = thin
    return cell


def merge(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def page(ws, *, landscape=True, freeze="A6", print_area=None, tab="1F4E79"):
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID if landscape else ws.PAPERSIZE_LETTER
    ws.page_setup.horizontalCentered = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.sheet_view.showGridLines = False
    if freeze:
        ws.freeze_panes = freeze
    if print_area:
        ws.print_area = print_area
    ws.sheet_properties.tabColor = tab
    ws.page_setup.horizontalDpi = 300
    ws.oddHeader.left.text = "CHTR Rec 1 · Attribution diagnostic TEMPLATE"
    ws.oddFooter.left.text = "Yellow = INPUT · Green = RETRIEVED Ex99.1 · Gray = formula · Not a filled mix"
    ws.oddFooter.right.text = "Page &P of &N"


def link_chart_title(chart, cell_ref: str):
    title = Title()
    text = Text()
    text.strRef = StrRef(cell_ref)
    title.tx = text
    chart.title = title


def style_series(series, hex_color: str, *, no_line=False):
    series.graphicalProperties.solidFill = hex_color
    if no_line:
        series.graphicalProperties.line.noFill = True
    else:
        series.graphicalProperties.line.solidFill = hex_color


def protect(ws):
    ws.protection.sheet = True
    ws.protection.enable()
    ws.protection.autoFilter = True
    ws.protection.sort = True
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def idx_q(sheet: str, row: int) -> str:
    """INDEX/MATCH the selected quarter from a row on a sheet with E5:P5 headers."""
    return (
        f'IFERROR(INDEX(\'{sheet}\'!$E${row}:$P${row},'
        f'MATCH(Param_Quarter,\'{sheet}\'!$E$5:$P$5,0)),"")'
    )


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def build() -> Path:
    wb = Workbook()

    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_cover = wb.create_sheet("Cover")
    ws_trk = wb.create_sheet("Net-loss tracker")
    ws_att = wb.create_sheet("Attribution")
    ws_seg = wb.create_sheet("Segments")
    ws_par = wb.create_sheet("Parameters")
    ws_asm = wb.create_sheet("Assumptions")
    ws_pb = wb.create_sheet("Playbook")

    build_parameters(ws_par)
    build_tracker(ws_trk)
    build_attribution(ws_att)
    build_segments(ws_seg)
    build_assumptions(ws_asm)
    build_playbook(ws_pb)
    build_cover(ws_cover)
    build_dashboard(ws_dash, wb)

    # Defined names (after sheets exist so refs resolve)
    names = {
        "Param_Quarter": "Parameters!$C$5",
        "Param_Dominant": "Parameters!$C$8",
        "Param_Mixed": "Parameters!$C$9",
        "Param_FWA": "Parameters!$C$10",
        "Param_Fiber": "Parameters!$C$11",
        "Param_Satellite": "Parameters!$C$12",
        "Param_Coequal": "Parameters!$C$13",
        "Param_CoreShare": "Parameters!$C$14",
        "Param_RateSkew": "Parameters!$C$15",
        "Param_TwoQ": "Parameters!$C$16",
        "Param_TrackerExists": "Parameters!$C$17",
        "Param_CMCSA_On": "Parameters!$C$18",
        "Param_InputMode": "Parameters!$C$19",
        "Param_CommitExcess": "Parameters!$C$22",
        "Dash_Print": "Dashboard!$C$8",
    }
    for name, ref in names.items():
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    # Print order: Dashboard first
    wb._sheets = [ws_dash, ws_cover, ws_trk, ws_att, ws_seg, ws_par, ws_asm, ws_pb]

    for ws in wb.worksheets:
        protect(ws)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


# ---------------------------------------------------------------------------
# Parameters
# ---------------------------------------------------------------------------

def build_parameters(ws):
    page(ws, landscape=False, freeze="A5", print_area="A1:F40", tab="ED7D31")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 62
    ws.column_dimensions["F"].width = 18

    put(ws, 1, 1, "PARAMETERS", kind="title")
    merge(ws, 1, 1, 1, 5)
    put(
        ws,
        2,
        1,
        "Yellow = INPUT. Blue fill = ESTIMATED template default the user can change. "
        "PPTX slide 5: ≥40% dominant · 25–40% mixed · <25% null. FWA signal >35–40%. "
        "Thresholds are judgment calls, set in advance (pptx notes: arbitrary-but-stated beats unstated).",
        kind="note",
    )
    merge(ws, 2, 1, 3, 5)
    ws.row_dimensions[2].height = 28

    put(ws, 4, 1, "ID", kind="head")
    put(ws, 4, 2, "Parameter", kind="head")
    put(ws, 4, 3, "Value", kind="head")
    put(ws, 4, 4, "Tag", kind="head")
    put(ws, 4, 5, "Source / why it is here", kind="head")

    rows = [
        (5, "P1", "Selected quarter", "Q1 2026", "INPUT",
         "Assignment anchor is Q1 2026. Q2 is a later public fact (pptx + six-recommendation-areas)."),
        (6, "P2", "Reason taxonomy", "PPTX 6-scenario", "INPUT",
         "PPTX scenarios 1–6. Alternate label only — codes stay R1–R4 + satellite + residual."),
        (7, "P3", "Highlight color (tables)", "2E75B6", "INPUT",
         "One accent. Charts use this hex for the story series; peers stay gray."),
        (8, "P4", "Dominant-driver threshold", 0.40, "ESTIMATED",
         "PPTX slide 5: one cause-cohort ≥40% of quarterly losses → fire that lead play at scale."),
        (9, "P5", "Mixed-picture floor", 0.25, "ESTIMATED",
         "PPTX slide 5: 25–40% mixed (top two as experiments). <25% → Scenario 6 null."),
        (10, "P6", "FWA scenario trigger", 0.35, "ESTIMATED",
         "PPTX S1: FWA-covered losses >35–40%; skew low-tenure / low-ARPU; exit cites price."),
        (11, "P7", "Fiber scenario trigger", 0.40, "ESTIMATED",
         "PPTX S2: fiber-DMA losses dominate; high-ARPU; exit cites speed / reliability / upload."),
        (12, "P8", "Satellite scenario trigger", 0.25, "ESTIMATED",
         "PPTX S3: losses map to satellite-only low-density blocks. No numeric cut in the deck — template default."),
        (13, "P9", "Co-equal band (pp of mix)", 0.05, "ESTIMATED",
         "GAP: co-equal cells fire the cheaper lever first. Two expensive levers do not fire in the same quarter."),
        (14, "P10", "Core-drives-print share", 0.70, "ESTIMATED",
         "WHERE overlay (decisions.md W1): if core is the Internet leak, fire the reason play on the core — do not add rural miles."),
        (15, "P11", "Segment rate-skew vs company", 1.50, "ESTIMATED",
         "Flag a segment whose net-loss rate is ≥ this multiple of the company residential rate."),
        (16, "P12", "Require two consecutive quarters before full-scale?", "Yes", "ESTIMATED",
         "PPTX slide 11 risk interlock: thresholds require two consecutive quarters before full-scale (experiments can start on one)."),
        (17, "P13", "Tracker already exists and is used quarterly?", "No — fill G1 first", "INPUT",
         "Kill rule 1 (six-recommendation-areas / decisions.md): if yes, skip Area 1 build and go to mix control."),
        (18, "P14", "Use Comcast benchmark?", "Yes", "INPUT",
         "Apples-to-apples is residential broadband (CHTR res. Internet vs CMCSA domestic resid. BB)."),
        (19, "P15", "Attribution input mode", "Counts", "INPUT",
         "Counts (preferred) = sampled disconnects. Percent allocates the selected base — not a measured rate."),
        (20, "P16", "If only % is filled, allocate against", "Disconnects (if known), else |public print|", "ESTIMATED",
         "Playbook: do not invent a mix. Allocating the net print is a PLACEHOLDER, not a churn rate."),
        (21, "P17", "Material share of peer excess", 0.25, "ESTIMATED",
         "GAP: a branch fires only if it is largest/co-equal AND closing the Comcast gap in that cell would absorb a material share of the excess."),
        (22, "P18", "Playbook committed excess (customers)", 54000, "RETRIEVED",
         "six-recommendation-areas.md / network_churn_model.py: 54k (round down from 54.4k) · 1.8× Comcast. Q1 identity only. Not in the pptx."),
    ]

    kind_map = {"INPUT": "input", "ESTIMATED": "est", "RETRIEVED": "retr"}
    fmt_map = {
        8: PCT, 9: PCT, 10: PCT, 11: PCT, 12: PCT, 13: PCT,
        14: PCT, 15: "0.00", 21: PCT, 22: NUM,
    }
    for r, lid, label, val, tag, src in rows:
        put(ws, r, 1, lid, kind="label", bold=True)
        put(ws, r, 2, label, kind="label")
        put(ws, r, 3, val, kind=kind_map[tag], fmt=fmt_map.get(r), bold=True, align=center)
        put(ws, r, 4, tag, kind="label")
        put(ws, r, 5, src, kind="note")
        ws.row_dimensions[r].height = 32

    # Reason names (editable taxonomy)
    put(ws, 24, 1, "REASON TAXONOMY (edit labels — codes stay fixed)", kind="sec")
    merge(ws, 24, 1, 24, 5)
    put(ws, 25, 1, "Code", kind="head")
    put(ws, 25, 2, "Dashboard label (INPUT)", kind="head")
    put(ws, 25, 3, "PPTX scenario", kind="head")
    put(ws, 25, 4, "GAP branch", kind="head")
    put(ws, 25, 5, "Lead play (do not fire from empty cells)", kind="head")
    play_rows = [
        ("R-FWA", "Competitive FWA", "S1 FWA-driven", "R2",
         "Convergence pricing / FWA-match in FWA zones; economics gate on MVNO unit cost."),
        ("R-FIB", "Competitive fiber (overbuild)", "S2 Fiber-driven", "R3",
         "Re-sequence DOCSIS 4.0 by exposure + symmetric win-back. Node-split only if sample names speed."),
        ("R-SAT", "LEO satellite", "S3 Satellite-driven", "W2 overlay",
         "Accelerate rural subsidized activation ahead of Starlink/Kuiper. Honor RDOF/BEAD."),
        ("R-MOV", "Move-out / housing", "S4 Housing-driven", "R1",
         "Mover capture: Community Solutions / MDU / instant-on. Not a save-offer problem."),
        ("R-NPY", "Non-pay / economic", "S5 Non-pay", "R4",
         "Payment arrangement / value tier at delinquency. Funded by care opex if this is the leak."),
        ("R-OTH", "Service / other", "Residual", "—",
         "Unmapped residual. If this is largest, the sample is incomplete — do not fire a plant program."),
    ]
    for i, (code, name, sc, gap, play) in enumerate(play_rows):
        r = 26 + i
        put(ws, r, 1, code, kind="label", bold=True)
        put(ws, r, 2, name, kind="input")
        put(ws, r, 3, sc, kind="label")
        put(ws, r, 4, gap, kind="label")
        put(ws, r, 5, play, kind="note")

    put(ws, 33, 1, "Dropdown lists (do not edit structure)", kind="sec")
    merge(ws, 33, 1, 33, 3)
    for i, q in enumerate(QUARTERS):
        put(ws, 34 + i, 1, q, kind="label")
    put(ws, 34, 2, "Yes", kind="label")
    put(ws, 35, 2, "No", kind="label")
    put(ws, 34, 3, "Yes — skip Area 1 build", kind="label")
    put(ws, 35, 3, "No — fill G1 first", kind="label")
    put(ws, 34, 4, "Counts", kind="label")
    put(ws, 35, 4, "Percent", kind="label")
    put(ws, 34, 5, "PPTX 6-scenario", kind="label")
    put(ws, 35, 5, "GAP R1–R4 only", kind="label")

    dv_q = DataValidation(type="list", formula1="$A$34:$A$45", allow_blank=False)
    dv_q.add("C5")
    ws.add_data_validation(dv_q)
    dv_tr = DataValidation(type="list", formula1="$C$34:$C$35", allow_blank=False)
    dv_tr.add("C17")
    ws.add_data_validation(dv_tr)
    dv_yn = DataValidation(type="list", formula1="$B$34:$B$35", allow_blank=False)
    dv_yn.add("C16")
    dv_yn.add("C18")
    ws.add_data_validation(dv_yn)
    dv_mode = DataValidation(type="list", formula1="$D$34:$D$35", allow_blank=False)
    dv_mode.add("C19")
    ws.add_data_validation(dv_mode)
    dv_tax = DataValidation(type="list", formula1="$E$34:$E$35", allow_blank=False)
    dv_tax.add("C6")
    ws.add_data_validation(dv_tax)

    put(ws, 47, 1,
        "Comcast Q1 2026 residential broadband −65,000 and begin 28,719,000 are RETRIEVED "
        "(CMCSA Ex99.1; network_churn_model.py). Other Comcast quarters stay INPUT. "
        "54,000 is the playbook commit (round down from 54.4k) — cite six-recommendation-areas.md, not the pptx.",
        kind="note")
    merge(ws, 47, 1, 48, 5)

    for r in range(1, 50):
        if ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = 16
    ws.row_dimensions[1].height = 24


# ---------------------------------------------------------------------------
# Net-loss tracker
# ---------------------------------------------------------------------------

def build_tracker(ws):
    page(ws, freeze="E6", print_area="A1:P32", tab="FFFF00")
    ws.auto_filter.ref = "A5:P20"
    widths = {"A": 8, "B": 46, "C": 14, "D": 38}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for i in range(N_Q):
        ws.column_dimensions[qletter(i)].width = 13

    put(ws, 1, 1, "G1 · QUARTERLY NET-LOSS TRACKER  —  fill this first", kind="title")
    merge(ws, 1, 1, 1, 8)
    put(
        ws,
        2,
        1,
        "six-recommendation-areas.md: “Fill the quarterly net-loss tracker first.” "
        "Unit = residential Internet. Prefer gross adds vs disconnects so net is a formula. "
        "Yellow = Charter internals (INPUT). Green = public Ex99.1 / trending (RETRIEVED). "
        "Do not treat rural CR +41k / +47k as this Internet cut. Do not fire a reason play from empty cells.",
        kind="note",
    )
    merge(ws, 2, 1, 3, 12)
    ws.row_dimensions[2].height = 32

    put(ws, 5, 1, "ID", kind="head")
    put(ws, 5, 2, "Line", kind="head")
    put(ws, 5, 3, "Tag", kind="head")
    put(ws, 5, 4, "Source / formula", kind="head")
    for i, q in enumerate(QUARTERS):
        put(ws, 5, qcol(i), q, kind="head")

    # Row map (locked for formulas on other sheets)
    # 6 begin INPUT/retr, 7 gross INPUT, 8 disc INPUT, 9 internal net formula,
    # 10 public print RETR, 11 print used, 12 company inet, 13 rural CR,
    # 14 CMCSA net, 15 CMCSA begin, 16 CMCSA rate, 17 CHTR rate, 18 ratio,
    # 19 expected, 20 excess, 21 eop identity, 22 reconcile flag
    lines = [
        (6, "T1", "Beginning residential Internet PSUs", "mixed",
         "Prior-quarter EOP. Prefill = trending identity (RETRIEVED). Overwrite if internal census differs."),
        (7, "T2", "Gross adds (internal)", "INPUT",
         "Not disclosed. Leave blank until the tracker exists. Sampled, not a guess."),
        (8, "T3", "Disconnects (internal)", "INPUT",
         "Not disclosed. Sample the disconnect — do not paste agent codes as the file (pptx slide 4)."),
        (9, "T4", "Internal net adds  =  T2 − T3", "DERIVED",
         "Blank until both T2 and T3 are filled. Zero is not a fill."),
        (10, "T5", "Public print — res. Internet net adds (Ex99.1)", "RETRIEVED",
         "CHTR Ex99.1 / Q1 trending. Q1 2026 = −117,000. Q2 2026 = −166,000. FY columns omitted."),
        (11, "T6", "Print used for diagnosis", "DERIVED",
         "Internal net if T2 and T3 filled; else public print. Never invent a mix on top of this."),
        (12, "T7", "Company Internet net adds (context, not this cut)", "RETRIEVED",
         "Ex99.1 total Internet. Q1 2026 −120,000 · Q2 2026 −172,000. Includes SMB."),
        (13, "T8", "Rural CR net adds (NOT an Internet cut)", "RETRIEVED",
         "Ex99.1 rural-footprint customer relationships. Q1 +41,000 · Q2 +47,000. Do not use as T6."),
        (14, "T9", "Comcast domestic resid. BB net adds", "mixed",
         "Q1 2026 −65,000 RETRIEVED (CMCSA Ex99.1). Other quarters INPUT. Leave blank if unused."),
        (15, "T10", "Comcast resid. BB beginning PSUs", "mixed",
         "Q1 2026 28,719,000 = 28,654k EOP − (−65k). Other quarters INPUT."),
        (16, "T11", "Comcast net-loss rate  =  |T9| / T10", "DERIVED",
         "Only when T9 and T10 are filled."),
        (17, "T12", "Charter res. net-loss rate  =  |T6| / T1", "DERIVED",
         "Net-loss rate, not gross churn. Gross churn is unused because it is not disclosed."),
        (18, "T13", "Rate ratio vs Comcast  =  T12 / T11", "DERIVED",
         "Q1 playbook commit 1.8× (round down from 1.87×)."),
        (19, "T14", "Expected CHTR losses at CMCSA rate  =  T11 × T1", "DERIVED",
         "Build 2 in ifxy-stock-signal.md / network_churn_model.py."),
        (20, "T15", "Excess vs Comcast  =  |T6| − T14", "DERIVED",
         "Q1 identity ≈ 54.4k; playbook commits 54k (Parameters P18). Cite six-recommendation-areas.md."),
        (21, "T16", "Ending PSUs  =  T1 + T6", "DERIVED",
         "Identity check vs trending EOP when public print is used."),
        (22, "T17", "Internal vs public reconcile", "DERIVED",
         "OK if internals blank or match public print; FLAG if they differ."),
    ]

    for r, lid, label, tag, src in lines:
        put(ws, r, 1, lid, kind="label", bold=True)
        put(ws, r, 2, label, kind="label")
        put(ws, r, 3, tag if tag != "mixed" else "see cells", kind="label")
        put(ws, r, 4, src, kind="note")
        ws.row_dimensions[r].height = 28

    for i, q in enumerate(QUARTERS):
        c = qcol(i)
        L = qletter(i)

        # T1 begin — retrieved where known, else INPUT blank
        if q in RES_BEGIN:
            put(ws, 6, c, RES_BEGIN[q], kind="retr", fmt=NUM, bold=True)
        else:
            put(ws, 6, c, None, kind="input", fmt=NUM)

        put(ws, 7, c, None, kind="input", fmt=NUM)  # gross
        put(ws, 8, c, None, kind="input", fmt=NUM)  # disc

        put(ws, 9, c,
            f'IF(OR({L}7="",{L}8=""),"",{L}7-{L}8)',
            kind="calc", fmt=NUM, bold=True)

        if q in RES_NET:
            put(ws, 10, c, RES_NET[q], kind="retr", fmt=NUM, bold=True)
        else:
            put(ws, 10, c, None, kind="input", fmt=NUM)

        put(ws, 11, c,
            f'IF({L}9<>"",{L}9,IF({L}10<>"",{L}10,""))',
            kind="calc", fmt=NUM, bold=True)

        if q in CO_INET:
            put(ws, 12, c, CO_INET[q], kind="retr", fmt=NUM)
        else:
            put(ws, 12, c, None, kind="input", fmt=NUM)

        if q in RURAL_CR:
            put(ws, 13, c, RURAL_CR[q], kind="retr", fmt=NUM)
        else:
            put(ws, 13, c, None, kind="input", fmt=NUM)

        if q in CMCSA_NET:
            put(ws, 14, c, CMCSA_NET[q], kind="retr", fmt=NUM, bold=True)
        else:
            put(ws, 14, c, None, kind="input", fmt=NUM)

        if q in CMCSA_BEGIN:
            put(ws, 15, c, CMCSA_BEGIN[q], kind="retr", fmt=NUM)
        else:
            put(ws, 15, c, None, kind="input", fmt=NUM)

        put(ws, 16, c,
            f'IF(OR({L}14="",{L}15="",{L}15=0),"",ABS({L}14)/{L}15)',
            kind="calc", fmt=PCT1)
        put(ws, 17, c,
            f'IF(OR({L}11="",{L}6="",{L}6=0),"",ABS({L}11)/{L}6)',
            kind="calc", fmt=PCT1)
        put(ws, 18, c,
            f'IF(OR({L}16="",{L}17="",{L}16=0),"",{L}17/{L}16)',
            kind="calc", fmt=MULT)
        put(ws, 19, c,
            f'IF(OR({L}16="",{L}6=""),"",{L}16*{L}6)',
            kind="calc", fmt=NUM)
        put(ws, 20, c,
            f'IF(OR({L}11="",{L}19=""),"",ABS({L}11)-{L}19)',
            kind="calc", fmt=NUM, bold=True)
        put(ws, 21, c,
            f'IF(OR({L}6="",{L}11=""),"",{L}6+{L}11)',
            kind="calc", fmt=NUM)
        put(ws, 22, c,
            f'IF({L}9="",IF({L}10="","PLACEHOLDER — no print","Public print only — internals empty"),'
            f'IF({L}10="","Internal only",'
            f'IF(ABS({L}9-{L}10)<500,"OK — internal matches public","FLAG — internal ≠ public print")))',
            kind="calc")

    put(ws, 24, 1, "HOW TO READ THIS SHEET", kind="sec")
    merge(ws, 24, 1, 24, 8)
    notes = [
        "1. Paste starting PSUs / gross adds / disconnects in yellow. Net (T4) and print used (T6) are formulas.",
        "2. Public −117,000 (Q1 2026) and −166,000 (Q2 2026) are RETRIEVED Ex99.1 residential Internet — not a filled attribution.",
        "3. Excess vs Comcast (T15) is the 54k figure’s identity for Q1 only (CMCSA rate × CHTR begin − |print|). Other quarters compute only when Comcast cells are filled.",
        "4. Rural CR is on this sheet so nobody confuses +41k / +47k with residential Internet. It is not T6.",
        "5. Kill: if Parameters says the tracker already exists, skip building G1 — paste the last eight quarters and go to Attribution.",
    ]
    for i, t in enumerate(notes):
        put(ws, 25 + i, 1, t, kind="note")
        merge(ws, 25 + i, 1, 25 + i, 12)

    put(ws, 31, 1, "Selected-quarter pull (engine — do not edit)", kind="sec")
    merge(ws, 31, 1, 31, 4)
    labels = [
        (32, "Begin PSUs"),
        (33, "Gross adds"),
        (34, "Disconnects"),
        (35, "Internal net"),
        (36, "Public print"),
        (37, "Print used"),
        (38, "Company Internet"),
        (39, "Rural CR"),
        (40, "CMCSA net"),
        (41, "CMCSA begin"),
        (42, "CMCSA rate"),
        (43, "CHTR rate"),
        (44, "Rate ratio"),
        (45, "Expected losses"),
        (46, "Excess vs CMCSA"),
        (47, "Ending PSUs"),
        (48, "Reconcile"),
    ]
    for r, lab in labels:
        put(ws, r, 1, lab, kind="label")
        src_row = r - 26  # 32->6
        put(ws, r, 2, f"={idx_q('Net-loss tracker', src_row)}",
            kind="calc", fmt=NUM if r not in (42, 43, 44, 48) else (MULT if r == 44 else (PCT1 if r in (42, 43) else None)))

    # hide engine rows from print by keeping them below print_area; leave visible for audit
    put(ws, 50, 1,
        "Engine rows 32–48 feed the Dashboard. Q3–Q4 2026 begin/print stay blank until Ex99.1 or internals exist.",
        kind="note")
    merge(ws, 50, 1, 50, 8)

    ws.row_dimensions[5].height = 22
    ws.auto_filter.ref = "A5:P5"


# ---------------------------------------------------------------------------
# Attribution
# ---------------------------------------------------------------------------

def build_attribution(ws):
    page(ws, freeze="E6", print_area="A1:P36", tab="FFFF00")
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 28
    for i in range(N_Q):
        ws.column_dimensions[qletter(i)].width = 13

    put(ws, 1, 1, "ATTRIBUTION · REASON PLAY  —  sampled disconnects, not agent codes", kind="title")
    merge(ws, 1, 1, 1, 8)
    put(
        ws,
        2,
        1,
        "PPTX cause-cohorts: FWA-covered · fiber-overbuilt · satellite-only · housing/move-out · non-pay · residual. "
        "GAP R1–R4 maps onto the same rows (satellite is pptx-only). "
        "Enter COUNTS of sampled disconnects (yellow). Residual is a formula. "
        "Empty cells are on purpose — the playbook will not fire from a guess.",
        kind="note",
    )
    merge(ws, 2, 1, 3, 12)
    ws.row_dimensions[2].height = 32

    put(ws, 5, 1, "ID", kind="head")
    put(ws, 5, 2, "Cause-cohort (label from Parameters)", kind="head")
    put(ws, 5, 3, "Tag", kind="head")
    put(ws, 5, 4, "PPTX / GAP", kind="head")
    for i, q in enumerate(QUARTERS):
        put(ws, 5, qcol(i), q, kind="head")

    # Rows 6-11 reason INPUT counts; 12 residual; 13 sum; 14 tracker disconnects; 15 mix check
    gap_tag = ["S1 / R2", "S2 / R3", "S3 / W2", "S4 / R1", "S5 / R4", "Residual"]
    for i, (name, tag) in enumerate(zip(REASONS, gap_tag)):
        r = 6 + i
        put(ws, r, 1, f"A{i+1}", kind="label", bold=True)
        put(ws, r, 2, f"=Parameters!B{26+i}", kind="calc")
        put(ws, r, 3, "INPUT", kind="label")
        put(ws, r, 4, tag, kind="note")
        for j in range(N_Q):
            put(ws, r, qcol(j), None, kind="input", fmt=NUM)

    put(ws, 12, 1, "A7", kind="label", bold=True)
    put(ws, 12, 2, "Residual (tracker disconnects − sum of A1–A6)", kind="label")
    put(ws, 12, 3, "DERIVED", kind="label")
    put(ws, 12, 4, "Largest residual is a data-quality flag", kind="note")
    for j in range(N_Q):
        L = qletter(j)
        put(ws, 12, qcol(j),
            f'IF({L}14="","",{L}14-SUM({L}6:{L}11))',
            kind="calc", fmt=NUM, bold=True)

    put(ws, 13, 1, "A8", kind="label", bold=True)
    put(ws, 13, 2, "Sum of entered reason counts (A1–A6)", kind="label")
    put(ws, 13, 3, "DERIVED", kind="label")
    put(ws, 13, 4, "Does not include residual", kind="note")
    for j in range(N_Q):
        L = qletter(j)
        put(ws, 13, qcol(j), f'IF(COUNT({L}6:{L}11)=0,"",SUM({L}6:{L}11))', kind="calc", fmt=NUM)

    put(ws, 14, 1, "A9", kind="label", bold=True)
    put(ws, 14, 2, "Tracker disconnects (T3) — or |print| placeholder", kind="label")
    put(ws, 14, 3, "DERIVED", kind="label")
    put(ws, 14, 4, "If T3 blank and mode is Percent, uses |T6| and labels PLACEHOLDER", kind="note")
    for j in range(N_Q):
        L = qletter(j)
        # Pull T3 and T6 from tracker same column
        put(ws, 14, qcol(j),
            f'IF(\'Net-loss tracker\'!{L}8<>"",\'Net-loss tracker\'!{L}8,'
            f'IF(AND(Parameters!$C$19="Percent",\'Net-loss tracker\'!{L}11<>""),'
            f'ABS(\'Net-loss tracker\'!{L}11),""))',
            kind="calc", fmt=NUM)

    put(ws, 15, 1, "A10", kind="label", bold=True)
    put(ws, 15, 2, "Mix status", kind="label")
    put(ws, 15, 3, "DERIVED", kind="label")
    put(ws, 15, 4, "Empty until a reason or residual exists", kind="note")
    for j in range(N_Q):
        L = qletter(j)
        put(ws, 15, qcol(j),
            f'IF(AND(COUNT({L}6:{L}11)=0,{L}14=""),"EMPTY — do not fire a play",'
            f'IF(\'Net-loss tracker\'!{L}8="",'
            f'"PLACEHOLDER — allocating |print|, not a churn mix",'
            f'IF(ABS(N({L}12))>0.02*MAX(1,N({L}14)),"FLAG — residual is material","OK — reasons reconcile")))',
            kind="calc")

    # % of disconnects
    put(ws, 17, 1, "SHARE OF DISCONNECTS (formula)", kind="sec")
    merge(ws, 17, 1, 17, 4)
    put(ws, 18, 1, "ID", kind="head")
    put(ws, 18, 2, "Cause-cohort", kind="head")
    put(ws, 18, 3, "Tag", kind="head")
    put(ws, 18, 4, "Share formula", kind="head")
    for i, q in enumerate(QUARTERS):
        put(ws, 18, qcol(i), q, kind="head")

    for i in range(7):  # 6 reasons + residual
        r = 19 + i
        src = 6 + i
        put(ws, r, 1, f"S{i+1}", kind="label", bold=True)
        put(ws, r, 2, f"=B{src}", kind="calc")
        put(ws, r, 3, "DERIVED", kind="label")
        put(ws, r, 4, f"A{i+1} / (A8+A7) or A7/base", kind="note")
        for j in range(N_Q):
            L = qletter(j)
            put(ws, r, qcol(j),
                f'IF(OR({L}14="",{L}14=0),"",N({L}{src})/{L}14)',
                kind="calc", fmt=PCT)

    # Contribution to net loss (negative numbers = losses)
    put(ws, 27, 1, "CONTRIBUTION TO THE NET PRINT (formula)", kind="sec")
    merge(ws, 27, 1, 27, 4)
    put(ws, 28, 1, "ID", kind="head")
    put(ws, 28, 2, "Cause-cohort", kind="head")
    put(ws, 28, 3, "Tag", kind="head")
    put(ws, 28, 4, "If gross adds exist, housing also shows failed-acquisition residual", kind="head")
    for i, q in enumerate(QUARTERS):
        put(ws, 28, qcol(i), q, kind="head")
    for i in range(7):
        r = 29 + i
        src = 6 + i
        put(ws, r, 1, f"C{i+1}", kind="label", bold=True)
        put(ws, r, 2, f"=B{src}", kind="calc")
        put(ws, r, 3, "DERIVED", kind="label")
        put(ws, r, 4, "− count  (disconnects reduce the print)", kind="note")
        for j in range(N_Q):
            L = qletter(j)
            put(ws, r, qcol(j),
                f'IF({L}{src}="","",-{L}{src})',
                kind="calc", fmt=NUM)

    # Selected-quarter engine for Dashboard (column B/C starting row 38)
    put(ws, 38, 1, "SELECTED-QUARTER ENGINE (Dashboard feed — do not edit)", kind="sec")
    merge(ws, 38, 1, 38, 6)
    put(ws, 39, 1, "#", kind="head")
    put(ws, 39, 2, "Reason", kind="head")
    put(ws, 39, 3, "Count", kind="head")
    put(ws, 39, 4, "Share", kind="head")
    put(ws, 39, 5, "Contribution", kind="head")
    put(ws, 39, 6, "Abs contrib", kind="head")
    put(ws, 39, 7, "Rank", kind="head")
    put(ws, 39, 8, "Is largest?", kind="head")
    put(ws, 39, 9, "Rest (gray)", kind="head")
    put(ws, 39, 10, "Largest (accent)", kind="head")

    for i in range(7):
        r = 40 + i
        put(ws, r, 1, i + 1, kind="label")
        put(ws, r, 2, f"=B{6+i}", kind="calc")
        put(ws, r, 3, f"={idx_q('Attribution', 6+i)}", kind="calc", fmt=NUM)
        put(ws, r, 4, f'IF(OR(C{r}="",$C$47="",$C$47=0),"",C{r}/$C$47)', kind="calc", fmt=PCT)
        put(ws, r, 5, f'IF(C{r}="","",-C{r})', kind="calc", fmt=NUM)
        put(ws, r, 6, f'IF(C{r}="",0,ABS(C{r}))', kind="calc", fmt=NUM)
        put(ws, r, 7, f'IF(SUM($F$40:$F$46)=0,"",RANK(F{r},$F$40:$F$46,0))', kind="calc")
        put(ws, r, 8, f'IF(G{r}=1,"LARGEST","")', kind="calc")
        put(ws, r, 9, f'IF(G{r}=1,0,IF(C{r}="",NA(),C{r}))', kind="calc", fmt=NUM)
        put(ws, r, 10, f'IF(G{r}=1,IF(C{r}="",NA(),C{r}),0)', kind="calc", fmt=NUM)

    put(ws, 47, 1, "Base", kind="label", bold=True)
    put(ws, 47, 2, "Disconnects / allocation base", kind="label")
    put(ws, 47, 3, f"={idx_q('Attribution', 14)}", kind="calc", fmt=NUM)
    put(ws, 48, 1, "Max share", kind="label", bold=True)
    put(ws, 48, 2, "Largest reason share of base", kind="label")
    put(ws, 48, 3, "=IF(COUNT(D40:D46)=0,\"\",MAX(D40:D46))", kind="calc", fmt=PCT)
    put(ws, 48, 4, "=IF(C48=\"\",\"\",INDEX(B40:B46,MATCH(C48,D40:D46,0)))", kind="calc")
    put(ws, 49, 1, "Entered?", kind="label", bold=True)
    put(ws, 49, 2, "Any reason count filled this quarter?", kind="label")
    put(ws, 49, 3, "=IF(COUNT(C40:C45)=0,\"No\",\"Yes\")", kind="calc")
    put(ws, 50, 1, "Second share", kind="label")
    put(ws, 50, 2, "2nd-largest share (mixed test)", kind="label")
    put(ws, 50, 3, "=IF(COUNT(D40:D46)<2,\"\",LARGE(D40:D46,2))", kind="calc", fmt=PCT)

    # Sorted view for chart categories (rank 1..7)
    put(ws, 52, 1, "SORTED FOR CHART (magnitude descending)", kind="sec")
    merge(ws, 52, 1, 52, 6)
    put(ws, 53, 1, "Rank", kind="head")
    put(ws, 53, 2, "Reason", kind="head")
    put(ws, 53, 3, "Count", kind="head")
    put(ws, 53, 4, "Share", kind="head")
    put(ws, 53, 5, "Rest (gray)", kind="head")
    put(ws, 53, 6, "Largest (accent)", kind="head")
    for k in range(7):
        r = 54 + k
        put(ws, r, 1, k + 1, kind="label")
        put(ws, r, 2,
            f'=IF($C$49="No","",IFERROR(INDEX($B$40:$B$46,MATCH({k+1},$G$40:$G$46,0)),""))',
            kind="calc")
        put(ws, r, 3,
            f'=IF(B{r}="","",IFERROR(INDEX($C$40:$C$46,MATCH({k+1},$G$40:$G$46,0)),""))',
            kind="calc", fmt=NUM)
        put(ws, r, 4,
            f'=IF(B{r}="","",IFERROR(INDEX($D$40:$D$46,MATCH({k+1},$G$40:$G$46,0)),""))',
            kind="calc", fmt=PCT)
        put(ws, r, 5, f'=IF(OR(B{r}="",C{r}=""),NA(),IF(A{r}=1,0,C{r}))', kind="calc", fmt=NUM)
        put(ws, r, 6, f'=IF(OR(B{r}="",C{r}=""),NA(),IF(A{r}=1,C{r},0))', kind="calc", fmt=NUM)

    # Exposure class (pptx sensor)
    put(ws, 63, 1, "EXPOSURE CLASS · pptx sensor (fiber-overbuilt / FWA / satellite-only / uncontested)", kind="sec")
    merge(ws, 63, 1, 63, 8)
    put(ws, 64, 1, "Primary exposure — assign each disconnect one class (MECE for the template; addresses can overlap in the real BDC join).", kind="note")
    merge(ws, 64, 1, 64, 8)
    put(ws, 65, 1, "ID", kind="head")
    put(ws, 65, 2, "Exposure class", kind="head")
    put(ws, 65, 3, "Tag", kind="head")
    put(ws, 65, 4, "What it tests", kind="head")
    for i, q in enumerate(QUARTERS):
        put(ws, 65, qcol(i), q, kind="head")
    for i, name in enumerate(EXPOSURE):
        r = 66 + i
        put(ws, r, 1, f"E{i+1}", kind="label", bold=True)
        put(ws, r, 2, name, kind="label")
        put(ws, r, 3, "INPUT", kind="label")
        put(ws, r, 4, "Disconnect counts in this class", kind="note")
        for j in range(N_Q):
            put(ws, r, qcol(j), None, kind="input", fmt=NUM)
    put(ws, 70, 1, "E5", kind="label", bold=True)
    put(ws, 70, 2, "Beginning PSUs in class — Uncontested (control)", kind="label")
    put(ws, 70, 3, "INPUT", kind="label")
    put(ws, 70, 4, "Needed for rate vs control (pptx slide 4 validation)", kind="note")
    for j in range(N_Q):
        put(ws, 70, qcol(j), None, kind="input", fmt=NUM)
    put(ws, 71, 1, "E6", kind="label", bold=True)
    put(ws, 71, 2, "Beginning PSUs — Fiber-overbuilt", kind="label")
    put(ws, 71, 3, "INPUT", kind="label")
    put(ws, 71, 4, "Optional — rate = disconnects / begin", kind="note")
    for j in range(N_Q):
        put(ws, 71, qcol(j), None, kind="input", fmt=NUM)
    put(ws, 72, 1, "E7", kind="label", bold=True)
    put(ws, 72, 2, "Beginning PSUs — FWA-covered", kind="label")
    put(ws, 72, 3, "INPUT", kind="label")
    put(ws, 72, 4, "Optional", kind="note")
    for j in range(N_Q):
        put(ws, 72, qcol(j), None, kind="input", fmt=NUM)
    put(ws, 73, 1, "E8", kind="label", bold=True)
    put(ws, 73, 2, "Beginning PSUs — Satellite-only", kind="label")
    put(ws, 73, 3, "INPUT", kind="label")
    put(ws, 73, 4, "Optional", kind="note")
    for j in range(N_Q):
        put(ws, 73, qcol(j), None, kind="input", fmt=NUM)

    put(ws, 75, 1, "EXPOSURE RATES vs UNCONTESTED CONTROL (pptx validation)", kind="sec")
    merge(ws, 75, 1, 75, 4)
    # selected quarter rates
    put(ws, 76, 1, "Class", kind="head")
    put(ws, 76, 2, "Disconnects (selected qtr)", kind="head")
    put(ws, 76, 3, "Begin PSUs", kind="head")
    put(ws, 76, 4, "Churn proxy (disc/begin)", kind="head")
    put(ws, 76, 5, "vs uncontested", kind="head")
    exp_disc_rows = [66, 67, 68, 69]
    exp_begin_rows = [71, 72, 73, 70]  # fiber, fwa, sat, uncontested
    for i, name in enumerate(EXPOSURE):
        r = 77 + i
        put(ws, r, 1, name, kind="label")
        put(ws, r, 2, f"={idx_q('Attribution', exp_disc_rows[i])}", kind="calc", fmt=NUM)
        put(ws, r, 3, f"={idx_q('Attribution', exp_begin_rows[i])}", kind="calc", fmt=NUM)
        put(ws, r, 4, f'IF(OR(B{r}="",C{r}="",C{r}=0),"",B{r}/C{r})', kind="calc", fmt=PCT1)
        put(ws, r, 5,
            f'IF(OR(D{r}="",$D$80=""),"",D{r}/$D$80)',
            kind="calc", fmt=MULT)

    put(ws, 82, 1,
        "Uncontested is the control group every other cohort is measured against (pptx slide 4). "
        "If fiber or FWA rate ≈ uncontested, the exposure map is not the attacker.",
        kind="note")
    merge(ws, 82, 1, 82, 6)

    ws.row_dimensions[5].height = 22


# ---------------------------------------------------------------------------
# Segments
# ---------------------------------------------------------------------------

def build_segments(ws):
    page(ws, freeze="E6", print_area="A1:P28", tab="FFFF00")
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 36
    for i in range(N_Q):
        ws.column_dimensions[qletter(i)].width = 13

    put(ws, 1, 1, "SEGMENT CUTS  —  WHERE and WHO overlay the reason tree", kind="title")
    merge(ws, 1, 1, 1, 8)
    put(
        ws,
        2,
        1,
        "PPTX: tenure · ARPU tier · exposure class (on Attribution) · rural as Scenario 3 (satellite/low-density). "
        "G1 / six-recommendation-areas also require: core vs rural Internet; AT&T ~27% / Verizon ~16% overlap vs rest "
        "(those shares are footprint mix, not losses); product mix at leave; sampled exit signal. "
        "Yellow = INPUT counts of disconnects (or net adds if labeled). Blank = not tracked.",
        kind="note",
    )
    merge(ws, 2, 1, 3, 12)
    ws.row_dimensions[2].height = 36

    def section(start, title, ids, names, note):
        put(ws, start, 1, title, kind="sec")
        merge(ws, start, 1, start, 8)
        put(ws, start + 1, 1, "ID", kind="head")
        put(ws, start + 1, 2, "Cut", kind="head")
        put(ws, start + 1, 3, "Tag", kind="head")
        put(ws, start + 1, 4, note, kind="head")
        for i, q in enumerate(QUARTERS):
            put(ws, start + 1, qcol(i), q, kind="head")
        for i, (lid, name) in enumerate(zip(ids, names)):
            r = start + 2 + i
            put(ws, r, 1, lid, kind="label", bold=True)
            put(ws, r, 2, name, kind="label")
            put(ws, r, 3, "INPUT", kind="label")
            put(ws, r, 4, "Disconnect counts", kind="note")
            for j in range(N_Q):
                put(ws, r, qcol(j), None, kind="input", fmt=NUM)
        tot_r = start + 2 + len(names)
        first = start + 2
        last = tot_r - 1
        put(ws, tot_r, 1, "", kind="label")
        put(ws, tot_r, 2, "Sum (formula)", kind="label", bold=True)
        put(ws, tot_r, 3, "DERIVED", kind="label")
        put(ws, tot_r, 4, "Should reconcile to tracker disconnects", kind="note")
        for j in range(N_Q):
            L = qletter(j)
            put(ws, tot_r, qcol(j),
                f'IF(COUNT({L}{first}:{L}{last})=0,"",SUM({L}{first}:{L}{last}))',
                kind="calc", fmt=NUM)
        return first, last, tot_r

    w_first, w_last, w_tot = section(
        5, "WHERE · core vs rural · overlap vs rest",
        ["W1", "W2", "W3", "W4", "W5"], WHERE,
        "AT&T ~27% / Verizon ~16% = footprint mix (10-K), not a loss split",
    )
    t_first, t_last, t_tot = section(
        14, "WHO · tenure at disconnect",
        ["N1", "N2", "N3"], TENURE,
        "PPTX S1 skews low-tenure; S2 skews established / high-ARPU",
    )
    p_first, p_last, p_tot = section(
        21, "WHO · product mix at leave",
        ["M1", "M2", "M3", "M4"], PRODUCT,
        "G1 cut. Tests whether video lock-in left and whether mobile holds",
    )
    a_first, a_last, a_tot = section(
        29, "WHO · ARPU tier (pptx value-cohort)",
        ["V1", "V2", "V3"], ARPU,
        "PPTX loss matrix = cause-cohort × value-cohort",
    )
    x_first, x_last, x_tot = section(
        36, "MEASURED EXIT SIGNAL (sampled survey / port-out — not agent codes)",
        ["X1", "X2", "X3", "X4", "X5"], EXIT,
        "PPTX slide 4: sample-audit codes against reality every quarter",
    )

    # Optional begin PSUs for core/rural rates
    put(ws, 45, 1, "OPTIONAL BEGIN PSUs FOR RATES (WHERE)", kind="sec")
    merge(ws, 45, 1, 45, 4)
    put(ws, 46, 1, "ID", kind="head")
    put(ws, 46, 2, "Cut", kind="head")
    put(ws, 46, 3, "Tag", kind="head")
    put(ws, 46, 4, "Beginning residential Internet PSUs in this cut", kind="head")
    for i, q in enumerate(QUARTERS):
        put(ws, 46, qcol(i), q, kind="head")
    for i, name in enumerate(["Core begin", "Rural begin"]):
        r = 47 + i
        put(ws, r, 1, f"B{i+1}", kind="label", bold=True)
        put(ws, r, 2, name, kind="label")
        put(ws, r, 3, "INPUT", kind="label")
        put(ws, r, 4, "Needed for a rate, not just a share", kind="note")
        for j in range(N_Q):
            put(ws, r, qcol(j), None, kind="input", fmt=NUM)

    # Selected-quarter diagnostic engine
    put(ws, 51, 1, "SELECTED-QUARTER DIAGNOSTIC (Dashboard feed — do not edit)", kind="sec")
    merge(ws, 51, 1, 51, 8)
    put(ws, 52, 1, "Cut", kind="head")
    put(ws, 52, 2, "Segment", kind="head")
    put(ws, 52, 3, "Disconnects", kind="head")
    put(ws, 52, 4, "Share of this cut", kind="head")
    put(ws, 52, 5, "Begin PSUs", kind="head")
    put(ws, 52, 6, "Rate", kind="head")
    put(ws, 52, 7, "vs company rate", kind="head")
    put(ws, 52, 8, "Flag", kind="head")

    # WHERE rows 53-57
    for i, name in enumerate(WHERE):
        r = 53 + i
        src = w_first + i
        put(ws, r, 1, "WHERE", kind="label")
        put(ws, r, 2, name, kind="label")
        put(ws, r, 3, f"={idx_q('Segments', src)}", kind="calc", fmt=NUM)
        put(ws, r, 4, f'IF(OR(C{r}="",SUM($C$53:$C$57)=0),"",C{r}/SUM($C$53:$C$57))', kind="calc", fmt=PCT)
        if i == 0:
            put(ws, r, 5, f"={idx_q('Segments', 47)}", kind="calc", fmt=NUM)
        elif i == 1:
            put(ws, r, 5, f"={idx_q('Segments', 48)}", kind="calc", fmt=NUM)
        else:
            put(ws, r, 5, "", kind="calc")
        put(ws, r, 6, f'IF(OR(C{r}="",E{r}="",E{r}=0),"",C{r}/E{r})', kind="calc", fmt=PCT1)
        put(ws, r, 7,
            f'IF(OR(F{r}="",\'Net-loss tracker\'!B43=""),"",F{r}/\'Net-loss tracker\'!B43)',
            kind="calc", fmt=MULT)
        put(ws, r, 8,
            f'IF(C{r}="","",IF(D{r}=MAX($D$53:$D$57),"Largest share",""))',
            kind="calc")

    put(ws, 59, 1, "Tenure", kind="head")
    put(ws, 59, 2, "Segment", kind="head")
    put(ws, 59, 3, "Disconnects", kind="head")
    put(ws, 59, 4, "Share", kind="head")
    for i, name in enumerate(TENURE):
        r = 60 + i
        put(ws, r, 1, "TENURE", kind="label")
        put(ws, r, 2, name, kind="label")
        put(ws, r, 3, f"={idx_q('Segments', t_first+i)}", kind="calc", fmt=NUM)
        put(ws, r, 4, f'IF(OR(C{r}="",SUM($C$60:$C$62)=0),"",C{r}/SUM($C$60:$C$62))', kind="calc", fmt=PCT)

    put(ws, 64, 1, "Product", kind="head")
    put(ws, 64, 2, "Segment", kind="head")
    put(ws, 64, 3, "Disconnects", kind="head")
    put(ws, 64, 4, "Share", kind="head")
    for i, name in enumerate(PRODUCT):
        r = 65 + i
        put(ws, r, 1, "PRODUCT", kind="label")
        put(ws, r, 2, name, kind="label")
        put(ws, r, 3, f"={idx_q('Segments', p_first+i)}", kind="calc", fmt=NUM)
        put(ws, r, 4, f'IF(OR(C{r}="",SUM($C$65:$C$68)=0),"",C{r}/SUM($C$65:$C$68))', kind="calc", fmt=PCT)

    put(ws, 70, 1, "Exit signal", kind="head")
    put(ws, 70, 2, "Signal", kind="head")
    put(ws, 70, 3, "Count", kind="head")
    put(ws, 70, 4, "Share", kind="head")
    for i, name in enumerate(EXIT):
        r = 71 + i
        put(ws, r, 1, "EXIT", kind="label")
        put(ws, r, 2, name, kind="label")
        put(ws, r, 3, f"={idx_q('Segments', x_first+i)}", kind="calc", fmt=NUM)
        put(ws, r, 4, f'IF(OR(C{r}="",SUM($C$71:$C$75)=0),"",C{r}/SUM($C$71:$C$75))', kind="calc", fmt=PCT)

    # Flags
    put(ws, 77, 1, "SEGMENT FLAGS", kind="sec")
    merge(ws, 77, 1, 77, 4)
    put(ws, 78, 1, "F1", kind="label", bold=True)
    put(ws, 78, 2, "WHERE driver", kind="label")
    put(ws, 78, 3,
        '=IF(COUNT(C53:C57)=0,"EMPTY — no WHERE fill",'
        'INDEX(B53:B57,MATCH(MAX(D53:D57),D53:D57,0))&" holds "&TEXT(MAX(D53:D57),"0%")&" of WHERE disconnects")',
        kind="calc")
    merge(ws, 78, 3, 78, 8)
    put(ws, 79, 1, "F2", kind="label", bold=True)
    put(ws, 79, 2, "Core vs rural rule", kind="label")
    put(ws, 79, 3,
        '=IF(COUNT(C53:C54)=0,"EMPTY",'
        'IF(D53>=Parameters!C14,"Core drives the print — do not add rural miles (W1 / kill rule 5)",'
        'IF(D54>=Parameters!C14,"Rural Internet (once filled) is the leak — take on lit plant, not more miles (W2)",'
        '"Neither cut clears the core-share threshold — read the reason play on both")))',
        kind="calc")
    merge(ws, 79, 3, 79, 8)
    put(ws, 80, 1, "F3", kind="label", bold=True)
    put(ws, 80, 2, "Overlap vs footprint", kind="label")
    put(ws, 80, 3,
        '=IF(COUNT(C55:C57)=0,"EMPTY — 27%/16% are footprint mix, not losses",'
        'IF(AND(ABS(D55-0.27)<0.05,ABS(D56-0.16)<0.05),'
        '"Losses follow footprint (~27% / ~16%) — overlap is not Root B; kill overlap GTM spend",'
        'IF(D55+D56>=0.5,"Overlap zips concentrate losses — concentrate the reason play there (W3)",'
        '"Rest-of-footprint leak — fiber (R3) is the wrong play off-overlap (W4)")))',
        kind="calc")
    merge(ws, 80, 3, 80, 8)
    put(ws, 81, 1, "F4", kind="label", bold=True)
    put(ws, 81, 2, "Tenure / ARPU skew (pptx S1 vs S2)", kind="label")
    put(ws, 81, 3,
        '=IF(COUNT(C60:C62)=0,"EMPTY",'
        'IF(D60>=0.4,"Low-tenure skew — consistent with FWA / price (S1)",'
        'IF(D62>=0.4,"Long-tenure skew — check fiber / high-ARPU (S2)",'
        '"Tenure is mixed")))',
        kind="calc")
    merge(ws, 81, 3, 81, 8)
    put(ws, 82, 1, "F5", kind="label", bold=True)
    put(ws, 82, 2, "Exit signal (sampled)", kind="label")
    put(ws, 82, 3,
        '=IF(COUNT(C71:C75)=0,"EMPTY — do not use agent codes as a substitute",'
        'INDEX(B71:B75,MATCH(MAX(D71:D75),D71:D75,0))&" is the largest audited exit signal")',
        kind="calc")
    merge(ws, 82, 3, 82, 8)
    put(ws, 83, 1, "F6", kind="label", bold=True)
    put(ws, 83, 2, "Product at leave", kind="label")
    put(ws, 83, 3,
        '=IF(COUNT(C65:C68)=0,"EMPTY",'
        'INDEX(B65:B68,MATCH(MAX(D65:D68),D65:D68,0))&" is the largest leave mix")',
        kind="calc")
    merge(ws, 83, 3, 83, 8)

    put(ws, 85, 1,
        "27% / 16% on Parameters are not typed here — they are 10-K footprint shares. "
        "Do not fire W3 on those shares alone (decisions.md). Rural +41k CR is not W2.",
        kind="note")
    merge(ws, 85, 1, 85, 8)

    # store first rows as comments for dashboard via fixed addresses — already used
    ws.row_dimensions[6].height = 22


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

def build_dashboard(ws, wb):
    page(ws, freeze="A5", print_area="A1:N58", tab="1F4E79")
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = "1:4"
    for col, w in {
        "A": 3, "B": 28, "C": 16, "D": 16, "E": 16, "F": 16,
        "G": 16, "H": 16, "I": 14, "J": 14, "K": 14, "L": 14, "M": 14, "N": 18,
    }.items():
        ws.column_dimensions[col].width = w

    put(ws, 1, 2, "CHARTER REC 1  ·  ATTRIBUTION DIAGNOSTIC", kind="title")
    merge(ws, 1, 2, 1, 10)
    # Formula-driven takeaway title (the claim)
    put(
        ws,
        2,
        2,
        '=IF(C15="","Fill the quarterly net-loss tracker first — attribution cells are empty on purpose.",C15)',
        kind="h2",
    )
    merge(ws, 2, 2, 2, 13)
    ws.row_dimensions[2].height = 28
    put(
        ws,
        3,
        2,
        '=CONCATENATE("Selected: ",Parameters!C5,"   ·   Units: residential Internet customers   ·   "'
        '&"Yellow INPUT on Net-loss tracker / Attribution / Segments   ·   "'
        '&"Public spine: Ex99.1 Q1 2026 res. Internet −117,000 · Q2 2026 −166,000 (RETRIEVED)   ·   "'
        '&"Template — not a filled mix")',
        kind="note",
    )
    merge(ws, 3, 2, 3, 13)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[3].height = 18

    # KPI cards
    cards = [
        (5, 2, "Res. Internet print",
         "='Net-loss tracker'!B37", NUM,
         "='Net-loss tracker'!B48"),
        (5, 5, "vs Comcast (same quarter)",
         '=IF(OR(Parameters!C18="No",\'Net-loss tracker\'!B40=""),"n/a",\'Net-loss tracker\'!B40)',
         NUM, '="CMCSA resid. BB"'),
        (5, 8, "Excess vs peer rate",
         '=IF(OR(Parameters!C18="No",\'Net-loss tracker\'!B46=""),"n/a",\'Net-loss tracker\'!B46)',
         NUM, '="Q1 playbook commit 54k"'),
        (5, 11, "Largest reason",
         '=IF(Attribution!C49="No","EMPTY",Attribution!D48)',
         None, '=IF(Attribution!C49="No","do not fire a play",TEXT(Attribution!C48,"0% of disconnects"))'),
    ]
    for r, c, title, val, fmt, sub in cards:
        put(ws, r, c, title, kind="label")
        merge(ws, r, c, r, c + 2)
        put(ws, r + 1, c, val, kind="kpi", fmt=fmt)
        merge(ws, r + 1, c, r + 1, c + 2)
        put(ws, r + 2, c, sub, kind="note")
        merge(ws, r + 2, c, r + 2, c + 2)
        ws.row_dimensions[r + 1].height = 26

    # Diagnosis traffic lights
    put(ws, 9, 2, "DIAGNOSIS PANEL  ·  thresholds live on Parameters (ESTIMATED defaults)", kind="sec")
    merge(ws, 9, 2, 9, 13)
    put(ws, 10, 2, "Test", kind="head")
    put(ws, 10, 3, "Read", kind="head")
    merge(ws, 10, 3, 10, 10)
    put(ws, 10, 11, "Light", kind="head")
    merge(ws, 10, 11, 10, 13)

    tests = [
        (11, "G1 tracker",
         '=Parameters!C17',
         '=IF(LEFT(Parameters!C17,3)="Yes","SKIP BUILD","FILL FIRST")'),
        (12, "Dominant driver (pptx >= P4)",
         '=IF(Attribution!C49="No","EMPTY — no reason fill",'
         'IF(Attribution!C48>=Parameters!C8,'
         '"DOMINANT — "&Attribution!D48&" holds "&TEXT(Attribution!C48,"0%")&" (>= "&TEXT(Parameters!C8,"0%")&")",'
         'IF(Attribution!C48>=Parameters!C9,'
         '"MIXED — largest is "&TEXT(Attribution!C48,"0%")&" (25-40% band). Run top two as experiments.",'
         '"NULL — no cohort >= "&TEXT(Parameters!C9,"0%")&". Scenario 6 fires.")))',
         '=IF(Attribution!C49="No","EMPTY",'
         'IF(Attribution!C48>=Parameters!C8,"RED — fire one play",'
         'IF(Attribution!C48>=Parameters!C9,"AMBER — experiment","GREEN — null / S6")))'),
        (13, "Scenario routed",
         "=C16",
         '=IF(Attribution!C49="No","HOLD","SEE PLAYBOOK")'),
        (14, "WHERE / WHO flag",
         "=Segments!C78",
         '=IF(LEFT(Segments!C78,5)="EMPTY","EMPTY","READ")'),
    ]
    for r, name, read, light in tests:
        put(ws, r, 2, name, kind="label", bold=True)
        put(ws, r, 3, read, kind="calc")
        merge(ws, r, 3, r, 10)
        put(ws, r, 11, light, kind="calc", bold=True, align=center)
        merge(ws, r, 11, r, 13)
        ws.row_dimensions[r].height = 22

    put(ws, 15, 2, "Takeaway (engine)", kind="label")
    put(
        ws,
        15,
        3,
        '=IF(Attribution!C49="No",'
        '"Fill the quarterly net-loss tracker first — attribution cells are empty on purpose.",'
        'IF(Attribution!C48>=Parameters!C8,'
        'Attribution!D48&" accounts for "&TEXT(Attribution!C48,"0%")&" of disconnects — above the "'
        '&TEXT(Parameters!C8,"0%")&" dominant-driver gate. Fire that one play on the core.",'
        'IF(Attribution!C48>=Parameters!C9,'
        '"Mixed picture: "&Attribution!D48&" is largest at "&TEXT(Attribution!C48,"0%")'
        '&" but below "&TEXT(Parameters!C8,"0%")&". Run the top two as controlled experiments.",'
        '"No dominant driver — no cohort reaches "&TEXT(Parameters!C9,"0%")'
        '&" of disconnects. Scenario 6 (null) is the recommendation, not a guess.")))',
        kind="calc",
    )
    merge(ws, 15, 3, 15, 13)
    ws.row_dimensions[15].hidden = True

    put(ws, 16, 2, "Scenario (engine)", kind="label")
    put(
        ws,
        16,
        3,
        '=IF(Attribution!C49="No","HOLD — empty cells. The playbook does not fire.",'
        'IF(AND(INDEX(Attribution!D40:D45,1)>=Parameters!C10,N(Segments!D71)>=0.3),'
        '"S1 FWA-driven — convergence pricing in FWA zones (kill: no churn split in 2 qtrs / negative bundle margin)",'
        'IF(INDEX(Attribution!D40:D45,2)>=Parameters!C11,'
        '"S2 Fiber-driven — re-sequence evolution on overlap + symmetric win-back (kill: completed zones bleed at same rate)",'
        'IF(INDEX(Attribution!D40:D45,3)>=Parameters!C12,'
        '"S3 Satellite-driven — accelerate rural activation (kill: new-passing penetration below underwriting)",'
        'IF(INDEX(Attribution!D40:D45,4)>=Parameters!C9,'
        '"S4 Housing / move-out — mover capture, not a save-offer (kill: port-out data contradicts the macro story)",'
        'IF(INDEX(Attribution!D40:D45,5)>=Parameters!C9,'
        '"S5 Non-pay — value tier at delinquency (kill: save-rate economics below LTV)",'
        'IF(Attribution!C48<Parameters!C9,'
        '"S6 Null — no dominant driver. Market-layer disclosure + spend-capped experiments.",'
        '"Largest cell is "&Attribution!D48&" — apply the matching lead play; do not fire two expensive levers."))))))',
        kind="calc",
    )
    merge(ws, 16, 3, 16, 13)
    ws.row_dimensions[16].hidden = True

    put(ws, 18, 2,
        '=IF(Attribution!C49="No",'
        '"Reason contribution is empty — fill Attribution before this chart can make a claim",'
        '"Largest leak is "&Attribution!D48&" at "&TEXT(Attribution!C48,"0%")'
        '&" of sampled disconnects  ·  gray = the rest  ·  accent = the story bar")',
        kind="h2")
    merge(ws, 18, 2, 18, 13)
    ws.row_dimensions[18].height = 22

    put(ws, 19, 2, "Reason (sorted)", kind="head")
    put(ws, 19, 3, "Disconnects", kind="head")
    put(ws, 19, 4, "Share", kind="head")
    put(ws, 19, 5, "Rest", kind="head")
    put(ws, 19, 6, "Largest", kind="head")
    for i in range(7):
        r = 20 + i
        put(ws, r, 2, f"=Attribution!B{54+i}", kind="calc")
        put(ws, r, 3, f"=Attribution!C{54+i}", kind="calc", fmt=NUM)
        put(ws, r, 4, f"=Attribution!D{54+i}", kind="calc", fmt=PCT)
        put(ws, r, 5, f"=Attribution!E{54+i}", kind="calc", fmt=NUM)
        put(ws, r, 6, f"=Attribution!F{54+i}", kind="calc", fmt=NUM)

    put(ws, 27, 2,
        "Source: Attribution A1-A7 (user-entered counts) · Residual is derived · "
        "Bars begin at zero · One accent (largest) vs gray · Not a pie",
        kind="note")
    merge(ws, 27, 2, 27, 6)

    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.style = 10
    chart.y_axis.scaling.min = 0
    chart.y_axis.numFmt = "#,##0"
    chart.x_axis.numFmt = "@"
    chart.legend.position = "b"
    data = Reference(ws, min_col=5, min_row=19, max_col=6, max_row=26)
    cats = Reference(ws, min_col=2, min_row=20, max_row=26)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.height = 8
    chart.width = 15
    if len(chart.series) >= 2:
        style_series(chart.series[0], GRAY)
        style_series(chart.series[1], ACCENT)
    link_chart_title(chart, "'Dashboard'!$B$18")
    chart.y_axis.title = None
    chart.x_axis.title = None
    ws.add_chart(chart, "H19")

    put(ws, 29, 2,
        '=IF(\'Net-loss tracker\'!J11="",'
        '"Residential Internet print — public spine only until internals fill",'
        '"Charter residential Internet stayed worse than Comcast in every quarter both prints exist")',
        kind="h2")
    merge(ws, 29, 2, 29, 13)

    put(ws, 30, 2, "Quarter", kind="head")
    for i, q in enumerate(QUARTERS):
        put(ws, 30, 3 + i, q, kind="head")
    put(ws, 31, 2, "CHTR res. Internet", kind="label", bold=True)
    put(ws, 32, 2, "CMCSA resid. BB", kind="label")
    put(ws, 33, 2, "Excess vs CMCSA rate", kind="label")
    for i in range(N_Q):
        L = qletter(i)
        put(ws, 31, 3 + i, f"=IF('Net-loss tracker'!{L}11=\"\",NA(),'Net-loss tracker'!{L}11)",
            kind="calc", fmt=NUM)
        put(ws, 32, 3 + i, f"=IF('Net-loss tracker'!{L}14=\"\",NA(),'Net-loss tracker'!{L}14)",
            kind="calc", fmt=NUM)
        put(ws, 33, 3 + i, f"=IF('Net-loss tracker'!{L}20=\"\",NA(),'Net-loss tracker'!{L}20)",
            kind="calc", fmt=NUM)

    line = LineChart()
    line.style = 10
    line.height = 7
    line.width = 18
    line.legend.position = "b"
    ldata = Reference(ws, min_col=2, min_row=31, max_col=2 + N_Q, max_row=32)
    lcats = Reference(ws, min_col=3, min_row=30, max_col=2 + N_Q)
    line.add_data(ldata, from_rows=True, titles_from_data=True)
    line.set_categories(lcats)
    if len(line.series) >= 2:
        style_series(line.series[0], ACCENT)
        style_series(line.series[1], GRAY)
        line.series[0].marker = Marker(symbol="circle", size=7)
        line.series[1].marker = Marker(symbol="circle", size=7)
        line.series[0].marker.graphicalProperties.solidFill = ACCENT
        line.series[1].marker.graphicalProperties.solidFill = GRAY
    link_chart_title(line, "'Dashboard'!$B$29")
    ws.add_chart(line, "B34")

    put(ws, 34, 11, "Axis note: line chart may zoom off zero because the story is movement in a tight negative band. Bars on this sheet still start at zero.", kind="note")
    merge(ws, 34, 11, 36, 13)

    put(ws, 48, 2, "KILL RULES AND UNWILLING-TO-CLAIM  ·  written before the recommendation", kind="sec")
    merge(ws, 48, 2, 48, 13)
    put(ws, 49, 2, "#", kind="head")
    put(ws, 49, 3, "Rule (from pptx + six-recommendation-areas / decisions.md)", kind="head")
    merge(ws, 49, 3, 49, 10)
    put(ws, 49, 11, "This quarter", kind="head")
    merge(ws, 49, 11, 49, 13)

    kills = [
        (50, "K1",
         "Tracker already exists and is used quarterly -> skip Area 1 build; go to mix control.",
         '=IF(LEFT(Parameters!C17,3)="Yes","KILL — skip G1 build","OPEN — fill G1")'),
        (51, "K2",
         "Fill shows move-out matching Comcast -> kill Area 2 as a Charter-specific print; Area 6 becomes the CEO-charge answer. Do not add rural miles.",
         '=IF(OR(Attribution!C49="No",\'Net-loss tracker\'!B46=""),"UNTESTED",'
         'IF(AND(Attribution!D48=Parameters!B29,\'Net-loss tracker\'!B46<=0),'
         '"KILL Area 2 — move-out / no excess","OPEN")'),
        (52, "K3",
         "Do not fire two expensive levers in the same quarter. Co-equal cells: cheaper first (Parameters P9).",
         '=IF(AND(Attribution!C49="Yes",Attribution!C50>=Parameters!C9,'
         'ABS(Attribution!C48-Attribution!C50)<=Parameters!C13),'
         '"CO-EQUAL — fire the cheaper lever only","OK")'),
        (53, "K4",
         "Do not treat call-center reason codes as the trigger. Sample the disconnect (pptx slide 4).",
         "PROCESS — not a cell"),
        (54, "K5",
         "Do not add rural miles to fix a core leak. Do not fire company-wide upgrade/rebuild as the default to move-out, FWA, nonpay, or off-overlap.",
         '=IF(LEFT(Segments!C79,4)="Core","CORE — no rural miles",'
         'IF(LEFT(Segments!C78,5)="EMPTY","UNTESTED","READ F2"))'),
        (55, "K6",
         "PPTX interlock: two consecutive quarters before full-scale (experiments may start on one).",
         '=IF(Parameters!C16="Yes","ARMED — full-scale needs 2 qtrs","OFF")'),
        (56, "K7",
         "Unwilling to claim: a filled mix, a gross churn rate, that -66k re-rates five-year ~80%, that mobile offsets Internet, that 54k sits in any one cell until the first fill.",
         "STANDING"),
    ]
    for r, kid, text, status in kills:
        put(ws, r, 2, kid, kind="label", bold=True)
        put(ws, r, 3, text, kind="note")
        merge(ws, r, 3, r, 10)
        put(ws, r, 11, status, kind="calc", bold=True, align=center)
        merge(ws, r, 11, r, 13)
        ws.row_dimensions[r].height = 28

    ws.conditional_formatting.add(
        "K11:M14",
        FormulaRule(formula=['LEFT(K11,4)="FILL"'], fill=fill_warn, font=font_b),
    )
    ws.conditional_formatting.add(
        "K11:M14",
        FormulaRule(formula=['LEFT(K11,3)="RED"'], fill=fill_fail, font=font_b),
    )
    ws.conditional_formatting.add(
        "K11:M14",
        FormulaRule(formula=['LEFT(K11,5)="AMBER"'], fill=fill_warn, font=font_b),
    )
    ws.conditional_formatting.add(
        "K11:M14",
        FormulaRule(formula=['LEFT(K11,5)="GREEN"'], fill=fill_pass, font=font_b),
    )
    ws.conditional_formatting.add(
        "K11:M14",
        FormulaRule(formula=['LEFT(K11,4)="SKIP"'], fill=fill_pass, font=font_b),
    )
    ws.conditional_formatting.add(
        "K50:M56",
        FormulaRule(formula=['LEFT(K50,4)="KILL"'], fill=fill_fail, font=font_b),
    )
    ws.conditional_formatting.add(
        "K50:M56",
        FormulaRule(formula=['LEFT(K50,4)="OPEN"'], fill=fill_warn, font=font_b),
    )

    put(ws, 58, 2,
        "Print-friendly: this sheet is the exhibit. Backup arithmetic lives on Net-loss tracker, Attribution, Segments, Assumptions. "
        "Do not present empty attribution cells as a measured mix. Title (row 2) updates from formulas when numbers are pasted.",
        kind="note")
    merge(ws, 58, 2, 58, 13)


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------

def build_cover(ws):
    page(ws, landscape=False, freeze="A5", print_area="A1:B48", tab="7F7F7F")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 88
    put(ws, 1, 1, "COVER  ·  HOW TO USE THIS TEMPLATE", kind="title")
    merge(ws, 1, 1, 1, 2)
    put(ws, 2, 1,
        "A diagnostic template for privatized / internal numbers. "
        "It is not a filled Charter mix and it does not fire Rec 1 from empty cells.",
        kind="h2")
    merge(ws, 2, 1, 2, 2)
    ws.row_dimensions[2].height = 28

    put(ws, 4, 1, "What this is", kind="sec")
    merge(ws, 4, 1, 4, 2)
    put(ws, 5, 1, "Job", kind="label", bold=True)
    put(ws, 5, 2,
        "Rec 1 sensor + playbook from CHTR_Rec1_AttributionPlaybook.pptx, wired to the G1 quarterly "
        "net-loss tracker language in Decisions/six-recommendation-areas.md and Decisions/decisions.md. "
        "Charter (or the case team) pastes internals; Excel diagnoses them.",
        kind="note")
    put(ws, 6, 1, "What it is not", kind="label", bold=True)
    put(ws, 6, 2,
        "Not a study. Not invented quarterly net-adds presented as Ex99.1. "
        "Not a reason play you can take to the CEO while Attribution is empty. "
        "Not a password-locked model.",
        kind="note")

    put(ws, 8, 1, "Color and tags", kind="sec")
    merge(ws, 8, 1, 8, 2)
    put(ws, 9, 1, "Yellow + blue font", kind="input")
    put(ws, 9, 2, "INPUT — user-entered privatized / internal. Unlock these; type over them.", kind="note")
    put(ws, 10, 1, "Green", kind="retr")
    put(ws, 10, 2, "RETRIEVED — public Ex99.1 / trending / CMCSA Ex99.1. Prefill is a citation, not an attribution.", kind="note")
    put(ws, 11, 1, "Blue fill", kind="est")
    put(ws, 11, 2, "ESTIMATED — template default thresholds from the pptx (40 / 25 / 35). Change them on Parameters before go-live.", kind="note")
    put(ws, 12, 1, "Gray", kind="calc")
    put(ws, 12, 2, "DERIVED — Excel formula. Do not overwrite. Recompute happens in the workbook, not in this cover.", kind="note")

    put(ws, 14, 1, "How the user fills it", kind="sec")
    merge(ws, 14, 1, 14, 2)
    steps = [
        ("1. Parameters",
         "Set the selected quarter (default Q1 2026 — the assignment anchor; Q2 is a later public fact). "
         "Answer tracker already exists? If yes, skip the G1 build. Confirm Comcast benchmark on, and leave thresholds at pptx defaults unless the executive team ratifies new ones."),
        ("2. Net-loss tracker first",
         "Paste beginning residential Internet PSUs, gross adds, and disconnects in yellow. "
         "Net is a formula (T2 - T3). Public prints are already filled where Ex99.1 exists "
         "(Q1 2026 -117,000; Q2 2026 -166,000). Do not invent internals to match them. "
         "Comcast Q1 -65,000 / begin 28,719,000 is filled; other Comcast quarters stay blank."),
        ("3. Attribution",
         "Paste sampled disconnect COUNTS by cause-cohort (FWA, fiber, satellite, move-out, non-pay, service/other). "
         "Not agent codes. Residual is a formula against tracker disconnects. "
         "Exposure-class counts (and optional begin PSUs) test the pptx sensor vs the uncontested control."),
        ("4. Segments",
         "Paste WHERE (core / rural / AT&T / Verizon / rest) and WHO (tenure, product at leave, ARPU, audited exit signal). "
         "27% / 16% are footprint mix — do not paste them as losses. Rural CR +41k is not the rural Internet cut."),
        ("5. Read Dashboard only after 2-4",
         "Row 2 is the takeaway (formula). Traffic lights apply pptx 40 / 25 / 35 gates. "
         "One accent bar vs gray. Kill rules sit on the same page. "
         "If Attribution is empty, the title will tell you to fill the tracker first — believe it."),
    ]
    for i, (h, t) in enumerate(steps):
        r = 15 + i
        put(ws, r, 1, h, kind="label", bold=True)
        put(ws, r, 2, t, kind="note")
        ws.row_dimensions[r].height = 48

    put(ws, 21, 1, "Sheet map", kind="sec")
    merge(ws, 21, 1, 21, 2)
    sheets = [
        ("Dashboard", "Print-first exhibit. Formula titles, KPIs, reason bars, CHTR vs CMCSA line, diagnosis, kill rules."),
        ("Cover", "This page."),
        ("Net-loss tracker", "G1. Quarters as columns. Internals yellow; public prints green; net / rates / excess formulas."),
        ("Attribution", "Reason counts, residual, shares, contribution, exposure class, sorted chart feed."),
        ("Segments", "WHERE / WHO input grids + flags (core vs rural, overlap, tenure, exit signal)."),
        ("Parameters", "Quarter, taxonomy, thresholds, Comcast toggle, highlight hex, dropdowns."),
        ("Assumptions", "One line per input, ID, RETRIEVED vs INPUT vs ESTIMATED, source."),
        ("Playbook", "Six pptx scenarios + GAP R1-R4: trigger, play, economics gate, proof, kill, fallback. Reference only."),
    ]
    for i, (n, d) in enumerate(sheets):
        put(ws, 22 + i, 1, n, kind="label", bold=True)
        put(ws, 22 + i, 2, d, kind="note")

    put(ws, 31, 1, "Public figures prefilled (RETRIEVED)", kind="sec")
    merge(ws, 31, 1, 31, 2)
    put(ws, 32, 1, "Residential Internet", kind="label", bold=True)
    put(ws, 32, 2,
        "Q1 2026 -117,000 · Q2 2026 -166,000 · plus 2024-2025 quarterly prints from the Q1 trending schedule. "
        "Company Internet -120,000 / -172,000 is context only.",
        kind="note")
    put(ws, 33, 1, "Begin PSUs", kind="label", bold=True)
    put(ws, 33, 2,
        "Q1 2026 begin 27,641,000 (YE2025 EOP). Other begins from trending EOP identity. Overwrite if the internal census differs.",
        kind="note")
    put(ws, 34, 1, "Comcast", kind="label", bold=True)
    put(ws, 34, 2,
        "Q1 2026 domestic residential BB -65,000; begin 28,719,000. Excess vs that rate is the 54k identity "
        "(commit 54,000, round down from 54.4k) — cited from six-recommendation-areas.md / network_churn_model.py, not the pptx.",
        kind="note")

    put(ws, 36, 1, "Could not map from the pptx alone", kind="sec")
    merge(ws, 36, 1, 36, 2)
    put(ws, 37, 1, "See return note", kind="label", bold=True)
    put(ws, 37, 2,
        "The pptx does not specify the G1 grid, the 54k Comcast excess, product-mix-at-leave, or a numeric satellite threshold. "
        "Those come from six-recommendation-areas.md / decisions.md (labeled). "
        "MVNO wholesale cost, LTV gates, port-out carrier shares, BDC location joins, and treatment/control results "
        "are open client items (pptx slide 12) — they are not fake-filled here.",
        kind="note")
    ws.row_dimensions[37].height = 48

    put(ws, 39, 1, "Question for the client", kind="sec")
    merge(ws, 39, 1, 39, 2)
    put(ws, 40, 1, "Binary", kind="label", bold=True)
    put(ws, 40, 2,
        "Will you stand up — or do you already have — a quarterly net-loss-by-category tracker "
        "for residential Internet on the core: yes or no?",
        kind="h2")
    merge(ws, 40, 1, 40, 2)
    ws.row_dimensions[40].height = 32

    for r in (5, 6, 32, 33, 34):
        ws.row_dimensions[r].height = 36


# ---------------------------------------------------------------------------
# Assumptions
# ---------------------------------------------------------------------------

def build_assumptions(ws):
    page(ws, freeze="A5", print_area="A1:G80", tab="548235")
    for col, w in {"A": 8, "B": 52, "C": 16, "D": 14, "E": 22, "F": 56, "G": 14}.items():
        ws.column_dimensions[col].width = w
    put(ws, 1, 1, "ASSUMPTION REGISTER  ·  one line per input", kind="title")
    merge(ws, 1, 1, 1, 6)
    put(ws, 2, 1,
        "Solving playbook: one assumption per line with a line ID; derived lines are formulas. "
        "Tag every input RETRIEVED or INPUT (user-entered). Thresholds are ESTIMATED template defaults. "
        "Dollar-weight is not meaningful until Attribution is filled — then the mix is 100% INPUT.",
        kind="note")
    merge(ws, 2, 1, 3, 6)

    headers = ["ID", "Line", "Value / pointer", "Tag", "Origin", "Formula or citation", "Confidence"]
    for i, h in enumerate(headers, 1):
        put(ws, 5, i, h, kind="head")

    rows = [
        ("T1", "Beginning res. Internet PSUs by quarter", "Net-loss tracker!E6:P6", "RETRIEVED / INPUT",
         "Ex99.1 trending EOP identity; overwrite if internal", "Prior EOP. Q1 2026 = 27,641,000", "High where green"),
        ("T2", "Gross adds", "Net-loss tracker!E7:P7", "INPUT", "Internal billing", "Blank until filled", "—"),
        ("T3", "Disconnects", "Net-loss tracker!E8:P8", "INPUT", "Sampled disconnect file, not agent codes", "Blank until filled", "—"),
        ("T4", "Internal net", "Net-loss tracker!E9:P9", "DERIVED", "T2 - T3", "=gross-disconnects", "Formula"),
        ("T5", "Public res. Internet print", "Net-loss tracker!E10:P10", "RETRIEVED",
         "CHTR Ex99.1 / Q1 2026 trending", "Q1 2026 -117,000 · Q2 2026 -166,000", "High"),
        ("T6", "Print used", "Net-loss tracker!E11:P11", "DERIVED", "T4 if both internals filled, else T5", "IF(T4<>\"\",T4,T5)", "Formula"),
        ("T7", "Company Internet print", "Net-loss tracker!E12:P12", "RETRIEVED", "Ex99.1 total Internet",
         "Q1 -120,000 · Q2 -172,000 — context, not T6", "High"),
        ("T8", "Rural CR adds", "Net-loss tracker!E13:P13", "RETRIEVED", "Ex99.1 rural-footprint CR",
         "Q1 +41,000 · Q2 +47,000 — NOT an Internet cut", "High / wrong unit"),
        ("T9", "CMCSA resid. BB net adds", "Net-loss tracker!E14:P14", "RETRIEVED / INPUT",
         "CMCSA Ex99.1 Q1 2026; other qtrs INPUT", "Q1 2026 -65,000", "High for Q1"),
        ("T10", "CMCSA begin PSUs", "Net-loss tracker!E15:P15", "RETRIEVED / INPUT",
         "28,654k EOP - (-65k)", "Q1 2026 28,719,000", "High for Q1"),
        ("T11", "CMCSA net-loss rate", "Net-loss tracker!E16:P16", "DERIVED", "|T9|/T10", "Q1 ~ 0.226%", "Formula"),
        ("T12", "CHTR res. net-loss rate", "Net-loss tracker!E17:P17", "DERIVED", "|T6|/T1", "Q1 ~ 0.423%", "Formula"),
        ("T13", "Rate ratio", "Net-loss tracker!E18:P18", "DERIVED", "T12/T11", "Unrounded ~1.87x; commit 1.8x", "Formula"),
        ("T14", "Expected losses at peer rate", "Net-loss tracker!E19:P19", "DERIVED", "T11*T1", "Q1 ~ 62.6k", "Formula"),
        ("T15", "Excess vs Comcast", "Net-loss tracker!E20:P20", "DERIVED", "|T6|-T14",
         "Q1 ~ 54.4k; P18 commits 54,000 (six-recommendation-areas.md)", "Formula"),
        ("P18", "Playbook committed excess", "Parameters!C22", "RETRIEVED",
         "six-recommendation-areas.md / network_churn_model.py — not the pptx", "54,000 customers", "Commit / round down"),
        ("A1-A6", "Disconnect counts by cause-cohort", "Attribution!E6:P11", "INPUT",
         "Sampled disconnect / port-out / audited survey", "Blank — PLACEHOLDER empty", "—"),
        ("A7", "Residual", "Attribution!E12:P12", "DERIVED", "T3 - SUM(A1:A6)", "Data-quality flag if material", "Formula"),
        ("E1-E4", "Exposure-class disconnects", "Attribution!E66:P69", "INPUT",
         "BDC join (pptx sensor)", "Blank", "—"),
        ("E5-E8", "Exposure-class begin PSUs", "Attribution!E70:P73", "INPUT",
         "Needed for rate vs uncontested control", "Blank", "—"),
        ("W1-W5", "WHERE disconnects", "Segments!E7:P11", "INPUT",
         "Core / rural / AT&T / Verizon / rest", "Blank. 27%/16% are footprint, not this row", "—"),
        ("N1-N3", "Tenure disconnects", "Segments!E16:P18", "INPUT", "Billing tenure", "Blank", "—"),
        ("M1-M4", "Product mix at leave", "Segments!E23:P26", "INPUT", "G1 cut", "Blank", "—"),
        ("V1-V3", "ARPU tier", "Segments!E31:P33", "INPUT", "PPTX value-cohort", "Blank", "—"),
        ("X1-X5", "Audited exit signal", "Segments!E38:P42", "INPUT",
         "Survey / port-out — not agent codes", "Blank", "—"),
        ("P4", "Dominant threshold", "Parameters!C8", "ESTIMATED", "PPTX slide 5", "40%", "Template default"),
        ("P5", "Mixed floor", "Parameters!C9", "ESTIMATED", "PPTX slide 5", "25%", "Template default"),
        ("P6", "FWA trigger", "Parameters!C10", "ESTIMATED", "PPTX S1 >35-40%", "35%", "Template default"),
        ("P7", "Fiber trigger", "Parameters!C11", "ESTIMATED", "PPTX S2 dominate", "40%", "Template default"),
        ("P8", "Satellite trigger", "Parameters!C12", "ESTIMATED",
         "PPTX S3 has no numeric cut — template default", "25%", "Low — user should ratify"),
        ("P9", "Co-equal band", "Parameters!C13", "ESTIMATED", "GAP cheaper-first", "5 pp", "Template default"),
        ("P10", "Core-share flag", "Parameters!C14", "ESTIMATED", "decisions.md W1", "70%", "Template default"),
        ("P12", "Two-quarter interlock", "Parameters!C16", "ESTIMATED", "PPTX slide 11", "Yes", "Template default"),
    ]
    for i, row in enumerate(rows):
        r = 6 + i
        for c, val in enumerate(row, 1):
            put(ws, r, c, val, kind="label" if c <= 4 else "note", bold=(c == 1))
        ws.row_dimensions[r].height = 20

    last = 6 + len(rows)
    put(ws, last + 1, 1, "WEIGHT", kind="sec")
    merge(ws, last + 1, 1, last + 1, 6)
    put(ws, last + 2, 1, "W1", kind="label", bold=True)
    put(ws, last + 2, 2, "Count of register lines by tag (this template, before a fill)", kind="label")
    put(ws, last + 2, 3,
        '=COUNTIF(D6:D' + str(last - 1) + ',"RETRIEVED")&" RETRIEVED  ·  "&'
        'COUNTIF(D6:D' + str(last - 1) + ',"INPUT")&" INPUT  ·  "&'
        'COUNTIF(D6:D' + str(last - 1) + ',"ESTIMATED")&" ESTIMATED  ·  "&'
        'COUNTIF(D6:D' + str(last - 1) + ',"DERIVED")&" DERIVED"',
        kind="calc")
    merge(ws, last + 2, 3, last + 2, 6)
    put(ws, last + 3, 1, "W2", kind="label", bold=True)
    put(ws, last + 3, 2, "Dollar-weighted once Attribution is filled", kind="label")
    put(ws, last + 3, 3,
        '=IF(Attribution!C49="No",'
        '"Not meaningful — mix is empty. Do not quote a retrieved % of a mix that does not exist.",'
        '"Once filled, 100% of the reason mix is INPUT (sampled disconnects). Public prints do not validate the mix.")',
        kind="calc")
    merge(ws, last + 3, 3, last + 3, 6)

    ws.auto_filter.ref = f"A5:G{last - 1}"


# ---------------------------------------------------------------------------
# Playbook
# ---------------------------------------------------------------------------

def build_playbook(ws):
    page(ws, freeze="A6", print_area="A1:G20", tab="2E75B6")
    for col, w in {"A": 6, "B": 22, "C": 36, "D": 40, "E": 32, "F": 36, "G": 32}.items():
        ws.column_dimensions[col].width = w
    put(ws, 1, 1, "PLAYBOOK ON A PAGE  ·  reference only — does not fire from empty cells", kind="title")
    merge(ws, 1, 1, 1, 7)
    put(ws, 2, 1,
        "PPTX slides 6-9 (scenarios 1-6) plus GAP R1-R4 / WHERE overlays from decisions.md. "
        "Budgets and owners are pre-authorized in the deck, not sized here. "
        "Dashboard routes a scenario; this sheet is the leave-behind.",
        kind="note")
    merge(ws, 2, 1, 3, 7)

    headers = ["#", "Trigger (signal)", "Lead play", "Economics gate", "Proof", "Kill", "Fallback"]
    for i, h in enumerate(headers, 1):
        put(ws, 5, i, h, kind="head")
    plays = [
        ("S1",
         "FWA-covered losses >35-40%; low-tenure / low-ARPU; audited exit cites price.",
         "Convergence pricing attack: internet + unlimited mobile at/below T-Mobile Home Internet, FWA-exposed zones only.",
         "Bundle discount < churn-avoided LTV. Needs MVNO wholesale cost/line (pptx open item).",
         "Churn delta, treated vs untreated FWA zones, 2 quarters, pre-registered.",
         "No churn separation after 2 quarters, or bundle margin negative after CBRS-offload.",
         "Value flanker tier — or concede the segment and route budget to S2/S6."),
        ("S2",
         "Losses skew to AT&T/Verizon/Frontier fiber DMAs; higher-ARPU; exit cites speed / reliability / upload.",
         "Re-sequence DOCSIS 4.0 / high-split by competitive exposure (not construction cost) + symmetric win-back.",
         "Re-sequencing premium vs LTV at risk. End-2027 completion is the binding constraint to test — not the closer.",
         "Churn in evolution-complete vs incomplete overbuilt zones (cleanest natural experiment).",
         "Parity product, continued losses — completed zones bleed at the same rate.",
         "Harvest, don't chase: stop defending unwinnable zones at full cost."),
        ("S3",
         "Losses map to low-density census blocks where Starlink (Kuiper next) is the only alternative.",
         "Accelerate rural subsidized build activation — penetrate new passings before satellite entrenches.",
         "Program already funded; variable is activation pace vs BEAD/RDOF milestone penalties.",
         "Penetration on new rural passings vs vintage cohorts. Rural CR already +41k Q1 / +47k Q2 (wrong unit for Internet).",
         "Penetration on new passings falls below the underwriting curve.",
         "Slow non-committed rural extensions; subsidy-committed builds continue."),
        ("S4",
         "Gross adds down, competitor churn flat; disconnects cluster on move-outs; losses spread across exposure classes.",
         "No attacker to fight. Mover capture: Community Solutions / MDU / instant-on self-install.",
         "Mover-capture cost per connect vs standard acquisition; MDU deals building-by-building.",
         "Connect share of new move-ins; gross-add recovery as housing turnover normalizes.",
         "Port-out destination data contradicts the macro story -> re-route to S1-S2.",
         "Disclose losses by cohort so the multiple stops pricing macro softness as share loss."),
        ("S5",
         "Losses skew to non-pay / downgrades; ACP-echo demographics. FY2025 60+ day delinquencies improved (~82k vs ~103k) — size before funding.",
         "Retention-priced value tier with credit-friendly terms at the point of delinquency.",
         "Save-rate x retained margin vs discount cost + bad-debt exposure.",
         "Save-rate on treated delinquent accounts vs control; 6-month survival.",
         "Save-rate economics below LTV — then let it churn.",
         "None — self-limiting."),
        ("S6",
         "No cause-cohort exceeds ~25-30% of losses. Plausibly the true state of the world.",
         "Change-what-the-market-prices track: re-segment disclosure (rural vs core), post-2027 capex guidance, plus spend-capped experiments from S1-S3.",
         "Disclosure costs ~nothing; experiments are spend-capped.",
         "Multiple response to disclosure; experiment read-outs feed the next cycle.",
         "A later quarter crosses a threshold -> that scenario fires. Null is a state, not a verdict.",
         "Not applicable — S6 is the floor."),
    ]
    for i, row in enumerate(plays):
        r = 6 + i
        for c, val in enumerate(row, 1):
            put(ws, r, c, val, kind="note" if c > 1 else "label", bold=(c == 1))
        ws.row_dimensions[r].height = 64

    put(ws, 13, 1, "GAP overlay (does not replace S1-S6)", kind="sec")
    merge(ws, 13, 1, 13, 7)
    put(ws, 14, 1, "R1-R4", kind="label", bold=True)
    put(ws, 14, 2,
        "If lost to move-out -> TOS/MDU (6 months). FWA -> 12-month rate-lock on and off overlap (hand on-overlap FWA to fiber if it is a stepping-stone). "
        "Fiber -> overlap-zip gig/symmetry; node-split only if the sample names speed. Nonpay -> payment arrangement / 90-day first-bill save. "
        "WHERE: fire on the core; do not add rural miles. WHO (mix/tenure) sizes the offer — it is not a fifth spend program. "
        "Funding: named pools P1-P4 only. Never two expensive levers in the same quarter.",
        kind="note")
    merge(ws, 14, 2, 14, 7)
    ws.row_dimensions[14].height = 56
    put(ws, 16, 1, "Timeline", kind="label", bold=True)
    put(ws, 16, 2,
        "PPTX slide 10: Q+1 sensor live · Q+2 first attribution read · Q+3 plays trigger · Q+4-5 proof or kill. "
        "six-recommendation-areas: G1 question this week; first fill in one quarter; reason-play 2-4 quarters after fill.",
        kind="note")
    merge(ws, 16, 2, 16, 7)
    put(ws, 18, 1, "Open items (pptx slide 12)", kind="label", bold=True)
    put(ws, 18, 2,
        "MVNO wholesale cost per line (gates S1) · any existing internal churn split by exposure class "
        "(accelerates Q+2 to Q+1) · whether 4.0-4.5x leverage and the buyback program are board constants (shapes S6). "
        "None of those are prefilled as fact.",
        kind="note")
    merge(ws, 18, 2, 18, 7)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    path = build()
    print(f"Wrote {path} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
