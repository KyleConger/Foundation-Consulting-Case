from pathlib import Path
from openpyxl import load_workbook

p = Path(r"c:\Users\Owner\Desktop\MAN6930 Case\Decisions\CHURN\CHTR Rec1 Attribution Dashboard.xlsx")
print("size", p.stat().st_size, "mtime", p.stat().st_mtime)


def balance(s: str) -> int:
    n = 0
    ins = False
    for ch in s:
        if ch == '"':
            ins = not ins
        elif not ins:
            if ch == "(":
                n += 1
            elif ch == ")":
                n -= 1
    return n


wb = load_workbook(p, data_only=False)
print("sheets", wb.sheetnames)
print("names:")
for name in wb.defined_names:
    print(" ", name, "=", wb.defined_names[name].attr_text)
ws = wb["Dashboard"]
print("freeze", ws.freeze_panes, "print", ws.print_area, "titles", ws.print_title_rows)
print("protect", ws.protection.sheet, "charts", len(ws._charts), "merges", len(list(ws.merged_cells.ranges)))
print("hidden", [r for r, d in ws.row_dimensions.items() if d.hidden])
print("fitH", ws.page_setup.fitToHeight, "fitW", ws.page_setup.fitToWidth)
print("--- unbal ---")
for name in wb.sheetnames:
    for row in wb[name].iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and balance(v) != 0:
                print("UNBAL", name, cell.coordinate, balance(v), v[:240])
print("--- dash keys ---")
for addr in [
    "B1", "B2", "B3", "B6", "E6", "H6", "K6", "B7", "C8", "C11", "C12",
    "C13", "C14", "C15", "C16", "B18", "B29", "K51", "K50",
]:
    print(addr, ws[addr].value)
print("used max", ws.max_row, ws.max_column)
