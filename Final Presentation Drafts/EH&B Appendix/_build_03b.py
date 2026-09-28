"""Appendix slide: brokerage concentration of Form 5500 compensation.

Numbers come only from Decisions/EH&B/MASTER.xlsx sheet BrkDn-M2
(PlanOptica filing-name shares and the Business Insurance revenue footnote).
Does not open PowerPoint and does not touch the final deck.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "Decisions" / "EH&B" / "MASTER.xlsx"
OUT = Path(__file__).resolve().parent / "03b Brokerage Concentration.pptx"

SLIDE_W = 12192000  # 13.333 in
SLIDE_H = 6858000  # 7.5 in
EMU_IN = 914400

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x00, 0x72, 0xCE)
BODY = RGBColor(0x5B, 0x66, 0x72)
MUTED = RGBColor(0x8A, 0x94, 0xA0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD7, 0xE0, 0xEA)
SOFT = RGBColor(0xF2, 0xF6, 0xFB)
TRACK = RGBColor(0xE6, 0xEC, 0xF2)
GRAY_BAR = RGBColor(0xB7, 0xC3, 0xCE)

PLANOPTICA_URL = "https://planoptica.com/welfare/broker"
FOOTER = (
    "Final Presentation  ·  “Changing the Trajectory”  ·  Charter Communications (CHTR)"
    "  ·  Anchored to Q1 2026; Q2 2026 labeled as later public fact where used"
)

# Published filing-name shares, tenths of a percent, PlanOptica order.
# Confirmed on BrkDn-M2 note B2.
NAMES = (
    ("Marsh & McLennan Agency", 71),
    ("Gallagher Benefit", 59),
    ("Lockton", 58),
    ("Mercer Health", 50),
    ("USI", 43),
    ("WTW US", 43),
    ("Hub Midwest", 41),
    ("Aon Consulting", 28),
    ("Alliant", 25),
    ("Brown & Brown", 22),
)
TOP10_TENTHS = 440  # 44.0%
TAIL_TENTHS = 560  # 56.0%
BROKER_NAMES = 43129
TAIL_NAMES = BROKER_NAMES - 10  # 43,119
# B5: MMA + Mercer + Gallagher + Lockton + WTW + Aon + NFP 1.6
NAMED_ROLLUP_TENTHS = 71 + 50 + 59 + 58 + 43 + 28 + 16  # 32.5%


def inches(emu: int) -> float:
    return emu / EMU_IN


def emu(inches_val: float) -> int:
    return int(round(inches_val * EMU_IN))


def text_width(text: str, pt: float, bold: bool) -> float:
    path = r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"
    font = ImageFont.truetype(path, pt * 96 / 72)
    return font.getlength(text) / 96


def _set_run(run, size, bold, color, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False


def add_text(slide, l, t, w, h, text, size, bold, color, align="left", anchor="t"):
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", {"t": "t", "ctr": "ctr", "b": "b"}[anchor])
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    _set_run(run, size, bold, color)
    return box


def add_runs(slide, l, t, w, h, paragraphs, anchor="t"):
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", {"t": "t", "ctr": "ctr", "b": "b"}[anchor])
    for i, (align, runs) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
        p.line_spacing = 1.0
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for text, size, bold, color in runs:
            run = p.add_run()
            run.text = text
            _set_run(run, size, bold, color)
    return box


def _no_shadow(shape):
    sp_pr = shape._element.spPr
    effect = sp_pr.find(qn("a:effectLst"))
    if effect is not None:
        sp_pr.remove(effect)


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE
    sh.line.width = Emu(12700)
    _no_shadow(sh)
    return sh


def add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    _no_shadow(sh)
    return sh


def _cell(row, idx):
    return "" if row[idx] is None else str(row[idx])


def verify_master() -> None:
    """Fail the build if the slide numbers drift from BrkDn-M2."""
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    ws = wb["BrkDn-M2"]
    by_id = {}
    for row in ws.iter_rows(values_only=True):
        if row[1]:
            by_id[str(row[1])] = row

    b2 = _cell(by_id["B2"], 3) + " " + _cell(by_id["B2"], 7)
    b3 = _cell(by_id["B3"], 3) + " " + _cell(by_id["B3"], 7)
    b5 = _cell(by_id["B5"], 3) + " " + _cell(by_id["B5"], 7)
    a5 = _cell(by_id["A5"], 3)

    assert "44.0" in b2
    assert "56.0" in b3
    assert "43,129" in b3
    assert "32.5" in b5
    assert "NFP 1.6" in b5 or "1.6" in b5
    assert "67" in a5
    for name, tenths in NAMES:
        token = f"{tenths / 10:.1f}"
        assert token in b2, token
    assert sum(t for _, t in NAMES) == TOP10_TENTHS
    assert TOP10_TENTHS + TAIL_TENTHS == 1000
    assert TAIL_NAMES == 43119
    assert NAMED_ROLLUP_TENTHS == 325
    assert BROKER_NAMES == 43129


def pct_label(tenths: int, decimals: bool) -> str:
    value = tenths / 10
    return f"{value:.1f}%" if decimals else f"{value:.0f}%"


def build() -> None:
    verify_master()

    margin_l = emu(0.55)
    content_w = emu(12.23)
    content_right = margin_l + content_w
    assert text_width(
        "The top 10 EH&B filing names hold 44% of Form 5500 compensation;", 22, True
    ) < 12.15
    assert text_width("a 43,119-name tail holds the rest.", 22, True) < 12.15

    subtitle = (
        "This measures compensation, not covered lives, so it is not literal ownership of the book."
    )
    rollup = (
        "A different basket: Marsh (MMA + Mercer), Gallagher, Lockton, WTW, Aon, and NFP "
        "sum to 32.5% of the same table (NFP is 1.6, outside the top 10)."
    )
    bi_line = (
        "A different metric: Business Insurance 2025 ranking of 2024 U.S. brokerage revenue "
        "— top 10 = 67% of the Top 100 (67.2% unrounded). All-lines revenue, not this table."
    )
    source_1 = (
        "Source: PlanOptica health-and-welfare broker table — Form 5500 compensation "
        "filed by plan sponsors with DOL/EBSA. " + PLANOPTICA_URL
    )
    source_2 = (
        "PlanOptica is a for-profit indexer. 43,129 filing names on the table; "
        "the other 43,119 hold 56%. Hub Midwest and WTW US are filing names, not full parents."
    )
    scale_note = "Bars start at zero. Scale is 0–100% of this compensation table."
    for text, pt, bold, limit in (
        (subtitle, 14, False, 12.15),
        (rollup, 12, False, 12.15),
        (bi_line, 12, False, 12.15),
        (source_1, 11, False, 12.15),
        (source_2, 11, False, 12.15),
        (scale_note, 11, False, 6.30),
        (FOOTER, 9, False, 11.55),
    ):
        width = text_width(text, pt, bold)
        if width >= limit:
            raise SystemExit(f"Line too wide ({width:.2f} in): {text}")

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    prs.core_properties.title = "EH&B brokerage concentration"
    blank = next(layout for layout in prs.slide_layouts if layout.name == "Blank")
    slide = prs.slides.add_slide(blank)
    for placeholder in list(slide.placeholders):
        placeholder._element.getparent().remove(placeholder._element)

    add_text(
        slide,
        margin_l,
        emu(0.32),
        content_w,
        emu(0.28),
        "APPENDIX  ·  EH&B — BROKERAGE CONCENTRATION",
        11,
        True,
        BLUE,
        anchor="b",
    )
    add_runs(
        slide,
        margin_l,
        emu(0.62),
        content_w,
        emu(0.78),
        [
            (
                "left",
                [
                    ("The top 10 EH&B filing names hold ", 22, True, NAVY),
                    ("44%", 22, True, BLUE),
                    (" of Form 5500 compensation;", 22, True, NAVY),
                ],
            ),
            (
                "left",
                [("a 43,119-name tail holds the rest.", 22, True, NAVY)],
            ),
        ],
    )
    add_text(slide, margin_l, emu(1.42), content_w, emu(0.32), subtitle, 14, False, BODY, anchor="t")

    # Left: top 10 vs tail, one scale from zero to 100.
    plot_top = emu(1.98)
    add_text(
        slide,
        margin_l,
        plot_top,
        emu(6.40),
        emu(0.26),
        "SHARE OF FORM 5500 COMPENSATION",
        11,
        True,
        BLUE,
        anchor="b",
    )

    label_w = emu(2.05)
    bar_x = margin_l + label_w + emu(0.12)
    bar_max = emu(3.85)
    assert bar_x + bar_max < emu(6.85)
    rows = (
        ("Top 10 filing names", TOP10_TENTHS, BLUE, BLUE),
        ("Other 43,119 names", TAIL_TENTHS, GRAY_BAR, NAVY),
    )
    row_h = emu(1.22)
    bar_h = emu(0.50)
    first_bar_top = None
    last_bar_bottom = None
    y = emu(2.40)
    for label, tenths, bar_color, text_color in rows:
        add_text(
            slide,
            margin_l,
            y,
            label_w,
            row_h,
            label,
            14,
            True,
            text_color,
            align="right",
            anchor="ctr",
        )
        bar_top = y + (row_h - bar_h) // 2
        add_rect(slide, bar_x, bar_top, bar_max, bar_h, TRACK)
        bar_w = int(round(bar_max * tenths / 1000))
        add_rect(slide, bar_x, bar_top, bar_w, bar_h, bar_color)
        label = pct_label(tenths, decimals=False)
        label_end = bar_x + bar_w + emu(0.08) + emu(text_width(label, 18, True))
        if label_end > emu(7.00):
            raise SystemExit(f"Bar label crosses into the name card: {label}")
        add_text(
            slide,
            bar_x + bar_w + emu(0.08),
            bar_top,
            emu(0.85),
            bar_h,
            label,
            18,
            True,
            text_color,
            anchor="ctr",
        )
        first_bar_top = bar_top if first_bar_top is None else first_bar_top
        last_bar_bottom = bar_top + bar_h
        y += row_h

    axis_y = y + emu(0.06)
    add_rect(slide, bar_x, first_bar_top, emu(0.015), axis_y - first_bar_top, LINE)
    add_rect(slide, bar_x, axis_y, bar_max, emu(0.015), LINE)
    for tick in (0, 25, 50, 75, 100):
        tx = bar_x + int(round(bar_max * tick / 100))
        add_rect(slide, tx, axis_y, emu(0.012), emu(0.08), LINE)
        add_text(
            slide,
            tx - emu(0.28),
            axis_y + emu(0.08),
            emu(0.56),
            emu(0.22),
            str(tick),
            10,
            False,
            MUTED,
            align="center",
        )
    add_text(
        slide,
        margin_l,
        axis_y + emu(0.32),
        emu(6.40),
        emu(0.24),
        scale_note,
        11,
        False,
        MUTED,
        anchor="t",
    )

    # Right: the ten names that add to 44%. Same unit, no second bar scale.
    card_x = emu(7.15)
    card_y = emu(1.98)
    card_w = content_right - card_x
    card_h = emu(3.78)
    add_round(slide, card_x, card_y, card_w, card_h, SOFT)
    inset = emu(0.20)
    inner_x = card_x + inset
    inner_w = card_w - 2 * inset
    add_text(
        slide,
        inner_x,
        card_y + emu(0.16),
        inner_w,
        emu(0.24),
        "INSIDE THE 44%",
        11,
        True,
        BLUE,
        anchor="b",
    )
    add_text(
        slide,
        inner_x,
        card_y + emu(0.40),
        inner_w,
        emu(0.22),
        "Filing-name share, descending",
        12,
        False,
        BODY,
        anchor="t",
    )

    name_top = card_y + emu(0.68)
    row_h_name = emu(0.248)
    pct_w = emu(0.72)
    name_w = inner_w - pct_w
    longest = max(text_width(name, 13, False) for name, _ in NAMES)
    if longest >= inches(name_w) - 0.05:
        raise SystemExit(f"Name column too narrow ({longest:.2f} in)")
    for i, (name, tenths) in enumerate(NAMES):
        yy = name_top + i * row_h_name
        add_text(slide, inner_x, yy, name_w, row_h_name, name, 13, False, NAVY, anchor="ctr")
        add_text(
            slide,
            inner_x + name_w,
            yy,
            pct_w,
            row_h_name,
            pct_label(tenths, decimals=True),
            13,
            True,
            NAVY,
            align="right",
            anchor="ctr",
        )

    sum_y = name_top + len(NAMES) * row_h_name + emu(0.06)
    add_rect(slide, inner_x, sum_y, inner_w, emu(0.012), LINE)
    add_text(
        slide,
        inner_x,
        sum_y + emu(0.06),
        name_w,
        emu(0.28),
        "Ten filing names",
        13,
        True,
        NAVY,
        anchor="ctr",
    )
    add_text(
        slide,
        inner_x + name_w,
        sum_y + emu(0.06),
        pct_w,
        emu(0.28),
        "44.0%",
        13,
        True,
        BLUE,
        align="right",
        anchor="ctr",
    )

    add_text(slide, margin_l, emu(5.90), content_w, emu(0.26), rollup, 12, False, BODY, anchor="t")
    add_text(slide, margin_l, emu(6.20), content_w, emu(0.26), bi_line, 12, False, MUTED, anchor="t")
    add_text(slide, margin_l, emu(6.52), content_w, emu(0.22), source_1, 11, False, MUTED, anchor="t")
    add_text(slide, margin_l, emu(6.74), content_w, emu(0.22), source_2, 11, False, MUTED, anchor="t")
    add_text(slide, margin_l, emu(7.08), emu(11.55), emu(0.28), FOOTER, 9, False, MUTED, anchor="b")
    add_text(
        slide,
        emu(12.35),
        emu(7.08),
        emu(0.45),
        emu(0.28),
        "1",
        10,
        False,
        MUTED,
        align="right",
        anchor="b",
    )

    # The sum row must sit inside the card.
    if sum_y + emu(0.36) > card_y + card_h - emu(0.08):
        raise SystemExit("Sum row overflows the name card")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
