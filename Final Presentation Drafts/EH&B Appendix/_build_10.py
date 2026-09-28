# -*- coding: utf-8 -*-
"""Appendix slide: Motley Fool and Simply Wall St perceptions of Charter.

Facts come only from the sell-side coverage canvas
(chtr-sellside-coverage.canvas.tsx, retrieved 21 Sep 2026).
Analysis/sellside-coverage/covering_analysts.csv has no row for either source.

Does not open PowerPoint and does not touch the final deck.
"""
from __future__ import annotations

from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLIDE_W = 12192000  # 13.333 in
SLIDE_H = 6858000  # 7.5 in
EMU_IN = 914400

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "10 Motley Fool and Simply Wall St.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x00, 0x72, 0xCE)
BODY = RGBColor(0x5B, 0x66, 0x72)
MUTED = RGBColor(0x8A, 0x94, 0xA0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD7, 0xE0, 0xEA)
SOFT = RGBColor(0xF4, 0xF7, 0xFA)
BLUE_SOFT = RGBColor(0xE8, 0xF3, 0xFB)

FOOTER = (
    "Final Presentation  ·  “Changing the Trajectory”  ·  "
    "Charter Communications (CHTR)  ·  Appendix"
)

# Canvas URLs, retrieved 21 Sep 2026. Paraphrases only; no article body text.
FOOL = (
    {
        "who": "Keith Noonan",
        "date": "24 Apr 2026",
        "stance": "Wait",
        "view": (
            "After the Q1 miss, wait for turnaround signs before buying "
            "the cheaper print."
        ),
        "url": "https://www.fool.com/investing/2026/04/24/why-charter-communications-stock-plummeted-today/",
    },
    {
        "who": "James Brumley",
        "date": "7 May 2026",
        "stance": "Fixed wireless",
        "view": (
            "Broadband no longer offsets video losses. Fixed wireless at "
            "T-Mobile and Verizon is the leak, and the migration can run a long time."
        ),
        "url": "https://www.fool.com/investing/2026/05/07/where-are-comcast-and-charters-internet-customers/",
    },
    {
        "who": "Brett Schafer",
        "date": "2 Jul 2026",
        "stance": "Buy the cash flow",
        "view": (
            "Would buy Charter over CoreWeave: wired broadband will not vanish "
            "overnight, and Charter is profitable at a cheaper valuation than an unprofitable AI host."
        ),
        "url": "https://www.fool.com/investing/2026/07/02/coreweave-is-about-to-join-the-nasdaq-100-heres-wh/",
    },
)

SWS = (
    {
        "who": "CHTR stock report",
        "date": "Retrieved 21 Sep 2026",
        "view": (
            "Screens cheap (valuation 5/6) against financial health 0/6; "
            "interest is not well covered. Recycles consensus, low, and high "
            "narratives and names no covering analyst."
        ),
        "url": "https://simplywall.st/stocks/us/media/nasdaq-chtr/charter-communications",
    },
    {
        "who": "Margin-pressure recap",
        "date": "26 Jul 2026",
        "view": (
            "Byline is Simply Wall St. Trailing net margin 9.1% versus 9.5% "
            "a year earlier, and a roughly 3× earnings multiple against a much "
            "larger DCF. Bulls on cheapness, bears on leverage."
        ),
        "url": "https://simplywall.st/stocks/us/media/nasdaq-chtr/charter-communications/news/charter-communications-chtr-stock-faces-margin-pressure-as-q",
    },
    {
        "who": "CFO-exit recap",
        "date": "8 Sep 2026",
        "view": (
            "Reviewed by Sasha Jovanovic, Simply Wall St staff, not a covering "
            "analyst. The CFO exit is framed as operational; the swing factor "
            "remains about $94 billion of debt and interest cover."
        ),
        "url": "https://simplywall.st/stocks/us/media/nasdaq-chtr/charter-communications/news/did-cfo-exit-just-shift-charter-communications-chtr-investme",
    },
)

TITLE = (
    "Outside the banks, Fool writers split between waiting out the broadband "
    "miss and buying the cash flow; Simply Wall St recycles cheap stock versus leverage"
)

BANNED = (
    "Moffett",
    "Harlalka",
    "Cahall",
    "Piecyk",
    "Diffley",
    "Hodulik",
    "Petti",
    "Swinburne",
    "Ehrlich",
    "Reif",
    "Michael Ng",
    "Rollins",
    "Louthan",
    "Goldman",
    "New Street",
    "Wells Fargo",
    "LightShed",
    "Morgan Stanley",
    "JPMorgan",
    "Citigroup",
    "Raymond James",
)


def emu(inches_val: float) -> int:
    return int(round(inches_val * EMU_IN))


def _font(pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    path = r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"
    return ImageFont.truetype(path, pt * 96 / 72)


def wrap_lines(text: str, width_in: float, pt: float, bold: bool) -> list[str]:
    font = _font(pt, bold)
    words = text.split()
    lines: list[str] = []
    cur = ""
    for word in words:
        trial = word if not cur else cur + " " + word
        if font.getlength(trial) / 96 <= width_in:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def _set_run(run, size, bold, color, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.underline = False
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag))
            rPr.append(el)
        el.set("typeface", "Calibri")


def _body(tf, anchor: str):
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", anchor)


def add_text(slide, l, t, w, h, paragraphs, anchor="t"):
    """paragraphs: list of (align, [(text, size, bold, color, italic, url|None), ...])."""
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    _body(tf, anchor)
    for i, (align, runs) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {
            "left": PP_ALIGN.LEFT,
            "right": PP_ALIGN.RIGHT,
            "center": PP_ALIGN.CENTER,
        }[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        pPr = p._p.get_or_add_pPr()
        if pPr.find(qn("a:buNone")) is None:
            pPr.append(pPr.makeelement(qn("a:buNone")))
        for text, size, bold, color, italic, url in runs:
            run = p.add_run()
            run.text = text
            _set_run(run, size, bold, color, italic)
            if url:
                run.hyperlink.address = url
                run.font.color.rgb = color
                run.font.underline = False
    return box


def _no_line(shape):
    shape.line.fill.background()
    sp_pr = shape._element.spPr
    effect = sp_pr.find(qn("a:effectLst"))
    if effect is not None:
        sp_pr.remove(effect)


def add_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        _no_line(sh)
    else:
        sh.line.color.rgb = line
        sh.line.width = Emu(12700)
        sp_pr = sh._element.spPr
        effect = sp_pr.find(qn("a:effectLst"))
        if effect is not None:
            sp_pr.remove(effect)
    return sh


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h)
    )
    try:
        sh.adjustments[0] = 0.06
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE
    sh.line.width = Emu(9525)
    sp_pr = sh._element.spPr
    effect = sp_pr.find(qn("a:effectLst"))
    if effect is not None:
        sp_pr.remove(effect)
    return sh


def run(text, size, bold, color, italic=False, url=None):
    return (text, size, bold, color, italic, url)


def build():
    title_lines = wrap_lines(TITLE, 12.45, 18, True)
    assert len(title_lines) <= 2, title_lines
    for item in FOOL + SWS:
        view_lines = wrap_lines(item["view"], 5.72, 12, False)
        assert len(view_lines) <= 3, (item["who"], view_lines)

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    prs.core_properties.title = "Motley Fool and Simply Wall St — appendix"
    prs.core_properties.subject = (
        "Charter perceptions outside sell-side coverage. "
        "Canvas retrieved 21 Sep 2026."
    )
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_text(
        slide,
        emu(0.42),
        emu(0.22),
        emu(12.45),
        emu(0.24),
        [("left", [run("APPENDIX  ·  PERCEPTIONS OUTSIDE THE BANKS", 11, True, BLUE)])],
        anchor="ctr",
    )
    add_text(
        slide,
        emu(0.42),
        emu(0.46),
        emu(12.45),
        emu(0.78),
        [
            (
                "left",
                [
                    run(
                        "Outside the banks, Fool writers split between waiting out the broadband miss and buying the cash flow",
                        18,
                        True,
                        NAVY,
                    )
                ],
            ),
            (
                "left",
                [
                    run(
                        "Simply Wall St recycles cheap stock versus leverage",
                        18,
                        True,
                        NAVY,
                    )
                ],
            ),
        ],
        anchor="t",
    )

    col_y = 1.32
    col_h = 4.58
    col_w = 6.16
    left_x = 0.42
    right_x = 6.74

    add_rect(slide, emu(left_x), emu(col_y), emu(col_w), emu(col_h), WHITE, LINE)
    add_rect(slide, emu(right_x), emu(col_y), emu(col_w), emu(col_h), SOFT, LINE)
    # One accent: the Fool bucket. Simply Wall St stays gray.
    add_rect(slide, emu(left_x), emu(col_y), emu(0.07), emu(col_h), BLUE)

    add_text(
        slide,
        emu(left_x + 0.22),
        emu(col_y + 0.16),
        emu(col_w - 0.44),
        emu(0.66),
        [
            ("left", [run("Motley Fool", 16, True, NAVY)]),
            (
                "left",
                [
                    run(
                        "Its own shop. These writers are not a covering desk. A Fool-hosted earnings transcript is Charter’s call, not a Fool view.",
                        11,
                        False,
                        BODY,
                    )
                ],
            ),
        ],
    )
    add_text(
        slide,
        emu(right_x + 0.22),
        emu(col_y + 0.16),
        emu(col_w - 0.44),
        emu(0.66),
        [
            ("left", [run("Simply Wall St", 16, True, NAVY)]),
            (
                "left",
                [
                    run(
                        "Aggregator recap of media and consensus. Not original research and not a covering desk.",
                        11,
                        False,
                        BODY,
                    )
                ],
            ),
        ],
    )

    def items(x, rows, stance: bool):
        top = col_y + 0.92
        block_h = 1.16
        for i, row in enumerate(rows):
            y = top + i * block_h
            if i:
                add_rect(
                    slide,
                    emu(x + 0.22),
                    emu(y),
                    emu(col_w - 0.44),
                    emu(0.01),
                    LINE,
                )
            head = [run(row["who"], 13, True, NAVY), run("   " + row["date"], 12, False, MUTED)]
            if stance:
                head.append(run("   ·   ", 12, False, MUTED))
                head.append(run(row["stance"], 12, True, BLUE))
            add_text(
                slide,
                emu(x + 0.22),
                emu(y + 0.08),
                emu(col_w - 0.44),
                emu(0.28),
                [("left", head)],
                anchor="ctr",
            )
            add_text(
                slide,
                emu(x + 0.22),
                emu(y + 0.38),
                emu(col_w - 0.44),
                emu(0.72),
                [("left", [run(row["view"], 12, False, BODY)])],
            )

    items(left_x, FOOL, stance=True)
    items(right_x, SWS, stance=False)

    # Source: canvas plus the public URLs the canvas lists. Full width so
    # the longest Simply Wall St links stay on the slide.
    src_y = 6.00
    url_rows = [
        (
            "left",
            [
                run(
                    "Source: chtr-sellside-coverage canvas, retrieved 21 Sep 2026. No row in Analysis/sellside-coverage/covering_analysts.csv. Sasha Jovanovic is Simply Wall St staff, not a sell-side analyst.",
                    8,
                    False,
                    MUTED,
                    True,
                )
            ],
        )
    ]
    for item in (*FOOL, *SWS):
        url_rows.append(("left", [run(item["url"], 8, False, MUTED, False, item["url"])]))
    source_box = add_text(
        slide,
        emu(0.42),
        emu(src_y),
        emu(12.48),
        emu(1.06),
        url_rows,
    )
    for p in source_box.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        ln = pPr.makeelement(qn("a:lnSpc"))
        spc = ln.makeelement(qn("a:spcPts"))
        spc.set("val", "1050")
        ln.append(spc)
        pPr.append(ln)
    add_text(
        slide,
        emu(0.42),
        emu(7.12),
        emu(12.45),
        emu(0.26),
        [("left", [run(FOOTER, 9, False, MUTED)])],
        anchor="ctr",
    )

    source_intro = (
        "Source: chtr-sellside-coverage canvas, retrieved 21 Sep 2026. "
        "No row in Analysis/sellside-coverage/covering_analysts.csv. "
        "Sasha Jovanovic is Simply Wall St staff, not a sell-side analyst."
    )
    assert len(wrap_lines(source_intro, 12.48, 8, False)) <= 2
    for item in (*FOOL, *SWS):
        width = _font(8, False).getlength(item["url"]) / 96
        assert width <= 12.48, (item["url"], width)

    text_blob = TITLE + " " + " ".join(r["who"] + r["view"] for r in FOOL + SWS)
    for name in BANNED:
        assert name not in text_blob, name

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
