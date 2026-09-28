"""Appendix slides: small vs large employment shares.

Numbers come only from Decisions/EH&B/MASTER.xlsx sheets Emp-US, Emp-States,
and Emp-MetrosFull (median note on that sheet). Does not open the final deck.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "Decisions" / "EH&B" / "MASTER.xlsx"
OUT = Path(__file__).resolve().parent / "01 Small Large Employment Shares.pptx"

SLIDE_W = 12192000
SLIDE_H = 6858000
MARGIN_L = 502920
TITLE_W = 11155680
KICKER_T = 292608
KICKER_H = 274320
TITLE_T = 566928
TITLE_H = 1005840
SOURCE_T = 6144768
SOURCE_H = 274320
FOOT_T = 6446520
FOOT_H = 274320
PAGE_L = 11430000
PAGE_W = 594360
INSET = 182880

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x00, 0x72, 0xCE)
BODY = RGBColor(0x5B, 0x66, 0x72)
MUTED = RGBColor(0x8A, 0x94, 0xA0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD7, 0xE0, 0xEA)
SOFT = RGBColor(0xF2, 0xF6, 0xFB)
GRAY_BAR = RGBColor(0xB7, 0xC3, 0xCE)

URL = "https://www.census.gov/data/datasets/2022/econ/susb/2022-susb.html"
FOOTER = (
    "Final Presentation  ·  “Changing the Trajectory”  ·  Charter Communications (CHTR)"
    "  ·  Anchored to Q1 2026; Q2 2026 labeled as later public fact where used"
)


def _set_run(run, size, bold, color, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False


def add_text(slide, l, t, w, h, text, size, bold, color, align="left", anchor="t"):
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", {"t": "t", "ctr": "ctr", "b": "b"}[anchor])
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    _set_run(run, size, bold, color)
    return box


def add_runs(slide, l, t, w, h, paragraphs, anchor="t"):
    """paragraphs: list of (align, list of (text, size, bold, color))."""
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", {"t": "t", "ctr": "ctr", "b": "b"}[anchor])
    for i, (align, runs) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for text, size, bold, color in runs:
            run = p.add_run()
            run.text = text
            _set_run(run, size, bold, color)
    return box


def _no_shadow(shape):
    sp_pr = shape._element.spPr
    effect = sp_pr.find(qn("a:effectLst"))
    if effect is not None:
        sp_pr.remove(effect)


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    try:
        sh.adjustments[0] = 0.05714
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE
    sh.line.width = Emu(12700)
    _no_shadow(sh)
    return sh


def add_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Emu(6350)
    _no_shadow(sh)
    return sh


def chrome(slide, kicker, title, source, page):
    add_text(slide, MARGIN_L, KICKER_T, TITLE_W, KICKER_H, kicker, 11, True, BLUE, anchor="b")
    add_text(slide, MARGIN_L, TITLE_T, TITLE_W, TITLE_H, title, 22, True, NAVY, anchor="t")
    add_text(slide, MARGIN_L, SOURCE_T, TITLE_W, SOURCE_H, source, 9, False, MUTED, anchor="t")
    add_text(slide, MARGIN_L, FOOT_T, 10881360, FOOT_H, FOOTER, 9, False, MUTED, anchor="b")
    add_text(slide, PAGE_L, FOOT_T, PAGE_W, FOOT_H, page, 10, False, MUTED, align="right", anchor="b")


def _rows(ws):
    data = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if row[1] is None:
            continue
        data.append(row)
    return data


def load():
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    us_rows = _rows(wb["Emp-US"])
    us = us_rows[0]
    # Geography, emp <500, emp 500+, total, share <500, share 500+
    us_small_emp, us_large_emp, us_total = int(us[1]), int(us[2]), int(us[3])
    us_small, us_large = float(us[4]), float(us[5])

    states = _rows(wb["Emp-States"])
    # rank, name, emp <500, emp 500+, total, share <500, share 500+
    by_name = {r[1]: r for r in states}
    high = states[:5]
    low = states[-5:]

    metros = []
    for row in _rows(wb["Emp-MetrosFull"]):
        if isinstance(row[4], (int, float)):
            metros.append(row)
    prec = sorted(float(r[4]) for r in metros)
    median_full = prec[len(prec) // 2]
    median_1 = round(median_full, 1)

    assert len(states) == 51
    assert us_small == 45.9 and us_large == 54.1
    assert by_name["Montana"][5] == 66.3
    assert by_name["Florida"][5] == 39.6
    assert len(metros) == 387
    assert median_1 == 49.8
    span = round(float(by_name["Montana"][5]) - float(by_name["Florida"][5]), 1)
    gap = round(median_1 - us_small, 1)
    assert span == 26.7
    assert gap == 3.9
    assert us_small_emp + us_large_emp == us_total

    return {
        "us_small": us_small,
        "us_large": us_large,
        "us_small_emp": us_small_emp,
        "us_large_emp": us_large_emp,
        "us_total": us_total,
        "high": high,
        "low": low,
        "montana_base": int(by_name["Montana"][4]),
        "florida_base": int(by_name["Florida"][4]),
        "n_states": len(states),
        "n_metros": len(metros),
        "median": median_1,
        "span": span,
        "gap": gap,
    }


def comma(n):
    return f"{int(n):,}"


def pct(n):
    return f"{float(n):.1f}%"


def build(d):
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    blank = min(prs.slide_layouts, key=lambda layout: len(layout.placeholders))
    slide_us(prs, blank, d)
    slide_states(prs, blank, d)
    slide_metros(prs, blank, d)
    prs.save(OUT)


def slide_us(prs, blank, d):
    s = prs.slides.add_slide(blank)
    chrome(
        s,
        "APPENDIX  ·  EH&B  ·  EMPLOYMENT SHARES  ·  UNITED STATES",
        "Large firms employ 54.1% of U.S. payroll workers; firms under 500 hold 45.9%",
        "Source: U.S. Census Bureau, Statistics of U.S. Businesses (SUSB) 2022, U.S. all-industry enterprise size. "
        + URL,
        "A1.1",
    )
    top = 1780000
    gap = 228600
    card_w = (TITLE_W - gap) // 2
    card_h = 2743200
    left = MARGIN_L
    right = MARGIN_L + card_w + gap
    add_round(s, left, top, card_w, card_h, WHITE)
    add_round(s, right, top, card_w, card_h, SOFT)

    def card_copy(x, kicker, number, number_color, workers):
        add_text(s, x + INSET, top + 228600, card_w - 2 * INSET, 274320, kicker, 11, True, BLUE, anchor="b")
        add_text(s, x + INSET, top + 594360, card_w - 2 * INSET, 731520, number, 40, True, number_color, anchor="t")
        add_text(
            s,
            x + INSET,
            top + 1463040,
            card_w - 2 * INSET,
            365760,
            "Share of U.S. SUSB employment",
            14,
            False,
            BODY,
            anchor="t",
        )
        add_text(
            s,
            x + INSET,
            top + 1920240,
            card_w - 2 * INSET,
            411480,
            f"{comma(workers)} workers",
            16,
            True,
            NAVY,
            anchor="t",
        )

    card_copy(left, "FIRMS WITH 500 OR MORE", pct(d["us_large"]), BLUE, d["us_large_emp"])
    card_copy(right, "FIRMS UNDER 500", pct(d["us_small"]), NAVY, d["us_small_emp"])

    band_t = top + card_h + 182880
    band_h = SOURCE_T - band_t - 137160
    add_round(s, MARGIN_L, band_t, TITLE_W, band_h, SOFT)
    add_runs(
        s,
        MARGIN_L + INSET,
        band_t + 137160,
        TITLE_W - 2 * INSET,
        band_h - 228600,
        [
            (
                "left",
                [
                    ("SUSB 2022 total  ", 13, True, NAVY),
                    (f"{comma(d['us_total'])} payroll workers.  ", 13, False, BODY),
                    ("Primary cut is the SBA threshold: enterprises with fewer than 500 employees.", 13, False, BODY),
                ],
            ),
            (
                "left",
                [
                    (
                        "Mid-March payroll employment at employer establishments, classified by enterprise size, "
                        "not establishment size. Excludes nonemployers, most government, and most of agriculture.",
                        12,
                        False,
                        BODY,
                    )
                ],
            ),
        ],
        anchor="ctr",
    )


def slide_states(prs, blank, d):
    s = prs.slides.add_slide(blank)
    chrome(
        s,
        "APPENDIX  ·  EH&B  ·  EMPLOYMENT SHARES  ·  STATES",
        f"Small-firm share spans {d['span']:.1f} points, from 66.3% in Montana to 39.6% in Florida",
        "Source: U.S. Census Bureau, SUSB 2022, 50 states and D.C. Highest and lowest 5 of 51. "
        f"Montana base {comma(d['montana_base'])} jobs; Florida base {comma(d['florida_base'])}. "
        + URL,
        "A1.2",
    )

    label_w = 2011680
    bar_x = MARGIN_L + label_w + 91440
    right = MARGIN_L + TITLE_W
    jobs_w = 1371600
    pct_w = 640080
    jobs_x = right - jobs_w
    pct_x = jobs_x - pct_w - 45720
    bar_max = pct_x - bar_x - 91440  # 0–100 scale; bars start at zero
    assert bar_max > 4000000

    y = 1680000
    add_text(s, bar_x, y, bar_max, 182880, "Share at firms under 500", 10, True, MUTED, anchor="b")
    add_text(s, jobs_x, y, jobs_w, 182880, "Employment", 10, True, MUTED, align="right", anchor="b")
    y += 200000

    row_h = 283464
    bar_h = 146304
    groups = (("HIGHEST 5", d["high"]), ("LOWEST 5", d["low"]))
    bar_tops = []
    bar_bottoms = []
    line_x = bar_x + int(round(bar_max * d["us_small"] / 100.0))

    for gi, (label, rows) in enumerate(groups):
        add_text(s, MARGIN_L, y, label_w, 182880, label, 10, True, BLUE, align="right", anchor="b")
        y += 192024
        for row in rows:
            name = row[1]
            share = float(row[5])
            jobs = int(row[4])
            accent = name in ("Montana", "Florida")
            color = BLUE if accent else NAVY
            bar_color = BLUE if accent else GRAY_BAR
            add_text(
                s, MARGIN_L, y, label_w, row_h, name, 13, accent, color, align="right", anchor="ctr"
            )
            bar_w = int(round(bar_max * share / 100.0))
            bar_top = y + (row_h - bar_h) // 2
            add_rect(s, bar_x, bar_top, max(bar_w, 0), bar_h, bar_color)
            bar_tops.append(bar_top)
            bar_bottoms.append(bar_top + bar_h)
            add_text(s, pct_x, y, pct_w, row_h, pct(share), 12, accent, color, align="right", anchor="ctr")
            add_text(s, jobs_x, y, jobs_w, row_h, comma(jobs), 12, False, BODY, align="right", anchor="ctr")
            y += row_h
        if gi == 0:
            # Open gap so the U.S. callout does not sit on a bar or a group label.
            add_text(
                s,
                line_x + 73152,
                y + 27432,
                1828800,
                row_h - 27432,
                f"U.S. {pct(d['us_small'])}",
                11,
                True,
                BODY,
                anchor="ctr",
            )
            y += row_h

    add_rect(s, bar_x, bar_tops[0], 9525, bar_bottoms[-1] - bar_tops[0], LINE)
    axis_y = y + 22860
    for tick in (0, 25, 50, 75, 100):
        tx = bar_x + int(round(bar_max * tick / 100.0))
        add_text(s, tx - 228600, axis_y, 457200, 164592, str(tick), 9, False, MUTED, align="center", anchor="t")

    line = s.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Emu(line_x),
        Emu(bar_tops[0]),
        Emu(line_x),
        Emu(bar_bottoms[-1]),
    )
    line.line.color.rgb = BODY
    line.line.width = Pt(1.25)
    line.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def slide_metros(prs, blank, d):
    s = prs.slides.add_slide(blank)
    chrome(
        s,
        "APPENDIX  ·  EH&B  ·  EMPLOYMENT SHARES  ·  METROS",
        f"The median metro small-firm share is {d['median']:.1f}%, {d['gap']:.1f} points above the U.S. 45.9%",
        "Source: U.S. Census Bureau, SUSB 2022, median of 387 metropolitan statistical areas "
        "(MSA file revised 2025-07-22). "
        + URL,
        "A1.3",
    )
    top = 1780000
    gap = 228600
    card_w = (TITLE_W - gap) // 2
    card_h = 2743200
    left = MARGIN_L
    right = MARGIN_L + card_w + gap
    add_round(s, left, top, card_w, card_h, SOFT)
    add_round(s, right, top, card_w, card_h, WHITE)

    def card_copy(x, kicker, number, number_color, line2, line3):
        add_text(s, x + INSET, top + 228600, card_w - 2 * INSET, 274320, kicker, 11, True, BLUE, anchor="b")
        add_text(s, x + INSET, top + 594360, card_w - 2 * INSET, 731520, number, 40, True, number_color, anchor="t")
        add_text(s, x + INSET, top + 1463040, card_w - 2 * INSET, 365760, line2, 14, False, BODY, anchor="t")
        add_text(s, x + INSET, top + 1920240, card_w - 2 * INSET, 457200, line3, 14, True, NAVY, anchor="t")

    card_copy(
        left,
        "METRO MEDIAN",
        pct(d["median"]),
        BLUE,
        "Share of employment at firms under 500",
        f"Median of {d['n_metros']} metropolitan areas",
    )
    card_copy(
        right,
        "UNITED STATES",
        pct(d["us_small"]),
        NAVY,
        "Share of employment at firms under 500",
        "Includes employment outside metros",
    )

    band_t = top + card_h + 182880
    band_h = SOURCE_T - band_t - 137160
    add_round(s, MARGIN_L, band_t, TITLE_W, band_h, SOFT)
    add_runs(
        s,
        MARGIN_L + INSET,
        band_t + 137160,
        TITLE_W - 2 * INSET,
        band_h - 228600,
        [
            (
                "left",
                [
                    (f"{d['median']:.1f}% − {d['us_small']:.1f}% = {d['gap']:.1f} points.  ", 14, True, NAVY),
                    (
                        "Both figures are SUSB 2022 shares at enterprises with fewer than 500 employees, shown to 1 decimal.",
                        14,
                        False,
                        BODY,
                    ),
                ],
            ),
            (
                "left",
                [
                    (
                        "Micro areas are excluded. The national share is not a metro median: it includes non-metro employment.",
                        13,
                        False,
                        BODY,
                    )
                ],
            ),
        ],
        anchor="ctr",
    )


if __name__ == "__main__":
    data = load()
    build(data)
    print(OUT)
