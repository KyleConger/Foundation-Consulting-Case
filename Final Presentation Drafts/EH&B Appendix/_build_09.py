# -*- coding: utf-8 -*-
"""Two appendix slides: EH&B mindset timeline, then informal sales interviews.

Writes the pptx with python-pptx only. Does not open PowerPoint and does not
touch the final deck.
"""
from __future__ import annotations

from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "09 EH&B Mindset and Sales Interviews.pptx"

SLIDE_W = 12192000  # 13.333... in, same EMU as the other appendix files
SLIDE_H = 6858000  # 7.5 in
EMU_IN = 914400

MARGIN_L = 502920
TITLE_W = 11155680
KICKER_T = 292608
KICKER_H = 274320
TITLE_T = 566928
TITLE_H = 1005840
SOURCE_T = 6144768
SOURCE_H = 274320
FOOT_T = 6446520
FOOT_H = 274320
PAGE_L = 11430000
PAGE_W = 594360

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x00, 0x72, 0xCE)
BODY = RGBColor(0x5B, 0x66, 0x72)
MUTED = RGBColor(0x8A, 0x94, 0xA0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD7, 0xE0, 0xEA)
SOFT = RGBColor(0xF2, 0xF6, 0xFB)

FOOTER = (
    "Final Presentation  ·  “Changing the Trajectory”  ·  Charter Communications (CHTR)"
    "  ·  Anchored to Q1 2026; Q2 2026 labeled as later public fact where used"
)

TITLE_1 = (
    "Brokers moved EH&B from a cost to manage",
    "into an investment in whether employees can perform",
)
TITLE_2 = (
    "Buyers stay for a price they can see for years;",
    "Spectrum’s quality story loses to the jump after month 12",
)

# Kyle's note, four eras only. Not the six-step method and not the privacy note.
ERAS = (
    (
        "Late 1990s–early 2000s",
        "Advisers connected employee health to absenteeism, disability costs, and lost productivity, not only insurance spend.",
    ),
    (
        "Mid-to-late 2000s",
        "Wellness, disease management, and health-data analytics treated benefits as an investment in workforce performance.",
    ),
    (
        "2010s",
        "Major brokerages positioned EH&B around engagement, retention, and organizational performance. The buyer moved from benefits procurement toward HR, finance, and executives.",
    ),
    (
        "2020 onward",
        "COVID accelerated mental health, workforce resilience, and benefits as a strategic tool.",
    ),
)

# Informal interview notes. Order follows the slide claim. East Coast stays an interviewee claim.
# Runs: (text, size, bold, color).
ROWS = (
    (
        "Xfinity’s offer",
        (
            ("Xfinity offers a ", 14, False, BODY),
            ("5-year price lock, unlimited for that period.", 14, True, BLUE),
        ),
    ),
    (
        "Quality vs. the jump",
        (
            (
                "Buyers the interviewee talks to believe Spectrum’s service quality is better than Xfinity. "
                "Spectrum has fewer usage fees and mostly unlimited data. ",
                14,
                False,
                BODY,
            ),
            (
                "Little reason to stay past 12 months, because the price jumps.",
                14,
                True,
                BLUE,
            ),
        ),
    ),
    (
        "Price sensitivity",
        (
            (
                "Buyers are extremely price sensitive and also want reliability. "
                "Charter’s weakness in this account is price raises, which give Xfinity five years to raise the cost of switching.",
                14,
                False,
                BODY,
            ),
        ),
    ),
    (
        "What buyers ask for",
        (
            (
                "Longer price guarantees, fewer jumps off the promo price, a clearer monthly cost, "
                "and lower switching costs — transparency that holds up over time.",
                14,
                False,
                BODY,
            ),
        ),
    ),
    (
        "Fiber is rare",
        (
            (
                "Xfinity fiber is rare and scattered. Non-fiber service, as the salesperson explains it, "
                "is shared with the next-door neighbor, which slows the line at busy hours.",
                14,
                False,
                BODY,
            ),
        ),
    ),
    (
        "Servicing the client",
        (
            (
                "Spectrum wins by being able to service the client. Rural quality falls off more than urban quality "
                "because maintenance lags, so wildlife, storms, and corrosion hit harder.",
                14,
                False,
                BODY,
            ),
        ),
    ),
    (
        "East Coast · interviewee’s claim",
        (
            (
                "Xfinity’s HOA and government deals helped Xfinity’s East Coast position.",
                14,
                False,
                BODY,
            ),
        ),
    ),
)


def emu(inches_val: float) -> int:
    return int(round(inches_val * EMU_IN))


def inches(value: int) -> float:
    return value / EMU_IN


def _font(pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    path = r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"
    return ImageFont.truetype(path, pt * 96 / 72)


def text_width(text: str, pt: float, bold: bool) -> float:
    return _font(pt, bold).getlength(text) / 96


def wrap_lines(text: str, width_in: float, pt: float, bold: bool) -> list[str]:
    """Wrap slightly inside the box so PowerPoint does not add a surprise line."""
    limit = width_in * 0.96
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    cur = words[0]
    for word in words[1:]:
        trial = cur + " " + word
        if text_width(trial, pt, bold) <= limit:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    lines.append(cur)
    for line in lines:
        if text_width(line, pt, bold) > width_in + 0.02:
            raise SystemExit(f"Word wider than box: {line!r}")
    return lines


def _set_run(run, size, bold, color, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False


def _anchor(tf, anchor: str):
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", {"t": "t", "ctr": "ctr", "b": "b"}[anchor])


def _space(paragraph, spacing):
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0 if spacing is None else Pt(spacing)


def add_text(slide, l, t, w, h, text, size, bold, color, align="left", anchor="t", spacing=None):
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    _anchor(tf, anchor)
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    _space(p, spacing)
    run = p.add_run()
    run.text = text
    _set_run(run, size, bold, color)
    return box


def add_runs(slide, l, t, w, h, paragraphs, anchor="t", spacing=None):
    """paragraphs: list of (align, list of (text, size, bold, color))."""
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    _anchor(tf, anchor)
    for i, (align, runs) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        _space(p, spacing)
        for text, size, bold, color in runs:
            run = p.add_run()
            run.text = text
            _set_run(run, size, bold, color)
    return box


def _no_shadow(shape):
    sp_pr = shape._element.spPr
    effect = sp_pr.find(qn("a:effectLst"))
    if effect is not None:
        sp_pr.remove(effect)


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE
    sh.line.width = Emu(6350)
    _no_shadow(sh)
    return sh


def add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    _no_shadow(sh)
    return sh


def add_oval(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(l), Emu(t), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    _no_shadow(sh)
    return sh


def chrome(slide, kicker, title_lines, source, page):
    add_text(slide, MARGIN_L, KICKER_T, TITLE_W, KICKER_H, kicker, 11, True, BLUE, anchor="b")
    add_runs(
        slide,
        MARGIN_L,
        TITLE_T,
        TITLE_W,
        TITLE_H,
        [("left", [(line, 22, True, NAVY)]) for line in title_lines],
        anchor="t",
    )
    add_text(slide, MARGIN_L, SOURCE_T, TITLE_W, SOURCE_H, source, 10, False, MUTED, anchor="t")
    add_text(slide, MARGIN_L, FOOT_T, 10881360, FOOT_H, FOOTER, 9, False, MUTED, anchor="b")
    add_text(slide, PAGE_L, FOOT_T, PAGE_W, FOOT_H, page, 10, False, MUTED, align="right", anchor="b")


def _blank(prs):
    layout = min(prs.slide_layouts, key=lambda item: len(item.placeholders))
    slide = prs.slides.add_slide(layout)
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    return slide


def _runs_text(runs) -> str:
    return "".join(run[0] for run in runs)


def slide_timeline(prs):
    slide = _blank(prs)
    chrome(
        slide,
        "APPENDIX  ·  EH&B  ·  HOW BENEFITS WERE SOLD",
        TITLE_1,
        "Source: Kyle’s note. Dates are approximate industry context from Kyle’s note, "
        "not a verified chronology for one firm.",
        "A9.1",
    )

    gap = emu(0.18)
    card_w = (TITLE_W - 3 * gap) // 4
    inset = emu(0.16)
    body_w_in = inches(card_w - 2 * inset)
    body_pt = 14
    body_spacing = 18  # absolute line height, matched to the text box
    line_counts = [len(wrap_lines(text, body_w_in, body_pt, False)) for _, text in ERAS]
    max_lines = max(line_counts)
    if max_lines > 8:
        raise SystemExit(f"Timeline card copy is too long: {line_counts}")

    label_h = emu(0.50)
    dot = emu(0.18)
    axis_gap = emu(0.10)
    below_axis = emu(0.22)
    text_h = emu(max_lines * body_spacing / 72)
    card_h = text_h + 2 * inset
    group_h = label_h + axis_gap + below_axis + card_h
    # Two 22 pt title lines occupy about 0.78 in. Keep the axis off that and off the source.
    band_top = TITLE_T + emu(0.86)
    band_bot = SOURCE_T - emu(0.18)
    if group_h > band_bot - band_top:
        raise SystemExit(
            f"Timeline group {inches(group_h):.2f} in does not fit "
            f"{inches(band_bot - band_top):.2f} in"
        )
    label_t = band_top + (band_bot - band_top - group_h) // 2
    axis_y = label_t + label_h + axis_gap
    card_t = axis_y + below_axis

    centers = []
    for i, (era, body) in enumerate(ERAS):
        x = MARGIN_L + i * (card_w + gap)
        era_lines = wrap_lines(era, inches(card_w) * 0.96, 14, True)
        if len(era_lines) > 2:
            raise SystemExit(f"Era label wraps too far: {era}")
        add_text(
            slide,
            x,
            label_t,
            card_w,
            label_h,
            era,
            14,
            True,
            NAVY,
            align="center",
            anchor="b",
            spacing=18,
        )
        add_round(slide, x, card_t, card_w, card_h, SOFT)
        add_text(
            slide,
            x + inset,
            card_t + inset,
            card_w - 2 * inset,
            text_h,
            body,
            body_pt,
            False,
            BODY,
            anchor="t",
            spacing=body_spacing,
        )
        centers.append(x + card_w // 2)

    line_h = emu(0.03)
    add_rect(
        slide,
        centers[0] - dot // 2,
        axis_y - line_h // 2,
        (centers[-1] - centers[0]) + dot,
        line_h,
        BLUE,
    )
    for cx in centers:
        add_oval(slide, cx - dot // 2, axis_y - dot // 2, dot, dot, BLUE)

    return line_counts


def slide_interviews(prs):
    slide = _blank(prs)
    chrome(
        slide,
        "APPENDIX  ·  EH&B  ·  INFORMAL INTERVIEWS",
        TITLE_2,
        "Source: Informal interviews with third-party salespeople, 2026 — not a filing or a survey.",
        "A9.2",
    )

    label_w = emu(2.85)
    gutter = emu(0.20)
    text_x = MARGIN_L + label_w + gutter
    text_w = TITLE_W - label_w - gutter
    text_w_in = inches(text_w)
    top = TITLE_T + emu(0.90)
    bottom = SOURCE_T - emu(0.16)
    body_spacing = 18

    measured = []
    for label, runs in ROWS:
        label_lines = wrap_lines(label, inches(label_w) * 0.96, 13, True)
        if len(label_lines) > 2:
            raise SystemExit(f"Row label wraps too far: {label}")
        plain = _runs_text(runs)
        body_lines = wrap_lines(plain, text_w_in, 14, True)
        measured.append((label, runs, len(label_lines), len(body_lines)))

    line_h = emu(body_spacing / 72)
    label_line_h = emu(16 / 72)
    heights = []
    for _, _, n_lab, n_body in measured:
        heights.append(max(n_body * line_h, n_lab * label_line_h, emu(0.32)))
    slack = (bottom - top) - sum(heights)
    gaps = len(ROWS) - 1
    if slack < gaps * emu(0.06):
        raise SystemExit(
            f"Interview rows overflow: slack {inches(slack):.2f} in"
        )
    row_gap = slack // gaps

    y = top
    for i, ((label, runs, _, _), h) in enumerate(zip(measured, heights)):
        add_text(
            slide,
            MARGIN_L,
            y,
            label_w,
            h,
            label,
            13,
            True,
            NAVY,
            anchor="t",
            spacing=16,
        )
        add_runs(
            slide,
            text_x,
            y,
            text_w,
            h,
            [("left", list(runs))],
            anchor="t",
            spacing=body_spacing,
        )
        y += h
        if i < len(ROWS) - 1:
            rule_y = y + row_gap // 2
            add_rect(slide, MARGIN_L, rule_y, TITLE_W, emu(0.01), LINE)
            y += row_gap

    return [n for _, _, _, n in measured]


def assert_copy():
    joined = " ".join(body for _, body in ERAS)
    for banned in ("Identify the prevailing", "privacy", "presenteeism", "54.1", "Form 5500"):
        if banned.lower() in joined.lower():
            raise SystemExit(f"Timeline picked up material that does not belong: {banned}")
    interview = " ".join(_runs_text(runs) for _, runs in ROWS)
    for banned in ("19", "27", "FCC", "Census", "percent", "%"):
        if banned in interview:
            raise SystemExit(f"Interview slide mixed in a filing fact: {banned}")


def main():
    assert_copy()
    for line in TITLE_1 + TITLE_2:
        width = text_width(line, 22, True)
        if width > inches(TITLE_W):
            raise SystemExit(f"Title wider than the slide: {line!r} ({width:.2f} in)")

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    prs.core_properties.title = "EH&B mindset and sales interviews"
    era_lines = slide_timeline(prs)
    row_lines = slide_interviews(prs)
    if len(prs.slides) != 2:
        raise SystemExit(f"Expected 2 slides, got {len(prs.slides)}")
    prs.save(OUT)
    print(f"saved {OUT}")
    print(f"timeline lines per card: {era_lines}")
    print(f"interview lines per row: {row_lines}")
    print("TITLE 1:", " / ".join(TITLE_1))
    print("TITLE 2:", " / ".join(TITLE_2))


if __name__ == "__main__":
    main()
